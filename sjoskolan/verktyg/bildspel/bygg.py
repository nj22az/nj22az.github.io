#!/usr/bin/env python3
"""Bygger webbildspel av veckornas presentationer.

För varje vecka-XX/aktuell/*.pptx (med en PDF bredvid, se veckosidorna):
- renderar varje bild ur PDF:en till bildspel/<id>/NN.webp (1280 px bred),
- hämtar bildernas text ur PowerPoint-filen, så att texten går att läsa, söka och använda med skärmläsare,
- kopplar bilder med ”övning N” i rubriken till motsvarande övning i Formelstöd och övningar,
- skriver bildspel/<id>/data.json som visaren (bildspel/index.html) läser.

    python3 sjoskolan/verktyg/bildspel/bygg.py            # bygger om bara ändrade presentationer
    python3 sjoskolan/verktyg/bildspel/bygg.py --alla     # bygger om allt
"""
import hashlib
import io
import json
import re
import sys
from pathlib import Path

import pymupdf
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'bildspel'
WIDTH = 1280
SKIP_PH = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE}


def shapes_flat(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shapes_flat(sh.shapes)
        else:
            yield sh


def run_text(runs):
    """Styckets text med innehållsmarkering: nedsänkta körningar blir _{…}, upphöjda ^{…}."""
    out = []
    for r in runs:
        rpr = r._r.find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
        base = int(rpr.get('baseline', '0')) if rpr is not None else 0
        out.append(f'_{{{r.text}}}' if base < 0 and r.text.strip() else f'^{{{r.text}}}' if base > 0 and r.text.strip() else r.text)
    return ''.join(out)


def slide_text(slide):
    title, items = '', []
    for sh in sorted(shapes_flat(slide.shapes), key=lambda s: ((s.top or 0) // 200000, s.left or 0)):
        if sh.is_placeholder and sh.placeholder_format.type in SKIP_PH:
            continue
        if sh.has_text_frame:
            paras = [run_text(p.runs).strip() for p in sh.text_frame.paragraphs]
            paras = [p for p in paras if p and not re.fullmatch(r'\d{1,3}', p)]
            is_title = sh.is_placeholder and sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
            if is_title and paras and not title:
                title = ' '.join(paras)
            else:
                items += paras
        elif getattr(sh, 'has_table', False) and sh.has_table:
            for row in sh.table.rows:
                cells = [t for t in (' '.join(run_text(p.runs) for p in c.text_frame.paragraphs).strip() for c in row.cells) if t]
                if cells:
                    items.append(' · '.join(cells))
    if not title and items:
        title = items.pop(0)
    return title, items


def _titles():
    """Titlar från veckosidornas data (verktyg/veckosidor/bygg.py) och vecka 40:s lektioner."""
    import importlib.util
    t = {}
    spec = importlib.util.spec_from_file_location('veckor', ROOT / 'verktyg' / 'veckosidor' / 'bygg.py')
    mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)
    for w in mod.VECKOR.values():
        for p in [x for d in w['delar'] for x in d['poster']] + w.get('fordjupning', []):
            if p.get('typ') == 'ppt':
                t[p['fil']] = p['titel']
    lek = ROOT / 'vecka-40' / 'aktuell' / 'lektioner.mjs'
    if lek.exists():
        for title, deck in re.findall(r"title:'([^']+)',deck:'([^']+)'", lek.read_text(encoding='utf-8')):
            t[deck] = title
    return t


TITLES = _titles()


def deck_title(pptx, first_slide_title):
    return TITLES.get(pptx.name) or first_slide_title or re.sub(r'^v\d+_\d+_|_elev$', '', pptx.stem).replace('_', ' ')


def exercise_ids(week_dir):
    f = week_dir / 'Formelstod_och_ovningar.html'
    return set(re.findall(r'<article class="exercise" id="([^"]+)"', f.read_text(encoding='utf-8'))) if f.exists() else set()


def build(pptx, force=False):
    week_dir = pptx.parent
    week = int(re.search(r'vecka-(\d+)', str(pptx)).group(1))
    pdf = pptx.with_suffix('.pdf')
    did = pptx.stem.removesuffix('_elev')
    out = OUT / did
    stamp = hashlib.sha1(pptx.read_bytes() + pdf.read_bytes()).hexdigest()[:16]
    data_file = out / 'data.json'
    if not force and data_file.exists() and json.loads(data_file.read_text()).get('stamp') == stamp:
        return did, False
    out.mkdir(parents=True, exist_ok=True)
    for old in out.glob('*.webp'):
        old.unlink()
    doc = pymupdf.open(pdf)
    prs = Presentation(pptx)
    slides_x = [s for s in prs.slides if s._element.get('show') != '0']
    if len(slides_x) != doc.page_count:
        raise SystemExit(f'{pptx.name}: {len(slides_x)} bilder i PowerPoint men {doc.page_count} sidor i PDF')
    chapter = re.match(r'(v\d+_\d+)', did).group(1)
    ex_ids = exercise_ids(week_dir)
    slides = []
    for i, (page, sx) in enumerate(zip(doc, slides_x), 1):
        z = WIDTH / page.rect.width
        pix = page.get_pixmap(matrix=pymupdf.Matrix(z, z))
        img = Image.open(io.BytesIO(pix.tobytes('png')))
        img.save(out / f'{i:02d}.webp', 'WEBP', quality=80, method=6)
        title, text = slide_text(sx)
        entry = {'title': title or f'Bild {i}', 'text': text}
        m = re.search(r'övning\s+(\d+)', title, re.I)
        if m and f'{chapter}-q{m.group(1)}' in ex_ids:
            entry['ex'] = f'../vecka-{week}/aktuell/Formelstod_och_ovningar.html#{chapter}-q{m.group(1)}'
        slides.append(entry)
    data = {
        'id': did, 'week': week, 'title': deck_title(pptx, slides[0]['title']), 'stamp': stamp,
        'w': img.width, 'h': img.height,
        'pdf': f'../vecka-{week}/aktuell/{pdf.name}', 'pptx': f'../vecka-{week}/aktuell/{pptx.name}',
        'weekUrl': f'../vecka-{week}/aktuell/',
        'exercises': f'../vecka-{week}/aktuell/Formelstod_och_ovningar.html#{chapter}' if any(e.startswith(chapter) for e in ex_ids) else None,
        'slides': slides,
    }
    data_file.write_text(json.dumps(data, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    return did, True


if __name__ == '__main__':
    force = '--alla' in sys.argv
    decks = []
    for pptx in sorted(ROOT.glob('vecka-*/aktuell/*.pptx')):
        if not pptx.with_suffix('.pdf').exists():
            print('saknar PDF, hoppar över:', pptx.name)
            continue
        did, built = build(pptx, force)
        d = json.loads((OUT / did / 'data.json').read_text())
        decks.append({'id': did, 'week': d['week'], 'title': d['title'], 'n': len(d['slides'])})
        print(('byggd  ' if built else 'oförändrad ') + did, len(d['slides']))
    (OUT / 'lista.json').write_text(json.dumps(decks, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
