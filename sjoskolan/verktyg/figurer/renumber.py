"""Sätt tryckta bildnummer efter bildernas faktiska ordning.
Hanterar sidfötter som ”Föreläsning 1A · 10 / 52” och rena nummerrutor (”10”).
Användning: renumber.py fil.pptx [fler.pptx]  (skriver över filerna)"""
import re, sys
from pptx import Presentation
FOOT = re.compile(r"^(.*·\s*)\d+\s*/\s*\d+\s*$")
ONLY = re.compile(r"^\s*\d+\s*$")
for f in sys.argv[1:]:
    p = Presentation(f); total = len(p.slides); n = 0
    for i, s in enumerate(p.slides, 1):
        for sh in s.shapes:
            if not sh.has_text_frame: continue
            runs = [r for para in sh.text_frame.paragraphs for r in para.runs]
            if len(runs) != 1: continue
            r = runs[0]; m = FOOT.match(r.text)
            if m:
                new = f"{m.group(1)}{i} / {total}"
            elif ONLY.match(r.text) and sh.top > p.slide_height * 0.85:  # nummerruta i sidfoten
                new = str(i)
            else: continue
            if new != r.text: r.text = new; n += 1
    p.save(f); print(f.split("/")[-1], total, "bilder,", n, "nummer ändrade")
