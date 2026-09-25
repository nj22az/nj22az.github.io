"""Hjälpfunktioner för att ändra text i befintliga presentationer med bevarat format."""
import sys, copy
from pptx import Presentation
from pptx.util import Pt

def shape(slide, sid):
    return next(s for s in slide.shapes if s.shape_id == sid)

def set_paras(sh, texts, size=None):
    """Ersätt styckenas text men behåll formatet från första stycket/körningen."""
    tf = sh.text_frame
    paras = tf.paragraphs
    tmpl = copy.deepcopy(paras[0]._p)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in first.runs[1:]:
        r._r.getparent().remove(r._r)
    first.runs[0].text = texts[0]
    for t in texts[1:]:
        np_ = copy.deepcopy(tmpl)
        tf._txBody.append(np_)
        from pptx.text.text import _Paragraph
        para = _Paragraph(np_, tf)
        for r in para.runs[1:]:
            r._r.getparent().remove(r._r)
        para.runs[0].text = t
    if size:
        for p in tf.paragraphs:
            for r in p.runs: r.font.size = Pt(size)

def add_para(sh, text, size=None, bold_prefix=None):
    tf = sh.text_frame
    src = tf.paragraphs[-1]._p
    np_ = copy.deepcopy(src)
    src.addnext(np_)
    from pptx.text.text import _Paragraph
    para = _Paragraph(np_, tf)
    for r in para.runs[1:]:
        r._r.getparent().remove(r._r)
    run = para.runs[0]
    if bold_prefix:
        run.text = bold_prefix
        run.font.bold = True
        r2 = copy.deepcopy(run._r); run._r.addnext(r2)
        para.runs[1].text = text; para.runs[1].font.bold = False
    else:
        run.text = text
    if size:
        for r in para.runs: r.font.size = Pt(size)

def drop_notes_line(prs):
    last = prs.slides[len(prs.slides) - 1]
    for sh in list(last.shapes):
        if sh.has_text_frame and 'bildanteckningarna' in sh.text_frame.text:
            sh._element.getparent().remove(sh._element)


def replace_words(prs, pairs):
    """Byt ord i alla textkörningar (formatet behålls). pairs: [(gammalt, nytt), ...] i ordning."""
    n = 0
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame: continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    t = r.text
                    for a, b in pairs: t = t.replace(a, b)
                    if t != r.text: r.text = t; n += 1
    return n


def subscript_tokens(prs, tokens):
    """Skriv index som riktigt nedsänkt text, t.ex. {'Uberöring': ('U', 'beröring')}. Hela ord matchas."""
    import copy, re
    from pptx.oxml.ns import qn
    pat = re.compile(r"(?<![\wåäöÅÄÖ])(" + "|".join(map(re.escape, tokens)) + r")(?![\wåäöÅÄÖ])")
    n = 0
    for s in prs.slides:
        for r in list(s._element.iter(qn("a:r"))):
            t = r.find(qn("a:t"))
            if t is None or not t.text or not pat.search(t.text): continue
            pieces = []
            for part in pat.split(t.text):
                if not part: continue
                if part in tokens: pieces += [(tokens[part][0], False), (tokens[part][1], True)]
                else: pieces.append((part, False))
            prev = r
            for txt, low in pieces:
                nr = copy.deepcopy(r); nr.find(qn("a:t")).text = txt
                if low:
                    pr = nr.find(qn("a:rPr"))
                    if pr is None: pr = nr.makeelement(qn("a:rPr"), {}); nr.insert(0, pr)
                    pr.set("baseline", "-25000")
                prev.addnext(nr); prev = nr
            r.getparent().remove(r); n += 1
    return n
