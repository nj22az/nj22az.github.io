"""Innehållsändringar i vecka 42 efter granskning (lärare, elinstallatör, sjöingenjör):
exempel och stödbilder som inte avslöjar övningssvar, fartygets larmvägar, enhetliga begrepp,
nedsänkta index och ingen hänvisning till saknade bildanteckningar.
Användning: textfix42.py <in-mapp> <ut-mapp>"""
import sys
from pptx import Presentation
from pptxtext import shape, set_paras, drop_notes_line, subscript_tokens, replace_words, replace_links
SUBS = {'Uberöring': ('U', 'beröring'), 'VA': ('V', 'A'), 'VB': ('V', 'B'), 'Ifel': ('I', 'fel'), 'Rslinga': ('R', 'slinga')}
WORDS = [('skydd mot tillkoppling', 'skydd mot återinkoppling'), ('Skydd mot tillkoppling', 'Skydd mot återinkoppling'),
         ('Landmatningen', 'Landanslutningen')]
REDOVISA = 'Redovisa din beräkning eller motivera ditt beslut.'

def reasoning_only(prs, calc):
    """Övningar utan beräkning får en enklare slutrad."""
    for n, s in enumerate(prs.slides, 1):
        if n in calc: continue
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text == REDOVISA: r.text = 'Motivera ditt beslut.'

OLD_TS = "https://www.transportstyrelsen.se/sv/sjofart/Fartyg/Fartygskonstruktion/Elinstallationer/Lagar-foreskrifter-och-standarder/"
NEW_TS = "https://www.transportstyrelsen.se/sv/om-oss/dina-rattigheter-lagar-och-regler/lagar-och-regler/regler-for-sjofart/regler-for-nationell-sjofart/regler-kompletterade-upplysningar/elektrisk-utrustning-och-elinstallationer/"  # den gamla adressen ger 404
src, dst = sys.argv[1], sys.argv[2]

f = "v42_01_Elektriska_risker_elev.pptx"; p = Presentation(f"{src}/{f}")
# Exemplen får inte vara övningarnas fall (granskning 2026-10-11): skadad kabel i vatten är övning 3.2, skadad prob är 3.5
set_paras(shape(p.slides[8], 10), ["En vinkelslip i verkstaden har en kabel som legat klämd under en tung låda. Ingen använder maskinen just nu."], size=18)
set_paras(shape(p.slides[8], 12), ["1. Observation: kabeln är tillplattad och manteln har en spricka.", "2. Möjlig händelse: någon håller i maskinen och rör en del där isolationen är trasig.", "3. Möjlig följd: strömgenomgång, eller att personen tappar maskinen och skadar sig."], size=15)
set_paras(shape(p.slides[8], 13), ["Ingen vet om kabeln är spänningssatt. Den synliga skadan räcker för att stoppa användningen."], size=14)
set_paras(shape(p.slides[23], 10), ["Ett låsbleck för frånskiljaren är böjt och går inte att stänga. Det upptäcks innan arbetet börjar."])
set_paras(shape(p.slides[23], 12), ["1. Ta det böjda låsblecket ur bruk och märk det enligt rutinen.", "2. Hämta ett helt låsbleck innan frånskiljaren låses.", "3. Följ upp hur skadan uppstod och om låsutrustningen kontrolleras före användning."])
set_paras(shape(p.slides[21], 3), ["Skydda dig själv och bryt energitillförseln om det kan ske säkert.", "Larma enligt fartygets nödinstruktion. I hamn: 112."])
set_paras(shape(p.slides[21], 4), ["Till sjöss: radiomedicinsk rådgivning (TMAS)"])
set_paras(shape(p.slides[21], 5), ["När platsen är säker: bedöm andning och ge HLR vid behov.", "Vid högspänning håll avstånd tills behörig person säkrat platsen."])
set_paras(shape(p.slides[24], 13), ["Arbetsgång: Vad händer i slingan när Rslinga är liten? Hitta inte på ett amperetal."])
set_paras(shape(p.slides[26], 12), ["Förutsättningar: Platsen kan vara spänningssatt. Larm enligt nödinstruktionen. Vid högspänning säkrar behörig person platsen."])
set_paras(shape(p.slides[28], 13), ["Arbetsgång: Vad rekommenderar du, och varför? Beskriv även rapporteringen enligt rutinen ombord."])
set_paras(shape(p.slides[34], 3), ["Vilka uppgifter saknas oftast när någon bara säger ”det är 24 V”?"])
reasoning_only(p, calc={17})
subscript_tokens(p, SUBS); replace_words(p, WORDS)
replace_links(p, OLD_TS, NEW_TS); drop_notes_line(p); p.save(f"{dst}/{f}")

f = "v42_02_Regler_ansvar_och_arbetsmetoder_elev.pptx"; p = Presentation(f"{src}/{f}")
# Exemplet får inte lista svaret på övning 2.7 (vilka fartygsdata som behövs)
set_paras(shape(p.slides[8], 10), ["Ett fartyg ska få en ny landanslutning. Underlaget säger bara ”enligt gällande regler”."])
set_paras(shape(p.slides[8], 12), ["1. Skriv upp vilka uppgifter om fartyget som saknas i underlaget. Gissa dem inte.", "2. Matcha uppgifterna mot föreskriftens tillämpningsområde.", "3. Dokumentera aktuell version och det avsnitt du har kontrollerat."])
set_paras(shape(p.slides[6], 5), ["SS-EN 50110-1: tre metoder (utan spänning, med spänning, nära spänning).", "ESA: fyra metoder (kompletterar med arbete utanför närområde)."])
set_paras(shape(p.slides[13], 10), ["Metoderna skiljer sig i hur nära spänning arbetet sker"])
set_paras(shape(p.slides[28], 10), ["Kontrollera föreskriftens status i regellistan"])
set_paras(shape(p.slides[28], 11), ["Transportstyrelsens regellista visar om en föreskrift gäller eller är upphävd. Ett nytt nummer kräver ändå kontroll av tillämpningen."])
reasoning_only(p, calc=set())
replace_words(p, WORDS)
replace_links(p, OLD_TS, NEW_TS); drop_notes_line(p); p.save(f"{dst}/{f}")

f = "v42_03_Riskbedomning_och_skydd_elev.pptx"; p = Presentation(f"{src}/{f}")
# Exemplet får inte ge något svar i övning 3.8 (12 och 4), och poängen är ingen kvotskala. Samma exempel på övningssidan.
set_paras(shape(p.slides[23], 10), ["En övningsmatris anger sannolikhet 5, konsekvens 3. Efter åtgärd bedöms sannolikheten till 2."])
set_paras(shape(p.slides[23], 12), ["1. Före: 5 · 3 = 15.", "2. Efter: 2 · 3 = 6.", "3. Poängen sjunker i denna matris. Det säger inte hur mycket säkrare arbetet har blivit."])
set_paras(shape(p.slides[11], 11), ["Fara är det som kan skada. Händelsen är det som kan inträffa. Konsekvensen är skadan som kan följa."])
set_paras(shape(p.slides[15], 11), ["En skylt visar ett budskap. Den hindrar inte automatiskt någon från att slå till."])
set_paras(shape(p.slides[32], 11), ["Jämför det som är verifierat med det som den valda arbetsmetoden kräver."])
set_paras(shape(p.slides[7], 3), ["Tvåpolig spänningsprovare väljs för systemets spänning och kategori.", "Låsning och märkning stöder skydd mot återinkoppling."])
reasoning_only(p, calc={30})
replace_words(p, WORDS)
replace_links(p, OLD_TS, NEW_TS); drop_notes_line(p); p.save(f"{dst}/{f}")
