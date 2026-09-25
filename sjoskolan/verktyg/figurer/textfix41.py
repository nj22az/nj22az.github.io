"""Innehållsändringar i vecka 41 (samma granskningsregler som vecka 40): exempel får inte avslöja
övningssvar, fartygsexempel där det passar, ingen hänvisning till saknade bildanteckningar.
Användning: textfix41.py <in-mapp> <ut-mapp>"""
import sys
from pptx import Presentation
from pptxtext import shape, set_paras, add_para, drop_notes_line, replace_words
WORDS = [('Trefastränaren', 'Trefasriggen'), ('trefastränare', 'trefasrigg'), ('SELV-tränare', 'SELV-rigg'), ('Avsedd tränare', 'Avsedd övningsrigg'),
         ('AC-tränarens', 'AC-riggens'), ('tränarens', 'riggens'), ('tränarkälla', 'övningskälla'), ('symmetrisk tränare', 'symmetrisk trefasrigg'),
         ('tränare', 'rigg'), ('källhöjning', 'källspänning'), ('skiljande kontroll', 'kontroll som skiljer mellan orsakerna'), ('Beräkna vinkel mellan', 'Beräkna vinkeln mellan')]
src, dst = sys.argv[1], sys.argv[2]

f = "v41_01_Trefassystemets_grunder_elev.pptx"; p = Presentation(f"{src}/{f}")
set_paras(shape(p.slides[6], 5), ["Ombord är 440 V vanligt. Generatorns fasspänning i Y är ≈ 254 V,", "men nätet har oftast ingen neutralledare. Vid ett jordfel kan en fas ha hela 440 V mot skrovet."])
set_paras(shape(p.slides[24], 13), ["Arbetsgång: Rita de tre visarna efter varandra, spets mot fot. Vad blir summan?"])
set_paras(shape(p.slides[26], 10), ["Nodlagen: Σ in = Σ ut"])
set_paras(shape(p.slides[26], 11), ["Iɴ är neutralledarens ström. Den för tillbaka strömmen från fas–neutral-lasterna."])
replace_words(p, WORDS)
drop_notes_line(p); p.save(f"{dst}/{f}")

f = "v41_02_Y_och_trefaseffekt_elev.pptx"; p = Presentation(f"{src}/{f}")
set_paras(shape(p.slides[6], 5), ["440 V över tre grenar på 55 Ω i Δ:", "Igren = 8,0 A och Iᴸ ≈ 13,9 A."])
set_paras(shape(p.slides[7], 5), ["440 V, 20 A och cos φ = 0,85 ger", "S ≈ 15,2 kVA och P ≈ 13,0 kW."])
set_paras(shape(p.slides[8], 10), ["Tre resistorer på 16 Ω ansluts till symmetriska 200 V mellan faserna."])
set_paras(shape(p.slides[8], 12), ["1. Y: Ugren = 200/√3 ≈ 115,5 V. Iᴸ ≈ 115,5/16 ≈ 7,22 A.", "2. Δ: Ugren = 200 V. Igren = 200/16 = 12,5 A.", "3. Δ: Iᴸ = √3 · 12,5 ≈ 21,65 A."])
set_paras(shape(p.slides[26], 10), ["Jämför nätets Uᴸ med båda märkvärdena"])
set_paras(shape(p.slides[26], 11), ["Δ/Y-märkningen anger vilken linjespänning som gäller för respektive koppling."])
replace_words(p, WORDS)
set_paras(shape(p.slides[20], 4), ["Nätets Uᴸ = högre märkvärde: Y"])
set_paras(shape(p.slides[20], 5), ["Nätets Uᴸ = lägre märkvärde: Δ.", "Ombord: en motor Δ/Y 440/760 V på 440 V-nät kopplas i Δ."])
set_paras(shape(p.slides[21], 5), ["Paxel = 11 kW, η = 0,90 ger Pᵢₙ ≈ 12,2 kW.", "Vid 440 V och PF = 0,85 blir Iᴸ ≈ 18,9 A."])
drop_notes_line(p); p.save(f"{dst}/{f}")

f = "v41_03_Fysisk_traff_och_matning_elev.pptx"; p = Presentation(f"{src}/{f}")
set_paras(shape(p.slides[6], 4), ["I = U/(R₁ + R₂)"])
set_paras(shape(p.slides[6], 5), ["Beräkna I och U₂ i övning 2 före mätningen.", "Instruktören granskar mätplan och anslutningar före start."])
set_paras(shape(p.slides[7], 4), ["Ugren = Uᴸ/√3 i symmetrisk Y"])
set_paras(shape(p.slides[7], 5), ["Förberedelse: övning 4 (AC) och 5 (trefas).", "Saknas trefasrigg: räkna på övningsdata eller i Trefaslabbet och dokumentera det som teori."])
set_paras(shape(p.slides[32], 13), ["Arbetsgång: Visa summan och jämför med källan. Vad visar jämförelsen, och vad visar den inte?"])
replace_words(p, WORDS)
drop_notes_line(p); p.save(f"{dst}/{f}")
print("ok")
