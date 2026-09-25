import sys, math, json, copy, re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
E=914400
def fit_box(shape, min_pt=16):
    """shrink run sizes until estimated text height fits the shape"""
    tf=shape.text_frame; w=shape.width/E; h=shape.height/E
    runs=[r for p in tf.paragraphs for r in p.runs]
    base=[r.font.size.pt if r.font.size else 24 for r in runs]
    if not runs: return
    scale=1.0
    while True:
        tot=0
        for p in tf.paragraphs:
            txt="".join(r.text for r in p.runs)
            sz=max([ (r.font.size.pt if r.font.size else 24) for r in p.runs] or [24])*scale if p.runs else 24*scale
            # current sizes are original; apply scale
            cpl=max(1,w*72/(sz*0.55))
            lines=max(1,math.ceil(len(txt)/cpl)) if txt else 1
            tot+=lines*sz*1.22/72
        if tot<=h or min(base)*scale<=min_pt: break
        scale-=0.03
    if scale<1.0:
        for r,b in zip(runs,base): r.font.size=Pt(round(b*scale*2)/2)
    return scale
UNIT=re.compile(r"(\d) (?=(?:V|Ω|ms|s|A|W|Hz|kW|kVA|VA|var|kvar|µF|mH|H|h|°|%|kWh|Wh)\b|(?:Ω|°|%))")
def keep_units(shape):
    """Icke-brytande mellanslag mellan tal och enhet samt efter likhetstecken."""
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            t=re.sub(r"(\d) (?=\d{3}\b)","\\1\u00a0",UNIT.sub("\\1\u00a0",r.text)).replace(" = ","\u00a0=\u00a0").replace("cos φ","cos\u00a0φ")
            if t!=r.text: r.text=t
def add_pic(slide,path,x,y,w,alt,max_h=None):
    from PIL import Image
    iw,ih=Image.open(path).size
    if max_h and w*ih/iw>max_h:  # höga figurer krymps så att de inte når ”Redovisa”-raden
        nw=max_h*iw/ih; x+= (w-nw)/2; w=nw
    pic=slide.shapes.add_picture(path,Inches(x),Inches(y),width=Inches(w))
    cNvPr=pic._element.nvPicPr.cNvPr; cNvPr.set("descr",alt); cNvPr.set("name","Figur: "+alt[:40])
    return pic
def process(src,dst,spec,figdir):
    prs=Presentation(src); log=[]
    for n,(kind,fname,alt) in spec.items():
        s=prs.slides[int(n)-1]
        if any(sh.name.startswith("Figur:") for sh in s.shapes):
            log.append(f"{n} har redan en figur, hoppar över"); continue
        texts=[sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        title=texts[0].text_frame.text
        body=[sh for sh in texts if sh.left/E<1.0 and sh.width/E>8 and sh.top/E>1.6]
        if kind=="ovn":
            tb=[sh for sh in body if 2.5<sh.top/E<3.0][0]
            tb.width=Inches(4.25); tb.height=Inches(2.85); keep_units(tb)
            sc=fit_box(tb)
            add_pic(s,f"{figdir}/{fname}.png",5.3,2.55,4.3,alt,max_h=3.1)
        else:
            scs=[]
            for sh in body:
                sh.width=Inches(4.35); keep_units(sh); scs.append(fit_box(sh))
            sc=min([x for x in scs if x] or [1])
            add_pic(s,f"{figdir}/{fname}.png",5.35,1.8,4.25,alt,max_h=4.6)
        log.append(f"{n} {title!r} scale={sc}")
    prs.save(dst); return log
if __name__=="__main__":
    spec=json.load(open(sys.argv[1]))
    for deck,sp in spec.items():
        print(deck); [print("  ",l) for l in process(sp["src"],sp["dst"],sp["slides"],sys.argv[2])]
