"""Innehållsändringar i vecka 44: teoriexempel som inte ger samma svar som övningarna,
nedsänkta index och ingen hänvisning till saknade bildanteckningar.
Användning: textfix44.py <in-mapp> <ut-mapp>"""
import sys
from pptx import Presentation
from pptxtext import shape, set_paras, drop_notes_line, subscript_tokens, replace_words
SUBS = {'Zfel': ('Z', 'fel'), 'Ifel': ('I', 'fel'), 'Rslinga': ('R', 'slinga'), 'Ulast': ('U', 'last'), 'Ukälla': ('U', 'källa'),
        'Ilåg': ('I', 'låg'), 'Ihög': ('I', 'hög'), 'Uhög': ('U', 'hög'), 'Ulåg': ('U', 'låg'), 'Utång': ('U', 'tång'),
        'Uprov': ('U', 'prov'), 'Iläck': ('I', 'läck'), 'ûverklig': ('û', 'verklig'), 'ûvisad': ('û', 'visad'),
        'xmax': ('x', 'max'), 'xmin': ('x', 'min'), 'kI': ('k', 'I'), 'Zs': ('Z', 's')}
src, dst = sys.argv[1], sys.argv[2]

f = "v44_01_Lagspanningssystem_elev.pptx"; p = Presentation(f"{src}/{f}")
set_paras(shape(p.slides[22], 5), ["230 V och Zs = 0,32 Ω ger cirka 720 A.", "Verklig dimensionering kräver skyddsdata och fler villkor."])
set_paras(shape(p.slides[23], 12), ["1. Iₖ ≈ 230/0,575 = 400 A.", "2. Om Zs ökar till 2,3 Ω blir Iₖ ≈ 230/2,3 = 100 A.", "3. Jämför båda strömmarna med aktuellt skydds utlösningskurva."])
set_paras(shape(p.slides[27], 3), ["Förenklad felström"])
set_paras(shape(p.slides[7], 5), ["Jordtagsresistanser kan begränsa felströmmen så att överströmsskydd ensamt inte ger avsedd bortkoppling.", "TT-system är ovanliga ombord."])
subscript_tokens(p, SUBS); drop_notes_line(p); p.save(f"{dst}/{f}")

f = "v44_02_Hogspanningssystem_elev.pptx"; p = Presentation(f"{src}/{f}")
set_paras(shape(p.slides[18], 4), ["Ordna dessa funktioner i en signalväg: brytare, mättransformator, skyddsrelä. Vad händer om skyddsreläet saknar hjälpspänning?"])
replace_words(p, [('UᴸPF', 'Uᴸ · PF')])
set_paras(shape(p.slides[23], 10), ["En ideal strömtransformator är märkt 300/5 A. Primärströmmen är 150 A. Endast beräkning på data."])
set_paras(shape(p.slides[23], 12), ["1. kI = 300/5 = 60.", "2. I₂ = 150/60 = 2,5 A.", "3. Kontroll: 2,5/5 = 150/300 = 0,50 av märkström på båda sidor."])
subscript_tokens(p, SUBS); drop_notes_line(p); p.save(f"{dst}/{f}")

f = "v44_03_Fordjupad_matteknik_elev.pptx"; p = Presentation(f"{src}/{f}")
set_paras(shape(p.slides[5], 5), ["Alla faser genom tången: summan är läckströmmen. Så spåras jordfel matning för matning.", "Vid frekvensomriktare: true RMS och mätning på nätsidan."])
set_paras(shape(p.slides[6], 5), ["Verifiera spänningslöshet. Koppla bort isolationsvakt, omriktare och elektronik.", "Normalt 500 V DC för 440 V-nät. Urladda till jord efter provningen."])
set_paras(shape(p.slides[20], 5), ["Flytande mätningar kräver differentialprob med rätt mätkategori.", "I IT-nät ger probjord på en fas ett jordfel och larm, och kortslutning om ett annat fel redan finns."])
set_paras(shape(p.slides[21], 5), ["1 % av 50,00 V + 3 steg à 0,01 V ger ±0,53 V.", "En full mätosäkerhetsbudget kräver en beskriven modell."])
subscript_tokens(p, SUBS); drop_notes_line(p); p.save(f"{dst}/{f}")
