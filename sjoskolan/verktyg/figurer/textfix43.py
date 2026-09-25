"""Innehållsändringar i vecka 43: teoriexempel som inte avslöjar övningssvar (fartygsvärden 440 V/60 Hz),
nedsänkta index och ingen hänvisning till saknade bildanteckningar.
Användning: textfix43.py <in-mapp> <ut-mapp>"""
import sys
from pptx import Presentation
from pptxtext import shape, set_paras, drop_notes_line, subscript_tokens
SUBS = {'IΔ': ('I', 'Δ'), 'Iut': ('I', 'ut'), 'Iretur': ('I', 'retur'), 'Paxel': ('P', 'axel'), 'Pförlust': ('P', 'förlust'),
        'Ulast': ('U', 'last'), 'Ukontakt': ('U', 'kontakt'), 'Ispole': ('I', 'spole'), 'Uspole': ('U', 'spole'), 'Rspole': ('R', 'spole'),
        'Vmatningssida': ('V', 'matningssida'), 'Vretur': ('V', 'retur'), 'Vföre': ('V', 'före'), 'Vefter': ('V', 'efter'),
        'Uöppen kontakt': ('U', 'öppen kontakt'), 'Ugren': ('U', 'gren')}
src, dst = sys.argv[1], sys.argv[2]

f = "v43_01_Komponenter_och_skydd_elev.pptx"; p = Presentation(f"{src}/{f}")
subscript_tokens(p, SUBS); drop_notes_line(p); p.save(f"{dst}/{f}")

f = "v43_02_Transformatorer_och_motorer_elev.pptx"; p = Presentation(f"{src}/{f}")
set_paras(shape(p.slides[5], 5), ["440 V till 110 V ger omsättningen 4:1.", "2 A på sekundären motsvarar idealt 0,50 A på primären."])
set_paras(shape(p.slides[6], 5), ["Ombord: 60 Hz och fyra poler ger nₛ = 1 800 r/min.", "Här är p antalet poler, inte antalet polpar."])
set_paras(shape(p.slides[7], 5), ["nₛ = 1 800 och n = 1 746 r/min ger s = 0,03 = 3 %.", "Eftersläpningen varierar med belastningen."])
set_paras(shape(p.slides[20], 5), ["Ombord: en 440/760 V Δ/Y-motor kan vara kandidat vid 440 V.", "Y-start minskar också tillgängligt startmoment."])
set_paras(shape(p.slides[23], 10), ["En åttapolig motor matas med 60 Hz. Rotorn går med 873 r/min."])
set_paras(shape(p.slides[23], 12), ["1. nₛ = 120 · 60/8 = 900 r/min.", "2. Skillnad = 900 − 873 = 27 r/min.", "3. s% = 100 · 27/900 = 3 %."])
set_paras(shape(p.slides[23], 13), ["Rotorn går långsammare än fältet vid den angivna motordriften. p är åtta poler, inte fyra polpar."])
subscript_tokens(p, SUBS); drop_notes_line(p); p.save(f"{dst}/{f}")

f = "v43_03_Elscheman_och_dokumentation_elev.pptx"; p = Presentation(f"{src}/{f}")
set_paras(shape(p.slides[21], 4), ["Full matningsspänning över en öppen kontakt kan vara normalt"])
subscript_tokens(p, SUBS); drop_notes_line(p); p.save(f"{dst}/{f}")
