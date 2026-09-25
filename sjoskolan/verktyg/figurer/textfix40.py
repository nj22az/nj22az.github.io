"""Innehållsändringar i vecka 40 efter granskningen 25 september 2026.
Körs på originalpresentationerna före insert.py. Användning: textfix40.py <in-mapp> <ut-mapp>"""
import sys, copy
from pptx import Presentation
from pptx.util import Pt

def shape(slide, sid):
    return next(s for s in slide.shapes if s.shape_id == sid)

def set_paras(sh, texts, size=None):
    """Ersätt styckenas text men behåll formatet från första stycket/körningen."""
    tf = sh.text_frame
    paras = tf.paragraphs
    tmpl = copy.deepcopy(paras[0]._p)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in first.runs[1:]:
        r._r.getparent().remove(r._r)
    first.runs[0].text = texts[0]
    for t in texts[1:]:
        np_ = copy.deepcopy(tmpl)
        tf._txBody.append(np_)
        from pptx.text.text import _Paragraph
        para = _Paragraph(np_, tf)
        for r in para.runs[1:]:
            r._r.getparent().remove(r._r)
        para.runs[0].text = t
    if size:
        for p in tf.paragraphs:
            for r in p.runs: r.font.size = Pt(size)

def add_para(sh, text, size=None, bold_prefix=None):
    tf = sh.text_frame
    src = tf.paragraphs[-1]._p
    np_ = copy.deepcopy(src)
    src.addnext(np_)
    from pptx.text.text import _Paragraph
    para = _Paragraph(np_, tf)
    for r in para.runs[1:]:
        r._r.getparent().remove(r._r)
    run = para.runs[0]
    if bold_prefix:
        run.text = bold_prefix
        run.font.bold = True
        r2 = copy.deepcopy(run._r); run._r.addnext(r2)
        para.runs[1].text = text; para.runs[1].font.bold = False
    else:
        run.text = text
    if size:
        for r in para.runs: r.font.size = Pt(size)

def drop_notes_line(prs):
    last = prs.slides[len(prs.slides) - 1]
    for sh in list(last.shapes):
        if sh.has_text_frame and 'bildanteckningarna' in sh.text_frame.text:
            sh._element.getparent().remove(sh._element)

src, dst = sys.argv[1], sys.argv[2]

p = Presentation(f"{src}/v40_01_Sinusformad_vaxelspanning_elev.pptx")
add_para(shape(p.slides[7], 5), "Ombord: 440 V ger û ≈ 622 V, 690 V ger û ≈ 976 V.")
add_para(shape(p.slides[18], 4), " Mät aldrig nätspänning med oscilloskopets jordade probe. Ombord är nätet ofta ett IT-system: en jordad mätklämma på en fas ger jordfel.", size=16, bold_prefix="Säkerhet:")
drop_notes_line(p)
p.save(f"{dst}/v40_01_Sinusformad_vaxelspanning_elev.pptx")

p = Presentation(f"{src}/v40_02_Reaktans_och_impedans_elev.pptx")
s8 = p.slides[7]
set_paras(shape(s8, 13), ["R = 15 Ω"]); set_paras(shape(s8, 14), ["Xᴸ = 20 Ω"]); set_paras(shape(s8, 15), ["|Z| = 25 Ω"])
set_paras(shape(s8, 20), ["R = 15 Ω och X = 20 Ω ger |Z| = 25 Ω.", "Vid U = 100 V RMS blir I = 4,0 A."])
set_paras(shape(p.slides[20], 5), ["R = 12 Ω och X = +5 Ω ger φ = +22,6°.", "X = −5 Ω ger φ = −22,6°."])
drop_notes_line(p)
p.save(f"{dst}/v40_02_Reaktans_och_impedans_elev.pptx")

p = Presentation(f"{src}/v40_03_Effekt_i_vaxelstromskretsar_elev.pptx")
set_paras(shape(p.slides[5], 5), ["Enfas, sinus: 230 V, 6,0 A, cos φ = 0,85", "ger S = 1,38 kVA och P = 1,17 kW."])
set_paras(shape(p.slides[6], 5), ["Vid cos φ = 0,85 är sin φ ≈ 0,53.", "S = 1,38 kVA ger Q ≈ 0,73 kvar."])
set_paras(shape(p.slides[18], 3), ["Energi vid landanslutning"])
set_paras(shape(p.slides[18], 4), ["Vid landanslutning tar en last P = 1,60 kW och S = 2,0 kVA. Den går i 3 h. Beräkna aktiv energi från land."])
set_paras(shape(p.slides[22], 5), ["I modellen P = 2,4 kW, U = 230 V:", "PF 0,60 ger 17,4 A. PF 1,00 ger 10,4 A."])
add_para(shape(p.slides[22], 5), "Ombord används kondensatorbatterier sällan: generatorn levererar Q själv. En kondensator kan hålla farlig spänning efter frånkoppling.", size=18)
drop_notes_line(p)
p.save(f"{dst}/v40_03_Effekt_i_vaxelstromskretsar_elev.pptx")
print("ok")
