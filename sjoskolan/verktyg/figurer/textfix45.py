"""Innehållsändringar i vecka 45: teoriexempel som inte ger samma svar som övningarna,
stödbilder som vägleder i stället för att ge slutsatsen, nedsänkta index och
ingen hänvisning till saknade bildanteckningar.
Användning: textfix45.py <in-mapp> <ut-mapp>"""
import sys
from pptx import Presentation
from pptxtext import shape, set_paras, drop_notes_line, subscript_tokens, replace_words
SUBS = {'Uresistor': ('U', 'resistor'), 'Ukontakt': ('U', 'kontakt'), 'Ukälla': ('U', 'källa'), 'Rtotal': ('R', 'total'), 'Rextra': ('R', 'extra'),
        'Ulast': ('U', 'last'), 'Unom': ('U', 'nom'), 'Um': ('U', 'm'), 'ΔUtotal': ('ΔU', 'total'), 'ΔUfram': ('ΔU', 'fram'), 'ΔUretur': ('ΔU', 'retur'),
        'ΔUkabel': ('ΔU', 'kabel'), 'Rfram': ('R', 'fram'), 'Rretur': ('R', 'retur'), 'Rföre': ('R', 'före'), 'Refter': ('R', 'efter'),
        'ΔUföre': ('ΔU', 'före'), 'ΔUefter': ('ΔU', 'efter'), 'Rslinga': ('R', 'slinga'), 'Urms': ('U', 'rms'), 'Uut': ('U', 'ut'), 'Uin': ('U', 'in'),
        'Rlast': ('R', 'last'), 'Rp': ('R', 'p'), 'Itotal': ('I', 'total')}
src, dst = sys.argv[1], sys.argv[2]

f = "v45_01_Systematisk_felsokning_elev.pptx"; p = Presentation(f"{src}/{f}")
set_paras(shape(p.slides[21], 5), ["0,15 Ω vid 8 A ger 1,2 V fall och 9,6 W värme.", "Kontrollerade felmoduler används i undervisning, aldrig avsiktligt glapp."])
set_paras(shape(p.slides[22], 5), ["Välj en mätpunkt och ett driftläge där de två hypoteserna", "förutsäger olika värden. Skriv förväntat värde före mätningen."])
set_paras(shape(p.slides[13], 12), ["Förutsättningar: Källan har 24 V. B har 0 V mot hel retur. En enkel kedja och samma referens antas."])
set_paras(shape(p.slides[14], 4), ["Källan visar 24 V men punkt B i mitten har 0 V mot hel retur. Vilken del prioriteras?"])
replace_words(p, [('I²Rextra', 'I² · Rextra')])
set_paras(shape(p.slides[28], 11), ["Vad förutsäger varje hypotes för källspänningen och spolspänningen, under START och efter släpp?"])
subscript_tokens(p, SUBS); drop_notes_line(p); p.save(f"{dst}/{f}")

f = "v45_02_Analys_av_matresultat_elev.pptx"; p = Presentation(f"{src}/{f}")
set_paras(shape(p.slides[6], 5), ["Krav 23,0–25,0 V, mätning 24,90 ±0,15 V: intervallet 24,75–25,05 V", "går över kravets övre gräns, så överensstämmelse kan inte visas."])
set_paras(shape(p.slides[8], 10), ["Krav 20,0 V ±2 %. Mätning 19,80 V ±0,15 V. Hela mätintervallet måste ligga inom kravet."])
set_paras(shape(p.slides[8], 11), ["Krav = 20,0 · (1 ± 0,02)", "Mätintervall = 19,80 ±0,15"])
set_paras(shape(p.slides[8], 12), ["1. Krav: 19,60–20,40 V.", "2. Mätintervall: 19,65–19,95 V.", "3. Hela mätintervallet ligger inom kravet: överensstämmelse visad."])
set_paras(shape(p.slides[8], 13), ["Med större felgräns eller en annan beslutsregel kan slutsatsen bli en annan."])
set_paras(shape(p.slides[23], 10), ["Källa 12 V, last 9 V och ström 1,5 A. Returens fall är 0,9 V. Samtidiga data."])
set_paras(shape(p.slides[23], 12), ["1. Totalt fall = 12 − 9 = 3 V.", "2. Framfall = 3 − 0,9 = 2,1 V. Rfram = 2,1/1,5 = 1,4 Ω.", "3. Rretur = 0,9/1,5 = 0,6 Ω."])
set_paras(shape(p.slides[23], 13), ["Kontroll: (1,4 + 0,6) Ω · 1,5 A = 3 V. Beräkningen lokaliserar fall till två delar."])
set_paras(shape(p.slides[22], 5), ["Ombord: avsluta arbetstillstånd och LOTO, kontrollera isolationsvaktens värde,", "informera vakthavande och bryggan, för in åtgärden i maskindagbok och underhållssystem."])
set_paras(shape(p.slides[7], 5), ["Pröva varje hypotes mot alla mätvärden,", "inte bara mot det första symtomet."])
subscript_tokens(p, SUBS); drop_notes_line(p); p.save(f"{dst}/{f}")

f = "v45_03_Repetition_infor_tentamen_elev.pptx"; p = Presentation(f"{src}/{f}")
set_paras(shape(p.slides[23], 10), ["Källan håller 12 V. Vid 3 A får lasten 10,5 V före och 11,85 V efter åtgärd (strömmen antas oförändrad)."])
set_paras(shape(p.slides[8], 10), ["18 V matar 2 kΩ och 2 kΩ i serie. En last på 2 kΩ ansluts över den nedre resistorn."])
set_paras(shape(p.slides[8], 11), ["Rp = R₂ · Rlast/(R₂ + Rlast)", "Uut = Uin · Rp/(R₁ + Rp)"])
set_paras(shape(p.slides[8], 12), ["1. Rp = 2 · 2/(2 + 2) = 1 kΩ.", "2. Uut = 18 · 1/(2 + 1) = 6 V.", "3. Itotal = 18/(2 + 1) = 6 mA. Varje nedre gren får 3 mA."])
set_paras(shape(p.slides[8], 13), ["Kontroll: 3 + 3 = 6 mA och fallet över övre resistorn är 12 V. 12 + 6 = 18 V."])
replace_words(p, [('R₂Rlast', 'R₂ · Rlast'), ('UinRp', 'Uin · Rp')])
set_paras(shape(p.slides[23], 12), ["1. Före: R = (12 − 10,5)/3 = 0,50 Ω.", "2. Efter: R = (12 − 11,85)/3 = 0,05 Ω.", "3. Samma ström och mätpunkter visar tydligt minskad resistans."])
set_paras(shape(p.slides[30], 10), ["Vilken väg ska hålla K1 efter START-släpp?"])
subscript_tokens(p, SUBS); drop_notes_line(p); p.save(f"{dst}/{f}")
