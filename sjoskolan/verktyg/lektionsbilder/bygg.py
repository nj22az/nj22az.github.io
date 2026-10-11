#!/usr/bin/env python3
"""Lektionsbilder: en kort presentation per lektion, alltid med samma upplägg (från vecka 42).

Mallen, i samma ordning varje gång, så att eleverna känner igen lektionen:
  titel · kursplanen i dag · startfråga · så går lektionen till · teori (en bild per avsnitt, i delsidans ordning och
  med delsidans rubriker, märkt ”Del N · avsnitt M”) · exempel ur delsidan · övningar tillsammans · paus ·
  avslutsfråga · sammanfattning · det här gör du nu · källor.

Bilderna visar inga övningstexter: övningarna görs på delsidan, där ledtrådar och facit finns. Start- och
avslutsfrågan och övningarnas titlar hämtas ur posterna. Placeringarna (innehall/placeringar/presentation-fragor.json)
skrivs om med de nya bild- och formnumren, så att `innehall.py bygg presentationer` fortsätter att hålla texten i takt
med databasen. Presentationens gamla övningsbilder tas bort ur placeringar/presentation.json.

Innehållet står i lektioner.py. Mallen (master och tema från Sjöskolans presentationer, sidfot och omslagsbilder)
ligger i mall/. Figurerna är veckans befintliga figurer (vecka-NN/aktuell/figurer/).

    python3 sjoskolan/verktyg/lektionsbilder/bygg.py v42_01          # pptx, pdf och webbildspel
    python3 sjoskolan/verktyg/lektionsbilder/bygg.py --alla
    python3 sjoskolan/verktyg/lektionsbilder/bygg.py v42_01 --utan-pdf
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

HAR = Path(__file__).resolve().parent
SJO = HAR.parents[1]
MALL = HAR / 'mall'
INNEHALL = SJO / 'innehall'
sys.path.insert(0, str(HAR))
sys.path.insert(0, str(SJO / 'verktyg' / 'veckosidor'))
import kursplan as K  # noqa: E402
import lektioner as L  # noqa: E402

INK = RGBColor(0x11, 0x11, 0x11)
DAMPAD = RGBColor(0x4A, 0x4A, 0x4A)
BLA = RGBColor(0x06, 0x4F, 0x91)
VIT = RGBColor(0xFF, 0xFF, 0xFF)
LINJE = RGBColor(0xD5, 0xDB, 0xE3)
YTA = RGBColor(0xF2, 0xF5, 0xF9)

# Innehållsytan: mellan rubriken och sidfoten (sidfoten börjar 6,67 tum).
X0, BREDD = 0.6, 8.8
Y0, Y1 = 1.5, 6.45


def farg(hexa):
    return RGBColor.from_string(hexa.lstrip('#').upper())


def omrade(lm):
    for _nyckel, namn, hexa, mal in K.OMRADEN:
        if lm in mal:
            return namn, hexa
    raise KeyError(lm)


# --- text -----------------------------------------------------------------------------------------------------
INDEX = re.compile(r'(_\{[^}]*\}|\^\{[^}]*\})')


def skriv(p, text, size, bold=False, color=INK, italic=False):
    """Text med index och fetstil: X_{L} blir nedsänkt, x^{2} upphöjd (samma markering som i innehållsdatabasen),
    **ord** blir fetstil."""
    fet = bold
    for del_ in re.split(r'(\*\*)', text):
        if del_ == '**':
            fet = not fet
            continue
        for bit in INDEX.split(del_):
            if not bit:
                continue
            r = p.add_run()
            baslinje = None
            if bit.startswith('_{'):
                bit, baslinje = bit[2:-1], '-25000'
            elif bit.startswith('^{'):
                bit, baslinje = bit[2:-1], '30000'
            r.text = bit
            f = r.font
            f.size, f.bold, f.italic, f.name = Pt(size), fet, italic, 'Calibri'
            f.color.rgb = color
            if baslinje:
                r._r.get_or_add_rPr().set('baseline', baslinje)


def ruta(s, x, y, w, h, stycken, size=20, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         efter=8, rad=1.08, namn=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if namn:
        tb.name = namn
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for st_i, st in enumerate(stycken):
        if isinstance(st, str):
            st = {'text': st}
        p = tf.paragraphs[0] if st_i == 0 else tf.add_paragraph()
        p.alignment = st.get('align', align)
        p.line_spacing = st.get('rad', rad)
        p.space_after = Pt(st.get('efter', efter))
        skriv(p, st['text'], st.get('size', size), st.get('bold', bold), st.get('color', color), st.get('italic', False))
    return tb


def platta(s, x, y, w, h, fyll, rundad=False, linje=None):
    form = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rundad else MSO_SHAPE.RECTANGLE,
                              Inches(x), Inches(y), Inches(w), Inches(h))
    form.fill.solid()
    form.fill.fore_color.rgb = fyll
    if linje:
        form.line.color.rgb = linje
        form.line.width = Pt(1)
    else:
        form.line.fill.background()
    form.shadow.inherit = False
    if rundad:
        form.adjustments[0] = 0.18
    form.text_frame.text = ''
    return form


def figur(s, fil, x, y, wmax, hmax, alt):
    w, h = Image.open(fil).size
    skala = min(wmax / w, hmax / h)
    wf, hf = w * skala, h * skala
    pic = s.shapes.add_picture(str(fil), Inches(x + (wmax - wf) / 2), Inches(y + (hmax - hf) / 2), Inches(wf), Inches(hf))
    pic.name = ('Figur: ' + alt)[:250]
    pic._element.nvPicPr.cNvPr.set('descr', alt)
    return pic


# --- gemensamma delar -------------------------------------------------------------------------------------------
def ny_bild(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # ”Tom”: inga platshållare


def sidfot(s, nr, etikett, avsnitt=None):
    """Sidfoten: Sjöskolans list, lektionen (och avsnittet på delsidan) och bildnumret. Delsidans generator
    (innehall/export/delsidor_veckor.py) hoppar över text under 6,9 tum, så sidfoten blir aldrig teori."""
    s.shapes.add_picture(str(MALL / 'sidfot.jpg'), 0, Inches(6.67), Inches(10), Inches(0.83))
    text = f'{etikett}  ·  {avsnitt}' if avsnitt else etikett
    ruta(s, 1.6, 7.0, 6.8, 0.34, [text], size=12, color=VIT, align=PP_ALIGN.CENTER, efter=0)
    ruta(s, 8.76, 7.0, 0.77, 0.34, [str(nr)], size=15, color=VIT, align=PP_ALIGN.RIGHT, efter=0)


def rubrik(s, text, accent):
    """Rubriken är bildens första text (delsidan läser den som avsnittets rubrik)."""
    ruta(s, X0, 0.36, BREDD, 0.8, [text], size=30, bold=True, anchor=MSO_ANCHOR.BOTTOM, efter=0)
    platta(s, X0 + 0.03, 1.2, 0.9, 0.07, accent)


def gradient(s, x, y, w, h):
    """Mörk ton bakom omslagets text (samma toning som Sjöskolans tidigare omslag)."""
    form = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    form.line.fill.background()
    form.shadow.inherit = False
    spPr = form._element.spPr
    for barn in list(spPr):
        if barn.tag in (qn('a:solidFill'), qn('a:noFill'), qn('a:gradFill')):
            spPr.remove(barn)
    A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    g = etree.SubElement(spPr, f'{{{A}}}gradFill', rotWithShape='1')
    lst = etree.SubElement(g, f'{{{A}}}gsLst')
    for pos, alfa in ((0, 70000), (45000, 55000), (75000, 0), (100000, 0)):
        gs = etree.SubElement(lst, f'{{{A}}}gs', pos=str(pos))
        c = etree.SubElement(gs, f'{{{A}}}srgbClr', val='001427')
        etree.SubElement(c, f'{{{A}}}alpha', val=str(alfa))
    etree.SubElement(g, f'{{{A}}}lin', ang='0', scaled='0')
    # gradFill ska stå före a:ln i spPr
    ln = spPr.find(qn('a:ln'))
    if ln is not None:
        spPr.remove(g)
        ln.addprevious(g)
    return form


# --- poster ur innehållsdatabasen -------------------------------------------------------------------------------
def post(id_):
    return json.loads((INNEHALL / 'ovningar' / f'{id_}.json').read_text(encoding='utf-8'))


def placeringar(namn):
    return json.loads((INNEHALL / 'placeringar' / f'{namn}.json').read_text(encoding='utf-8'))


def fraga_for(lek, roll):
    for pl in placeringar('presentation-fragor')['placeringar']:
        if pl['plats'] == lek['fil'] and pl.get('roll') == roll:
            return pl, post(pl['ovning'])
    raise KeyError(f'{lek["fil"]}: ingen {roll}fråga i placeringar/presentation-fragor.json')


def ovningar_for(lek, nummer):
    """Övningarnas nummer och titel ur delsidans placeringar (kurs-formelstod) och posterna."""
    ut = []
    for pl in placeringar('kurs-formelstod')['placeringar']:
        if pl.get('del') == lek['id'] and pl.get('nummer') in [f'Övning {n}' for n in nummer]:
            ut.append((pl['nummer'], post(pl['ovning'])['titel']))
    ordning = {f'Övning {n}': i for i, n in enumerate(nummer)}
    ut.sort(key=lambda x: ordning[x[0]])
    if len(ut) != len(nummer):
        raise KeyError(f'{lek["id"]}: hittar inte alla övningar {nummer}')
    return ut


# --- bildtyper --------------------------------------------------------------------------------------------------
def b_titel(prs, lek, nr):
    s = ny_bild(prs)
    s.shapes.add_picture(str(MALL / lek['omslag']), 0, 0, Inches(10), Inches(6.67))
    gradient(s, 0, 0, 10, 6.67)
    ruta(s, 0.6, 1.1, 6.6, 1.0, [lek['titel']], size=44, bold=True, color=VIT, efter=0)
    ruta(s, 0.6, 2.15, 6.2, 1.0, [lek['underrubrik']], size=22, color=VIT, efter=0)
    ruta(s, 0.6, 3.25, 6.4, 0.45, [lek['kicker']], size=16, color=VIT, efter=0)
    ruta(s, 0.6, 5.7, 6.4, 0.4, [f'Del {lek["del"]} av {lek["delar"]} · Elteknik och ellära'], size=14, color=VIT, efter=0)
    sidfot(s, nr, lek['etikett'])


def b_kursplan(prs, lek, nr, accent):
    s = ny_bild(prs)
    rubrik(s, 'Kursplanen i dag', accent)
    y = Y0 + 0.05
    for lm in lek['mal']:
        namn, hexa = omrade(lm)
        chip = platta(s, X0, y + 0.04, 1.05, 0.46, farg(hexa), rundad=True)
        tf = chip.text_frame
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        skriv(p, lm, 16, bold=True, color=INK)
        ruta(s, X0 + 1.3, y, BREDD - 1.3, 0.95, [K.ENKEL[lm], {'text': namn, 'size': 14, 'color': DAMPAD}], size=22, efter=2)
        y += 1.05
    y += 0.15
    ruta(s, X0, y, BREDD, 0.45, ['Efter lektionen kan du'], size=22, bold=True, color=BLA, efter=0)
    y += 0.5
    ruta(s, X0, y, BREDD, Y1 - y, [f'{i}. {t}' for i, t in enumerate(lek['efter'], 1)], size=21, efter=12)
    sidfot(s, nr, lek['etikett'])


def b_fraga(prs, lek, nr, accent, roll, instruktion=None):
    pl, p = fraga_for(lek, roll)
    s = ny_bild(prs)
    rubrik(s, 'Startfråga' if roll == 'start' else 'Avslutsfråga', accent)
    fraga = ruta(s, X0, 2.0, BREDD, 2.2, [p['uppgift']['fraga']], size=32, bold=True, anchor=MSO_ANCHOR.MIDDLE, rad=1.05)
    former = {'fraga': str(fraga.shape_id)}
    if roll == 'avslut' and p['uppgift'].get('instruktion'):
        ins = ruta(s, X0, 4.45, BREDD, 1.6, [p['uppgift']['instruktion']], size=20, color=DAMPAD)
        former['instruktion'] = str(ins.shape_id)
    if instruktion:
        ruta(s, X0, 5.55, BREDD, 0.6, [instruktion], size=18, color=BLA)
    sidfot(s, nr, lek['etikett'])
    return pl, former


def b_upplagg(prs, lek, nr, accent, rader):
    s = ny_bild(prs)
    rubrik(s, 'Så går lektionen till', accent)
    radh = min(0.62, (Y1 - Y0 - 0.1) / len(rader))
    for i, (tid, text) in enumerate(rader):
        y = Y0 + 0.05 + i * radh
        if text.startswith('Paus'):
            platta(s, X0, y, BREDD, radh - 0.06, YTA)
        ruta(s, X0 + 0.1, y, 1.1, radh - 0.06, [tid], size=18, bold=True, color=BLA, anchor=MSO_ANCHOR.MIDDLE, efter=0)
        ruta(s, X0 + 1.35, y, BREDD - 1.45, radh - 0.06, [text], size=18, anchor=MSO_ANCHOR.MIDDLE, efter=0)
    sidfot(s, nr, lek['etikett'])


def _figurfil(lek, namn):
    fil = SJO / lek['figurmapp'] / namn
    if not fil.exists():
        raise FileNotFoundError(fil)
    return fil


def b_teori(prs, lek, nr, accent, b):
    """Teoribild: nyckelmening (blå) och korta rader, figur till höger eller (figur_under) under texten i full bredd."""
    s = ny_bild(prs)
    etikett = b.get('etikett') or (f'Del {lek["del"]} · avsnitt {b["avsnitt"]}' if b.get('avsnitt') else None)
    rubrik(s, b['rubrik'], accent)
    under = b.get('figur_under')
    bredd = 4.45 if b.get('figur') and not under else BREDD
    y0 = Y0 + 0.05
    stycken = []
    if b.get('scenario'):
        if b.get('figur') and not under:
            # Exempel med figur: fallet står i full bredd under rubriken, kolumnerna börjar under det.
            ruta(s, X0, y0, BREDD, 0.8, [{'text': b['scenario'], 'size': 18, 'italic': True, 'color': DAMPAD, 'efter': 0}])
            y0 += 0.85
        else:
            stycken.append({'text': b['scenario'], 'size': 19, 'italic': True, 'color': DAMPAD, 'efter': 12})
    mellan = b.get('mellan', 10)
    if b.get('nyckel') and not b.get('nyckel_sist'):
        stycken.append({'text': b['nyckel'], 'size': b.get('nyckelstorlek', 25), 'bold': True, 'color': BLA, 'efter': 12, 'rad': 1.0})
    storlek = b.get('size', 20)
    for t in b.get('punkter', []):
        stycken.append({'text': t, 'size': storlek, 'efter': mellan})
    if b.get('nyckel') and b.get('nyckel_sist'):
        stycken.append({'text': b['nyckel'], 'size': 22, 'bold': True, 'color': BLA, 'efter': 0, 'rad': 1.0})
    if b.get('slutsats'):  # fet men inte blå: delsidan läser den som vanlig text, inte som avsnittets formel
        stycken.append({'text': b['slutsats'], 'size': storlek, 'bold': True, 'efter': 0})
    texthojd = b.get('texthojd', 1.9) if under else Y1 - y0
    ruta(s, X0, y0, bredd, texthojd, stycken)
    if b.get('figur') and under:
        y = y0 + texthojd + 0.1
        figur(s, _figurfil(lek, b['figur']), X0, y, BREDD, Y1 - y, b['alt'])
    elif b.get('figur'):
        figur(s, _figurfil(lek, b['figur']), 5.3, y0, 4.15, Y1 - y0 - 0.1, b['alt'])
    sidfot(s, nr, lek['etikett'], etikett)


def b_ovningar(prs, lek, nr, accent, b):
    s = ny_bild(prs)
    rubrik(s, 'Övningar tillsammans', accent)
    ruta(s, X0, Y0 + 0.05, BREDD, 0.55, [f'Öppna del {lek["del"]} och gå till 4 Övningar.'], size=22, bold=True, color=BLA, efter=0)
    y = Y0 + 0.75
    for nummer, titel in ovningar_for(lek, b['nummer']):
        platta(s, X0, y, BREDD, 0.62, YTA)
        ruta(s, X0 + 0.15, y, 1.9, 0.62, [nummer], size=20, bold=True, anchor=MSO_ANCHOR.MIDDLE, efter=0)
        ruta(s, X0 + 2.1, y, BREDD - 2.2, 0.62, [titel], size=20, anchor=MSO_ANCHOR.MIDDLE, efter=0)
        y += 0.74
    ruta(s, X0, y + 0.15, BREDD, Y1 - y - 0.15, [b['instruktion'], {'text': lek['delsida_url'], 'size': 14, 'color': DAMPAD}], size=19, efter=8)
    sidfot(s, nr, lek['etikett'], f'Del {lek["del"]} · 4 Övningar')


def b_paus(prs, lek, nr, accent, b):
    s = ny_bild(prs)
    platta(s, 0, 0, 10, 6.67, YTA)
    platta(s, 4.55, 2.2, 0.9, 0.08, accent)
    ruta(s, 1.0, 2.45, 8.0, 1.0, ['Paus'], size=48, bold=True, align=PP_ALIGN.CENTER, efter=0)
    ruta(s, 1.0, 3.5, 8.0, 0.6, [b['text']], size=24, color=DAMPAD, align=PP_ALIGN.CENTER, efter=0)
    sidfot(s, nr, lek['etikett'])


def b_lista(prs, lek, nr, accent, b):
    s = ny_bild(prs)
    rubrik(s, b['rubrik'], accent)
    stycken = [{'text': f'{i}. {t}', 'efter': 14} for i, t in enumerate(b['punkter'], 1)]
    if b.get('not'):
        stycken.append({'text': b['not'], 'size': 17, 'color': DAMPAD, 'efter': 0})
    ruta(s, X0, Y0 + 0.1, BREDD, Y1 - Y0 - 0.1, stycken, size=b.get('size', 21))
    sidfot(s, nr, lek['etikett'])


def b_kallor(prs, lek, nr, accent, b):
    s = ny_bild(prs)
    rubrik(s, 'Källor och fortsatt läsning', accent)
    # En textruta per källa: delsidan slår annars ihop rader utan slutpunkt till ett stycke.
    for i, (t, url) in enumerate(b['kallor']):
        tb = ruta(s, X0, Y0 + 0.1 + i * 0.5, BREDD, 0.45, [t], size=17, efter=0)
        if url:
            for r in tb.text_frame.paragraphs[0].runs:
                r.hyperlink.address = url
                r.font.color.rgb = BLA
    sidfot(s, nr, lek['etikett'])


# --- bygg -------------------------------------------------------------------------------------------------------
def bygg_deck(lek):
    prs = Presentation(str(MALL / 'mall.pptx'))
    accent = farg(dict((k, h) for k, _n, h, _m in K.OMRADEN)[lek['omrade']])
    frageplacering = {}
    nr = 0
    for typ, b in [('titel', None), ('kursplan', None)] + lek['bilder']:
        nr += 1
        if typ == 'titel':
            b_titel(prs, lek, nr)
        elif typ == 'kursplan':
            b_kursplan(prs, lek, nr, accent)
        elif typ in ('start', 'avslut'):
            pl, former = b_fraga(prs, lek, nr, accent, typ, (b or {}).get('instruktion'))
            frageplacering[typ] = (nr, former)
        elif typ == 'upplagg':
            b_upplagg(prs, lek, nr, accent, b)
        elif typ == 'teori':
            b_teori(prs, lek, nr, accent, b)
        elif typ == 'ovningar':
            b_ovningar(prs, lek, nr, accent, b)
        elif typ == 'paus':
            b_paus(prs, lek, nr, accent, b)
        elif typ == 'lista':
            b_lista(prs, lek, nr, accent, b)
        elif typ == 'kallor':
            b_kallor(prs, lek, nr, accent, b)
        else:
            raise ValueError(typ)
    fil = SJO / lek['fil']
    prs.save(str(fil))
    return fil, nr, frageplacering


def skriv_placeringar(lek, frageplacering):
    """Start- och avslutsfrågans bild och former; presentationens gamla övningsbilder tas bort."""
    p = INNEHALL / 'placeringar' / 'presentation-fragor.json'
    d = json.loads(p.read_text(encoding='utf-8'))
    for pl in d['placeringar']:
        if pl['plats'] == lek['fil'] and pl.get('roll') in frageplacering:
            pl['bild'], pl['former'] = frageplacering[pl['roll']]
    p.write_text(json.dumps(d, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    p = INNEHALL / 'placeringar' / 'presentation.json'
    d = json.loads(p.read_text(encoding='utf-8'))
    fore = len(d['placeringar'])
    d['placeringar'] = [pl for pl in d['placeringar'] if pl['plats'] != lek['fil']]
    if len(d['placeringar']) != fore:
        p.write_text(json.dumps(d, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def pdf(fil):
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', str(fil.parent), str(fil)],
                   check=True, capture_output=True)


def main(argv):
    alla = '--alla' in argv
    utan_pdf = '--utan-pdf' in argv
    valda = [a for a in argv if not a.startswith('--')]
    lektioner = [lek for lek in L.LEKTIONER if alla or lek['id'] in valda]
    if not lektioner:
        sys.exit(f'Ange en lektion ({", ".join(x["id"] for x in L.LEKTIONER)}) eller --alla')
    for lek in lektioner:
        fil, n, fp = bygg_deck(lek)
        skriv_placeringar(lek, fp)
        print(f'{lek["id"]}: {n} bilder → {fil.relative_to(SJO)}')
        if not utan_pdf:
            pdf(fil)
    if not utan_pdf:
        subprocess.run([sys.executable, str(SJO / 'verktyg' / 'bildspel' / 'bygg.py')], check=True)


if __name__ == '__main__':
    main(sys.argv[1:])
