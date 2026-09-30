#!/usr/bin/env python3
"""Lärarkopior med manus för vecka 41 och framåt: elevens presentation + talaranteckningar, krypterad.

Anteckningarna byggs ur tre källor:
- bildens egen text (vad du säger),
- utbildningsguiden (larare/vecka-NN.html, krypterad): upplägget, avsnitten under Så förklarar du (kärnan, ordningen,
  tavlan, bilden att använda, kontrollfrågor) och elevernas vanliga frågor,
- innehållsdatabasen: övningens facit, svar och första ledtråd på övningsbilderna.

    LARARLOSEN=… python3 sjoskolan/verktyg/larare/lararmanus.py 41 42

Skriver larare/filer/<presentation>_larare.pptx.enc. Klartexten skrivs aldrig till disk i repot. Vecka 40 har eget
manus (lararnoter.py) och är låst.
"""
import html
import io
import json
import os
import re
import sys
from pathlib import Path

from pptx import Presentation

SJO = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(SJO / 'innehall' / 'lib'), str(SJO / 'innehall'), str(SJO / 'verktyg' / 'larare')]
import krypto  # noqa: E402
from atkomst import Atkomst  # noqa: E402
from lararnoter import stycke  # noqa: E402

WEBB = 'https://nj22az.github.io/sjoskolan'
# Guidens avsnitt → (presentation, bild). Avsnittet läggs i anteckningarna till den bild där det hör hemma.
AVSNITT = {
    '41': {'Fas och huvudspänning': ('v41_01', 7), 'Neutralströmmen': ('v41_01', 21), 'Y och Δ': ('v41_02', 6),
           'Bruten neutral': ('v41_01', 22)},
    '42': {'Hur el skadar': ('v42_01', 6), 'Arbetsmetoder och ansvar': ('v42_02', 21), 'Riskbedömning och CAT': ('v42_03', 7)},
    '43': {'Skyddsdata': ('v43_01', 7), 'Motor och transformator': ('v43_02', 6), 'Hållkretsen': ('v43_03', 8)},
    '44': {'IT-nätet ombord': ('v44_01', 21), 'Isolationsmätning': ('v44_03', 7)},
    '45': {'Felsökning': ('v45_01', 6), 'Mätresultat mot krav': ('v45_02', 7)},
}
NOTATION = [('Uᴸ', 'U_{L}'), ('Uꜰ', 'U_{F}'), ('Ugren', 'U_{gren}'), ('Igren', 'I_{gren}'), ('Iᴸ', 'I_{L}'), ('IN ', 'I_{N} ')]


def ren(s):
    s = html.unescape(re.sub(r'<[^>]+>', '', s)).strip()
    for a, b in NOTATION:
        s = s.replace(a, b)
    return re.sub(r'\s+', ' ', s)


def guide(vecka, losen):
    """Utbildningsguiden: upplägg (rader), avsnitt {rubrik: [rader]} och vanliga frågor [rader]."""
    import subprocess
    import tempfile
    with tempfile.TemporaryDirectory() as d:  # klartexten ligger aldrig i repot och tas bort direkt
        tmp = Path(d) / 'guide.html'
        subprocess.run(['node', str(SJO / 'verktyg' / 'larare' / 'las.mjs'), 'unlock', str(SJO / 'larare' / f'vecka-{vecka}.html'), str(tmp)],
                       check=True, capture_output=True, env={**os.environ, 'LARARLOSEN': losen})
        t = tmp.read_text(encoding='utf-8')
    upplagg = [f'{ren(a)}: {ren(b)}' for a, b in re.findall(r'<tr><td>(.*?)</td><td>(.*?)</td></tr>', t.split('Förslag till upplägg')[1].split('</table>')[0], re.S)]
    avsnitt = {}
    del_ = t.split('Så förklarar du')[1].split('Vanliga frågor')[0]
    for rubrik, kropp in re.findall(r'<h3>(.*?)</h3>(.*?)(?=<h3>|$)', del_, re.S):
        rader = []
        for m in re.finditer(r'<p><strong>(.*?):</strong>(.*?)</p>|<ol[^>]*>(.*?)</ol>|<dl class="qa">(.*?)</dl>', kropp, re.S):
            if m.group(1) and ren(m.group(2)):
                rader.append(f'{ren(m.group(1))}: {ren(m.group(2))}')
            elif m.group(3):
                rader += [f'{i}. {ren(x)}' for i, x in enumerate(re.findall(r'<li>(.*?)</li>', m.group(3), re.S), 1)]
            elif m.group(4):
                rader += [f'Fråga: {ren(q)} Svar: {ren(s)}' for q, s in re.findall(r'<dt>(.*?)</dt><dd>(.*?)</dd>', m.group(4), re.S)]
        avsnitt[ren(rubrik)] = rader
    fragor = [f'{ren(q)} {ren(s)}' for q, s in re.findall(r'<dt>(.*?)</dt><dd>(.*?)</dd>', t.split('Vanliga frågor')[1].split('Labbar')[0], re.S)]
    return upplagg, avsnitt, fragor


def texter(bild):
    ut = []
    for sh in bild.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t and not re.fullmatch(r'\d+', t):
                ut.append(re.sub(r'\s*\n\s*', ' ', t))
    return ut


def svar_text(p):
    l = p.get('losning') or {}
    rader = []
    if l.get('text'):
        rader.append(f'Facit: {l["text"]}')
    for s in l.get('svar') or []:
        v = s['varde']
        v = f'{v:g}'.replace('.', ',') if isinstance(v, (int, float)) else v
        rader.append(f'Svar: {s["storhet"]} = {v} {s.get("enhet", "")}'.strip())
    led = [x['text'] for x in p.get('ledtradar') or []]
    if led:
        rader.append(f'Om de fastnar (ledtråd 1): {led[0]}')
    return rader or ['Facit finns i lärarguiden (facit och bedömning).']


def manus(vecka, deck, n, bild, ovn, stod, avs, upplagg, fragor, antal):
    t = texter(bild)
    titel = t[0] if t else ''
    kropp = t[1:]
    rader = []
    if n == 1:
        rader = [f'Titelbild. Säg: ”Idag: {titel.lower()}.”', '', 'Veckans upplägg enligt guiden:'] + upplagg
    elif n in ovn:
        pl, p = ovn[n]
        rader = [f'{pl["nummer"]}: {p["titel"]}. Läs uppgiften högt. Eleverna räknar själva eller i par, 3–5 minuter.',
                 'Rutin: alla skriver på miniwhiteboard, säg ”Visa!”, se hur många som har rätt, gå sedan igenom.', ''] + svar_text(p)
    elif n in stod:
        pl, p = stod[n]
        rader = [f'Stödbild till {pl["nummer"].lower()}. Visa den först när eleverna har försökt en stund.',
                 'Den ger samband, förutsättningar och arbetsgång, inte svaret. Säg: ”Vilket samband behöver ni?”']
    elif titel.startswith('Exempel'):
        rader = ['Genomräknat exempel. Räkna stegen på tavlan tillsammans och låt klassen säga nästa steg innan du visar det.',
                 '', 'På bilden:'] + kropp
    elif titel == 'Pass 2':
        rader = ['Pass 2. Säg vad som kommer: ' + ', '.join(kropp[:1]) + '.', 'Kort paus före om passet börjar efter rast.']
    elif titel.startswith('Källor'):
        rader = ['Källor. Visa kort. Hänvisningen ska gå att följa.']
    elif titel.startswith('Avslut'):
        rader = ['Avslut. Säg vad som ska vara gjort före nästa lektion och påminn om inlämningen.', '', 'På bilden:'] + kropp
        if fragor:
            rader += ['', 'Vanliga frågor från elever:'] + fragor
    else:
        rader = ['Säg:'] + kropp
    for rubrik, (d, b) in AVSNITT.get(vecka, {}).items():
        if deck.startswith(d) and b == n and rubrik in avs:
            rader += ['', f'Ur guiden, {rubrik}:'] + avs[rubrik]
    rader += ['', f'Elevernas sida: {WEBB}/vecka-{vecka}/aktuell/ · Bild {n} av {antal}']
    return rader


def main():
    losen = os.environ.get('LARARLOSEN')
    if not losen:
        raise SystemExit('Ange lärarlösenordet i LARARLOSEN.')
    import innehall as I
    a = Atkomst(I.bygg_db(I.katalog_eller_avbryt()), 'larare')
    placeringar = json.loads((SJO / 'innehall' / 'placeringar' / 'presentation.json').read_text())['placeringar']
    for vecka in sys.argv[1:]:
        if int(vecka) <= 40:
            raise SystemExit('vecka 40 och tidigare ändras inte här')
        upplagg, avs, fragor = guide(vecka, losen)
        for fil in sorted((SJO / f'vecka-{vecka}' / 'aktuell').glob('v*_elev.pptx')):
            prs = Presentation(fil)
            bilder = list(prs.slides)
            ovn, stod = {}, {}
            for pl in placeringar:
                if pl['plats'].endswith(fil.name) and not pl.get('kontroll'):
                    p = a.ovning(pl['ovning'])
                    ovn[pl['bild']] = (pl, p)
                    if pl.get('stodbild'):
                        stod[pl['stodbild']] = (pl, p)
            for n, bild in enumerate(bilder, 1):
                tf = bild.notes_slide.notes_text_frame
                tf.clear()
                for j, r in enumerate(manus(vecka, fil.name, n, bild, ovn, stod, avs, upplagg, fragor, len(bilder))):
                    stycke(tf, r, j == 0)
            ut = io.BytesIO()
            prs.save(ut)
            mal = SJO / 'larare' / 'filer' / fil.name.replace('_elev.pptx', '_larare.pptx.enc')
            krypto.kryptera(ut.getvalue(), mal, losen)
            print(f'{mal.relative_to(SJO)}: {len(bilder)} bilder med manus, {len(ovn)} övningar med facit')


if __name__ == '__main__':
    main()
