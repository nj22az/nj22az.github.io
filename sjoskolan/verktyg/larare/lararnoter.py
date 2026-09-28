#!/usr/bin/env python3
"""Lärarkopiorna av vecka 40:s presentationer: elevens presentation + lärarmanus i anteckningarna, krypterad.

Manuset kommer ur vecka-40/aktuell/lektioner.mjs (note), kontrollfrågan och svaret ur innehållsdatabasen (check, answer).
Elevens presentationer har inga anteckningar; lärarkopian skrivs till larare/filer/*_larare.pptx.enc (NJENC1).

    LARARLOSEN=… python3 sjoskolan/verktyg/larare/lararnoter.py

Kör efter varje ändring i lektionerna eller elevpresentationerna. Klartexten skrivs aldrig till disk i repot.
"""
import io
import json
import os
import re
import subprocess
import sys
from pathlib import Path

from pptx import Presentation

SJO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(SJO / 'innehall' / 'lib'))
import krypto  # noqa: E402

WEBB = 'https://nj22az.github.io/sjoskolan'


def stycke(tf, text, forsta):
    """Ett stycke i anteckningarna. X_{L} skrivs med nedsänkt index."""
    p = tf.paragraphs[0] if forsta else tf.add_paragraph()
    for i, del_ in enumerate(re.split(r'_\{([^{}]*)\}', text)):
        if not del_:
            continue
        r = p.add_run()
        r.text = del_
        if i % 2:
            r.font._element.set('baseline', '-25000')


def manus(l, s):
    rader = s.get('note', '').split('\n')
    if s.get('check'):
        if s.get('tur'):
            rader += ['', f'Frågan på bilden: {s["check"]}', f'Svar (visas vid klick): {s["answer"]}']
        elif s['check'].startswith('Räkna innan'):
            rader += ['', f'Facit: {s["answer"]}']
        else:
            rader += ['', f'Kontrollfrågan: {s["check"]}', f'Svar: {s["answer"]}']
    if s['id'] == 'mal':
        rader += ['', f'Webbgenomgång: {WEBB}/vecka-40/aktuell/Genomgang.html?del={l["id"]}&avsnitt=mal',
                  f'Guidad labb: {WEBB}/vaxelstromslabbet/?lage=guidad&del={l["id"]}',
                  f'Inlämningar under lektionen: {WEBB}/larare/resultat.html?vy=presentera']
    return rader


def main():
    losen = os.environ.get('LARARLOSEN')
    if not losen:
        raise SystemExit('Ange lärarlösenordet i LARARLOSEN.')
    js = "import('%s').then(m=>console.log(JSON.stringify(m.LESSONS)))" % (SJO / 'vecka-40' / 'aktuell' / 'lektioner.mjs').as_uri()
    lektioner = json.loads(subprocess.run(['node', '--input-type=module', '-e', js], capture_output=True, text=True, check=True).stdout)
    for l in lektioner:
        elev = SJO / 'vecka-40' / 'aktuell' / l['deck']
        prs = Presentation(elev)
        bilder = list(prs.slides)
        if len(bilder) != len(l['slides']) + 1:
            raise SystemExit(f'{elev.name}: {len(bilder)} bilder, lektionen har {len(l["slides"]) + 1}. Lägg in bilderna först.')
        for i, bild in enumerate(bilder):
            rader = [f'Titelbild. Säg: ”Idag: {l["title"].lower()}.” Klicka till bild 2.'] if i == 0 else manus(l, l['slides'][i - 1])
            tf = bild.notes_slide.notes_text_frame
            tf.clear()
            for j, t in enumerate(rader):
                stycke(tf, t, j == 0)
        ut = io.BytesIO()
        prs.save(ut)
        mal = SJO / 'larare' / 'filer' / l['deck'].replace('_elev.pptx', '_larare.pptx.enc')
        krypto.kryptera(ut.getvalue(), mal, losen)
        print(f'{mal.relative_to(SJO)}: {len(bilder)} bilder med manus')


if __name__ == '__main__':
    main()
