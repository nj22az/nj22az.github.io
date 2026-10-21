#!/usr/bin/env python3
"""Check protected Week 44 teaching and unchanged public slides. No private text is logged."""
import hashlib
import io
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation

SJO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(SJO / 'innehall' / 'lib'))
import krypto


def guide():
    with tempfile.TemporaryDirectory() as d:
        target = Path(d) / 'guide.html'
        subprocess.run(['node', str(SJO / 'verktyg/larare/las.mjs'), 'unlock',
                        str(SJO / 'larare/vecka-44.html'), str(target)],
                       check=True, capture_output=True)
        return target.read_text()


def public_snapshot():
    # Entire student-facing course, assessment, database and simulator files.
    return {str(p.relative_to(SJO)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in SJO.rglob('*') if p.is_file()
            and not any(x in p.parts for x in ('larare', 'verktyg', '__pycache__', 'ut', 'build', '.skyddat'))}


def decks():
    students = sorted((SJO / 'vecka-44/aktuell').glob('v44_*_elev.pptx'))
    analysis = SJO / 'vecka-45/aktuell/v45_02_Analys_av_matresultat_elev.pptx'
    if analysis.exists():
        students.append(analysis)
    for student in students:
        teacher = SJO / 'larare/filer' / student.name.replace('_elev.pptx', '_larare.pptx.enc')
        yield student, teacher


def check():
    text = guide()
    for stale in ('Kapacitans 1', 'Friskt nät:', 'Jordfel L1:',
                  'Visar övervakningen den sämsta fasen?'):
        assert stale not in text, 'Old assessment explanation remains in protected guide'
    for required in ('2,5 MΩ', 'TN-system', 'TT-system', 'IT-nätet ombord',
                     'verifiera', 'Bedömningsgräns', 'Bekräftade villkor'):
        assert required in text, 'Required guide section missing'
    assert 250 / 0.00010 == 2500000
    assert 2500000 * 0.00010 == 250
    assert 230 / (40 + 10) == 4.6
    for student, teacher in decks():
        data = krypto.dekryptera(teacher, os.environ['LARARLOSEN'])
        prs = Presentation(io.BytesIO(data))
        source = Presentation(student)
        assert len(prs.slides) == len(source.slides), 'Teacher/student slide counts differ'
        # Notes may change. Every visible slide XML must be identical to the current student deck.
        with zipfile.ZipFile(student) as a, zipfile.ZipFile(io.BytesIO(data)) as b:
            for name in a.namelist():
                if re.fullmatch(r'ppt/slides/slide\d+\.xml', name):
                    assert a.read(name) == b.read(name), 'Visible slide changed'
        notes = '\n'.join(s.notes_slide.notes_text_frame.text for s in prs.slides)
        assert notes.count('Bedömningsgräns:') == len(prs.slides)
        for stale in ('Kapacitans 1', 'Friskt nät:', 'Jordfel L1:',
                      'Visar övervakningen den sämsta fasen?'):
            assert stale not in notes, 'Old assessment explanation remains in slide notes'
        if student.name.startswith('v44_01'):
            assert 'Ur guiden, IT-nätet ombord:' in prs.slides[8].notes_slide.notes_text_frame.text
            assert 'Ur guiden, TT-system:' in prs.slides[7].notes_slide.notes_text_frame.text
        if student.name.startswith('v44_03'):
            assert '2,5 MΩ' in prs.slides[6].notes_slide.notes_text_frame.text
        print(f'{student.stem}: {len(prs.slides)} slides, unchanged visible XML, protected notes verified')


if __name__ == '__main__':
    if '--before' in sys.argv:
        # Inspect actual published binaries without exposing their private solutions.
        for student, teacher in decks():
            data = krypto.dekryptera(teacher, os.environ['LARARLOSEN'])
            prs = Presentation(io.BytesIO(data))
            notes = '\n'.join(s.notes_slide.notes_text_frame.text for s in prs.slides)
            print(f'{student.stem}: published teacher slides={len(prs.slides)}, '
                  f'current student slides={len(Presentation(student).slides)}, '
                  f'old guide example in notes={any(t in notes for t in ("Kapacitans 1", "Friskt nät:", "Visar övervakningen den sämsta fasen?"))}')
    else:
        if '--check-only' not in sys.argv:
            before = public_snapshot()
            # Only regenerate Week 44 here. Analysis may be checked separately after a targeted Week 45 build.
            subprocess.run([sys.executable, str(SJO / 'verktyg/larare/lararmanus.py'), '44'], check=True)
            assert public_snapshot() == before, 'Student/assessment source changed'
        check()
