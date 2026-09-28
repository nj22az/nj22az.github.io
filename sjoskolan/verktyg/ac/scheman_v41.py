#!/usr/bin/env python3
"""Vecka 41: diagram och kopplingsscheman i presentationerna som bilder (samma skäl som scheman_v38.py).

v41_01 bild 6: tre sinusspänningar (var ett PowerPoint-diagram).
v41_02 bild 6: Y- och Δ-koppling (var lösa linjer).
v41_03 bild 21: Station C, hållkretsen (var lösa linjer).

Formerna i figurens område (diagram, linjer och deras etikettrutor) tas bort och ersätts av en bild på samma plats.
Rubrik, text och formel på bilden behålls.

    python3 sjoskolan/verktyg/ac/scheman_v41.py
"""
import math
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent))
from scheman_v38 import INK, BLA, GRON, RITA, instrument, ln, motstand, prick, t  # noqa: E402,F401

SJO = Path(__file__).resolve().parents[2]
V41 = SJO / 'vecka-41' / 'aktuell'
UT = V41 / 'figurer'
ORANGE = '#a94d0a'
EMU = 914400


def svg(b, h, inner):
    return f'<svg xmlns="http://www.w3.org/2000/svg" width="{b}" height="{h}" viewBox="0 0 {b} {h}"><rect width="{b}" height="{h}" fill="#fff"/>{inner}</svg>'


def sinus():
    """L1, L2 och L3 över 20 ms (50 Hz), normaliserade, 120° isär."""
    b, h, v, r, top, bot = 880, 300, 70, 110, 30, 250
    X = lambda ms: v + ms / 20 * (b - v - r)
    Y = lambda u: top + (1 - u) / 2 * (bot - top)
    s = ''
    for ms in (0, 5, 10, 15, 20):
        s += ln((X(ms), top), (X(ms), bot), farg='#cad8e2', w=1.5) + t(X(ms), bot + 24, str(ms), 17)
    for u, txt in ((1, '1'), (0, '0'), (-1, '−1')):
        s += ln((v, Y(u)), (X(20), Y(u)), farg='#cad8e2', w=1.5) + t(v - 10, Y(u) + 6, txt, 17, anchor='end')
    s += t(X(20), bot + 48, 'Tid (ms)', 17, anchor='end') + t(v - 50, top - 8, 'u/û', 17, anchor='start')
    for k, (namn, farg) in enumerate((('L1', BLA), ('L2', ORANGE), ('L3', GRON))):
        pts = [(X(ms), Y(math.sin(2 * math.pi * 50 * ms / 1000 - k * 2 * math.pi / 3))) for ms in [i / 10 for i in range(0, 201)]]
        s += ln(*[(round(x, 1), round(y, 1)) for x, y in pts], farg=farg, w=3.5)
        s += ln((b - r + 30, 80 + k * 40), (b - r + 60, 80 + k * 40), farg=farg, w=4) + t(b - r + 70, 87 + k * 40, namn, 19, farg, 'start', 700)
    return svg(b, h + 30, s)


def gren(x1, y1, x2, y2):
    """En lastgren: ledare med ett motstånd på mitten, längs linjen."""
    dx, dy = x2 - x1, y2 - y1
    L = math.hypot(dx, dy)
    ux, uy = dx / L, dy / L
    a = (x1 + ux * (L / 2 - 30), y1 + uy * (L / 2 - 30))
    c = (x1 + ux * (L / 2 + 30), y1 + uy * (L / 2 + 30))
    vinkel = math.degrees(math.atan2(dy, dx))
    mx, my = (x1 + x2) / 2, (y1 + y2) / 2
    return (ln((x1, y1), a) + ln(c, (x2, y2))
            + f'<rect x="{mx - 30}" y="{my - 11}" width="60" height="22" fill="#fff" stroke="{INK}" stroke-width="3" transform="rotate({vinkel:.1f} {mx} {my})"/>')


def y_delta():
    """Y: tre grenar mellan fas och stjärnpunkt. Δ: tre grenar mellan två faser."""
    b, h = 880, 330
    s = ''
    # Y till vänster
    cx, cy = 220, 170
    p = {'L1': (90, 50), 'L2': (350, 50), 'L3': (220, 300)}
    for namn, (x, y) in p.items():
        s += gren(x, y, cx, cy) + prick(x, y)
    s += prick(cx, cy) + t(90, 36, 'L1', 20, vikt=700) + t(350, 36, 'L2', 20, vikt=700) + t(248, 306, 'L3', 20, anchor='start', vikt=700)
    s += t(cx + 18, cy + 6, 'N', 18, BLA, 'start') + t(cx, 325, 'Y', 24, BLA, vikt=700)
    # Δ till höger
    q = {'L1': (530, 60), 'L2': (790, 60), 'L3': (660, 285)}
    s += gren(*q['L1'], *q['L2']) + gren(*q['L2'], *q['L3']) + gren(*q['L3'], *q['L1'])
    for x, y in q.values():
        s += prick(x, y)
    s += t(530, 44, 'L1', 20, vikt=700) + t(790, 44, 'L2', 20, vikt=700) + t(688, 292, 'L3', 20, anchor='start', vikt=700)
    s += t(660, 325, 'Δ', 24, BLA, vikt=700)
    return svg(b, h + 10, s)


def kontakt(x, y, nc=False):
    """Kontakt mellan x och x + 60 på höjden y. NO: öppen i vila. NC: sluten i vila (bladet vilar mot anslaget)."""
    s = ln((x, y), (x + 8, y)) + ln((x + 52, y), (x + 60, y)) + prick(x + 8, y)
    if nc:
        s += ln((x + 52, y), (x + 52, y - 16)) + ln((x + 8, y), (x + 58, y - 12))
    else:
        s += ln((x + 8, y), (x + 50, y - 22))
    return s


def hallkrets():
    """+12 V – S0 STOPP (NC) – S1 START (NO) parallellt med K1 (NO) – spolen K1 – 0 V."""
    b, h, y, y2 = 880, 300, 110, 220
    s = t(40, y - 20, '+12 V', 20, anchor='start', vikt=700) + t(840, y - 20, '0 V', 20, anchor='end', vikt=700)
    s += ln((40, y), (150, y)) + kontakt(150, y, nc=True) + ln((210, y), (330, y))
    s += kontakt(330, y) + ln((390, y), (560, y))
    s += ln((290, y), (290, y2), (330, y2)) + kontakt(330, y2) + ln((390, y2), (500, y2), (500, y))
    s += prick(290, y) + prick(500, y)
    s += f'<rect x="560" y="{y - 26}" width="90" height="52" fill="#fff" stroke="{INK}" stroke-width="3"/>' + t(605, y + 8, 'K1', 22, vikt=700)
    s += ln((650, y), (840, y))
    s += t(180, y + 42, 'S0 STOPP', 18, BLA, vikt=700) + t(180, y + 64, 'NC, sluten i vila', 16)
    s += t(360, y - 58, 'S1 START', 18, BLA, vikt=700) + t(360, y - 36, 'NO, öppen i vila', 16)
    s += t(360, y2 + 36, 'K1 hållkontakt', 18, BLA, vikt=700) + t(360, y2 + 58, 'NO, sluts när K1 drar', 16)
    s += t(605, y + 56, 'Spole K1', 18, BLA, vikt=700)
    return svg(b, h, s)


FIGURER = [
    ('v41_01_Trefassystemets_grunder_elev.pptx', 6, 'sinus', sinus),
    ('v41_02_Y_och_trefaseffekt_elev.pptx', 6, 'y-delta', y_delta),
    ('v41_03_Fysisk_traff_och_matning_elev.pptx', 21, 'hallkrets', hallkrets),
]


def ar_ritning(sh):
    tag = sh._element.tag.split('}')[1]
    prst = sh._element.xpath('string(.//a:prstGeom/@prst)')
    return tag in ('graphicFrame', 'cxnSp') or bool(sh._element.xpath('.//a:custGeom')) or prst in ('line', 'straightConnector1', 'ellipse') \
        or sh.name.startswith('schema-')


def inuti(sh, box, marg=0.15 * EMU):
    x0, y0, x1, y1 = box
    return sh.left >= x0 - marg and sh.top >= y0 - marg and sh.left + sh.width <= x1 + marg and sh.top + sh.height <= y1 + marg


def main():
    UT.mkdir(exist_ok=True)
    filer = []
    for _, n, namn, f in FIGURER:
        fil = UT / f'{namn}.svg'
        fil.write_text(f())
        filer.append(str(fil))
    rot = subprocess.run(['npm', 'root', '-g'], capture_output=True, text=True, check=True).stdout.strip()
    skript = UT / '.rita.cjs'
    skript.write_text(RITA)
    try:
        subprocess.run(['node', str(skript), rot, *filer], check=True)
    finally:
        skript.unlink()
    for deck, n, namn, f in FIGURER:
        fil = V41 / deck
        prs = Presentation(fil)
        bild = prs.slides[n - 1]
        # Figurbandet: mellan brödtexten (slutar 2,6 tum) och formeln (börjar 5,24 tum). Allt som börjar i bandet hör
        # till figuren (diagram, linjer, etiketter, bildtext) och ersätts av bilden.
        y0, y1 = int(2.55 * EMU), int(5.2 * EMU)
        bort = [sh for sh in bild.shapes if y0 <= sh.top < y1]
        if not any(ar_ritning(sh) for sh in bort):
            raise SystemExit(f'{deck} bild {n}: ingen ritning i figurbandet')
        for sh in bort:
            for rid in sh._element.xpath('.//c:chart/@r:id'):
                bild.part.drop_rel(rid)
            sh._element.getparent().remove(sh._element)
        box = (int(0.5 * EMU), int(2.62 * EMU), int(9.5 * EMU), int(5.18 * EMU))
        from PIL import Image
        im = Image.open(UT / f'{namn}.png')
        w, h = box[2] - box[0], box[3] - box[1]
        skala = min(w / im.width, h / im.height)
        pw, ph = int(im.width * skala), int(im.height * skala)
        pic = bild.shapes.add_picture(str(UT / f'{namn}.png'), Emu(box[0] + (w - pw) // 2), Emu(box[1] + (h - ph) // 2), Emu(pw), Emu(ph))
        pic.name = f'schema-{namn}'
        pic._element.nvPicPr.cNvPr.set('descr', f.__doc__)
        prs.save(fil)
        print(f'{deck} bild {n}: {len(bort)} former ersatta av bilden {namn}.png')


if __name__ == '__main__':
    sys.exit(main())
