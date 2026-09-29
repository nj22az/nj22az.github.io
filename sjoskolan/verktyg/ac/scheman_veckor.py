#!/usr/bin/env python3
"""Vecka 41–45: diagram och kopplingsscheman i presentationerna som bilder (samma skäl som scheman_v38.py).

v41_01 bild 6: tre sinusspänningar; figuren sätts nu ihop med generatorn i 3D av verktyg/figurer/figurer3d.py (sinus() används därifrån).
v41_02 bild 6: Y- och Δ-koppling. v43_03 bild 8: hållkretsen (v41_03 bild 21 visar i stället startarens delar, verktyg/figurer/figs41.py).
v43_03 bild 6: symboler. v43_03 bild 7: huvudströmsschema för en motor. v44_01 bild 6: enlinjeschema.

Formerna i figurens område (diagram, linjer och deras etikettrutor) tas bort och ersätts av en bild på samma plats.
Rubrik, text och formel på bilden behålls.

    python3 sjoskolan/verktyg/ac/scheman_veckor.py
"""
import math
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent))
from scheman_v38 import INK, BLA, GRON, RITA, bild6, instrument, ln, motstand, prick, t  # noqa: E402,F401

SJO = Path(__file__).resolve().parents[2]
ORANGE = '#a94d0a'
EMU = 914400


def svg(b, h, inner):
    return f'<svg xmlns="http://www.w3.org/2000/svg" width="{b}" height="{h}" viewBox="0 0 {b} {h}"><rect width="{b}" height="{h}" fill="#fff"/>{inner}</svg>'


def sinus():
    """L1, L2 och L3 över 20 ms (50 Hz) med toppvärdet 1, 120° isär. Streckad: linjespänningen U12 = L1 − L2 med toppvärdet √3 ≈ 1,73."""
    b, h, v, r, top, bot = 880, 330, 70, 130, 36, 280
    R3 = math.sqrt(3)
    X = lambda ms: v + ms / 20 * (b - v - r)
    Y = lambda u: top + (R3 - u) / (2 * R3) * (bot - top)
    s = ''
    for ms in (0, 5, 10, 15, 20):
        s += ln((X(ms), top), (X(ms), bot), farg='#cad8e2', w=1.5) + t(X(ms), bot + 24, str(ms), 17)
    for u, txt in ((R3, '√3'), (1, '1'), (0, '0'), (-1, '−1'), (-R3, '−√3')):
        s += ln((v, Y(u)), (X(20), Y(u)), farg='#cad8e2', w=1.5) + t(v - 10, Y(u) + 6, txt, 17, anchor='end')
    s += t(X(20), bot + 48, 'Tid (ms)', 17, anchor='end') + t(v - 50, top - 16, 'u/û', 17, anchor='start')
    kurvor = (('L1', BLA, 1, 0, ''), ('L2', ORANGE, 1, 1, ''), ('L3', GRON, 1, 2, ''), ('U₁₂', INK, R3, None, '8 6'))
    for k, (namn, farg, a, fas, streck) in enumerate(kurvor):
        f = (lambda x: math.sin(x) - math.sin(x - 2 * math.pi / 3)) if fas is None else (lambda x, fas=fas: math.sin(x - fas * 2 * math.pi / 3))
        pts = [(X(ms), Y(f(2 * math.pi * 50 * ms / 1000))) for ms in [i / 10 for i in range(0, 201)]]
        linje = ln(*[(round(x, 1), round(y, 1)) for x, y in pts], farg=farg, w=3 if streck else 3.5)
        if streck:
            linje = linje.replace('/>', f' stroke-dasharray="{streck}"/>', 1)
        s += linje
        mark = ln((b - r + 30, 70 + k * 40), (b - r + 60, 70 + k * 40), farg=farg, w=4)
        if streck:
            mark = mark.replace('/>', f' stroke-dasharray="6 4"/>', 1)
        s += mark + t(b - r + 70, 77 + k * 40, namn, 19, farg, 'start', 700)
    return svg(b, h + 30, s)


def gren(x1, y1, x2, y2, farg=INK):
    """En lastgren: ledare med ett motstånd på mitten, längs linjen."""
    dx, dy = x2 - x1, y2 - y1
    L = math.hypot(dx, dy)
    ux, uy = dx / L, dy / L
    a = (x1 + ux * (L / 2 - 30), y1 + uy * (L / 2 - 30))
    c = (x1 + ux * (L / 2 + 30), y1 + uy * (L / 2 + 30))
    vinkel = math.degrees(math.atan2(dy, dx))
    mx, my = (x1 + x2) / 2, (y1 + y2) / 2
    return (ln((x1, y1), a, farg=farg) + ln(c, (x2, y2), farg=farg)
            + f'<rect x="{mx - 30}" y="{my - 11}" width="60" height="22" fill="#fff" stroke="{farg}" stroke-width="3" transform="rotate({vinkel:.1f} {mx} {my})"/>')


def ti(x, y, bas, index, size=26, farg=INK, anchor='middle', vikt=400):
    """Text med nedsänkt index, till exempel U_{gren}."""
    return t(x, y, f'{bas}<tspan dy="{size * 0.25:.0f}" font-size="{size * 0.7:.0f}">{index}</tspan>', size, farg, anchor, vikt)


def pil2(p1, p2, farg=BLA):
    """Dubbelpil för en spänning mellan två punkter."""
    (x1, y1), (x2, y2) = p1, p2
    L = math.hypot(x2 - x1, y2 - y1); ux, uy = (x2 - x1) / L, (y2 - y1) / L; nx, ny = -uy, ux
    spets = lambda x, y, sx, sy: f'<path d="M{x},{y} L{x - 14 * sx + 6 * nx:.1f},{y - 14 * sy + 6 * ny:.1f} L{x - 14 * sx - 6 * nx:.1f},{y - 14 * sy - 6 * ny:.1f} Z" fill="{farg}"/>'
    return ln(p1, p2, farg=farg, w=2.5) + spets(x2, y2, ux, uy) + spets(x1, y1, -ux, -uy)


def y_delta():
    """Y: varje gren ligger mellan en fas och stjärnpunkten, U_{gren} = U_{L}/√3. Δ: varje gren ligger mellan två faser, U_{gren} = U_{L}."""
    b, h = 880, 410
    s = ''
    # Y till vänster
    cx, cy = 220, 215
    p = {'L1': (80, 95), 'L2': (360, 95), 'L3': (220, 355)}
    for namn, (x, y) in p.items():
        s += gren(x, y, cx, cy, BLA if namn == 'L1' else INK) + prick(x, y)
    s += prick(cx, cy) + t(66, 101, 'L1', 32, anchor='end', vikt=700) + t(374, 101, 'L2', 32, anchor='start', vikt=700) + t(240, 350, 'L3', 32, anchor='start', vikt=700)
    s += t(cx + 22, cy + 8, 'stjärnpunkt', 24, '#6b6b6b', 'start')
    s += pil2((92, 55), (348, 55)) + ti(220, 36, 'U', 'L', 34, BLA)
    s += ti(122, 205, 'U', 'gren', 34, BLA, 'end') + t(220, 395, 'Y', 34, BLA, vikt=700)
    # Δ till höger
    q = {'L1': (530, 95), 'L2': (810, 95), 'L3': (670, 335)}
    s += gren(*q['L1'], *q['L2'], BLA) + gren(*q['L2'], *q['L3']) + gren(*q['L3'], *q['L1'])
    for x, y in q.values():
        s += prick(x, y)
    s += t(516, 101, 'L1', 32, anchor='end', vikt=700) + t(824, 101, 'L2', 32, anchor='start', vikt=700) + t(690, 342, 'L3', 32, anchor='start', vikt=700)
    s += pil2((542, 55), (798, 55)) + ti(670, 36, 'U', 'L', 34, BLA)
    s += ti(670, 140, 'U', 'gren', 34, BLA) + t(670, 395, 'Δ', 34, BLA, vikt=700)
    return svg(b, h, s)


def kontakt(x, y, nc=False):
    """Kontakt mellan x och x + 60 på höjden y. NO: öppen i vila. NC: sluten i vila (bladet vilar mot anslaget)."""
    s = ln((x, y), (x + 8, y)) + ln((x + 52, y), (x + 60, y)) + prick(x + 8, y)
    if nc:
        s += ln((x + 52, y), (x + 52, y - 16)) + ln((x + 8, y), (x + 58, y - 12))
    else:
        s += ln((x + 8, y), (x + 50, y - 22))
    return s


def hallkrets():
    """+12 V – S0 STOPP (NC) – S1 START (NO) parallellt med K1 (NO) – spolen K1 – 0 V."""
    b, h, y, y2 = 880, 300, 110, 220
    s = t(40, y - 20, '+12 V', 20, anchor='start', vikt=700) + t(840, y - 20, '0 V', 20, anchor='end', vikt=700)
    s += ln((40, y), (150, y)) + kontakt(150, y, nc=True) + ln((210, y), (330, y))
    s += kontakt(330, y) + ln((390, y), (560, y))
    s += ln((290, y), (290, y2), (330, y2)) + kontakt(330, y2) + ln((390, y2), (500, y2), (500, y))
    s += prick(290, y) + prick(500, y)
    s += f'<rect x="560" y="{y - 26}" width="90" height="52" fill="#fff" stroke="{INK}" stroke-width="3"/>' + t(605, y + 8, 'K1', 22, vikt=700)
    s += ln((650, y), (840, y))
    s += t(180, y + 42, 'S0 STOPP', 18, BLA, vikt=700) + t(180, y + 64, 'NC, sluten i vila', 16)
    s += t(360, y - 58, 'S1 START', 18, BLA, vikt=700) + t(360, y - 36, 'NO, öppen i vila', 16)
    s += t(360, y2 + 36, 'K1 hållkontakt', 18, BLA, vikt=700) + t(360, y2 + 58, 'NO, sluts när K1 drar', 16)
    s += t(605, y + 56, 'Spole K1', 18, BLA, vikt=700)
    return svg(b, h, s)


def huvudstrom():
    """Huvudströmsschema: L1–L3 genom säkringar F1, kontaktor K1 och överlastskydd F2 till motorn M 3~."""
    b, h = 880, 330
    s = ''
    for k, y in enumerate((60, 120, 180)):
        s += t(40, y + 7, f'L{k + 1}', 20, vikt=700) + ln((70, y), (160, y))
        s += f'<rect x="160" y="{y - 11}" width="60" height="22" fill="none" stroke="{INK}" stroke-width="3"/>' + ln((150, y), (230, y)) + ln((230, y), (330, y))
        s += prick(330, y) + ln((330, y), (382, y - 24)) + prick(392, y) + ln((392, y), (470, y))
        s += f'<rect x="470" y="{y - 11}" width="60" height="22" fill="#fff" stroke="{INK}" stroke-width="3"/>'
        s += f'<path d="M478,{y} h10 v-6 h12 v12 h12 v-6 h10" fill="none" stroke="{INK}" stroke-width="2"/>'
        s += ln((530, y), (680 - math.sqrt(max(0, 70 ** 2 - (y - 120) ** 2)), y))
    s += f'<circle cx="720" cy="120" r="70" fill="#fff" stroke="{INK}" stroke-width="3"/>' + t(720, 118, 'M', 34, vikt=700) + t(720, 150, '3~', 20)
    s += ln((330, 36), (330, 204), farg='#9fb4c4', w=1.5).replace('stroke-linecap="round"', 'stroke-dasharray="6 5"')
    s += t(190, 240, 'F1', 20, BLA, vikt=700) + t(190, 264, 'Säkringar', 17)
    s += t(360, 240, 'K1', 20, BLA, vikt=700) + t(360, 264, 'Kontaktor', 17)
    s += t(500, 240, 'F2', 20, BLA, vikt=700) + t(500, 264, 'Överlastskydd', 17)
    s += t(720, 240, 'Motor', 17)
    return svg(b, h, s)


def enlinje():
    """Enlinjeschema: källa, skydd, fördelning och last i en kedja, med mätpunkt A och B."""
    b, h = 880, 260
    s = t(440, 30, 'Enlinjemodell', 20, vikt=700)
    lador = [(40, 'Källa'), (250, 'Skydd'), (460, 'Fördelning'), (670, 'Last')]
    for x, namn in lador:
        s += f'<rect x="{x}" y="90" width="170" height="80" rx="6" fill="#edf4f9" stroke="{BLA}" stroke-width="3"/>' + t(x + 85, 138, namn, 20, vikt=700)
    for x in (210, 420, 630):
        s += ln((x, 130), (x + 40, 130))
    for x, namn in ((230, 'Mätpunkt A'), (650, 'Mätpunkt B')):
        s += prick(x, 130) + ln((x, 125), (x, 72), farg=GRON, w=2) + t(x, 64, namn, 17, GRON, vikt=700)
    s += t(440, 220, 'Driftläge och matningsväg avgör vilken väg energin tar.', 17)
    return svg(b, h, s)


def symboler():
    """Grafiska symboler: motstånd, NO-kontakt, spänningsmätare, säkring, reläspole K1 och strömmätare."""
    return bild6()


# (vecka, presentation, bild, namn, ritfunktion)
FIGURER = [
    # v41_01 bild 6 (sinus) ritas av verktyg/figurer/figurer3d.py tillsammans med generatorn i 3D.
    (41, 'v41_02_Y_och_trefaseffekt_elev.pptx', 6, 'y-delta', y_delta),
    (43, 'v43_03_Elscheman_och_dokumentation_elev.pptx', 6, 'symboler', symboler),
    (43, 'v43_03_Elscheman_och_dokumentation_elev.pptx', 7, 'huvudstrom', huvudstrom),
    (43, 'v43_03_Elscheman_och_dokumentation_elev.pptx', 8, 'hallkrets', hallkrets),
    (44, 'v44_01_Lagspanningssystem_elev.pptx', 6, 'enlinje', enlinje),
]


def ar_ritning(sh):
    tag = sh._element.tag.split('}')[1]
    prst = sh._element.xpath('string(.//a:prstGeom/@prst)')
    return tag in ('graphicFrame', 'cxnSp') or bool(sh._element.xpath('.//a:custGeom')) or prst in ('line', 'straightConnector1', 'ellipse') \
        or sh.name.startswith('schema-')


def inuti(sh, box, marg=0.15 * EMU):
    x0, y0, x1, y1 = box
    return sh.left >= x0 - marg and sh.top >= y0 - marg and sh.left + sh.width <= x1 + marg and sh.top + sh.height <= y1 + marg


def main():
    rot = subprocess.run(['npm', 'root', '-g'], capture_output=True, text=True, check=True).stdout.strip()
    for vecka in sorted({v for v, *_ in FIGURER}):
        UT = SJO / f'vecka-{vecka}' / 'aktuell' / 'figurer'
        UT.mkdir(exist_ok=True)
        filer = []
        for v, _, n, namn, f in FIGURER:
            if v == vecka:
                fil = UT / f'{namn}.svg'
                fil.write_text(f())
                filer.append(str(fil))
        skript = UT / '.rita.cjs'
        skript.write_text(RITA)
        try:
            subprocess.run(['node', str(skript), rot, *filer], check=True)
        finally:
            skript.unlink()
    for vecka, deck, n, namn, f in FIGURER:
        UT = SJO / f'vecka-{vecka}' / 'aktuell' / 'figurer'
        fil = SJO / f'vecka-{vecka}' / 'aktuell' / deck
        prs = Presentation(fil)
        bild = prs.slides[n - 1]
        # Figurbandet: mellan brödtexten (slutar 2,6 tum) och formeln (börjar 5,24 tum). Allt som börjar i bandet hör
        # till figuren (diagram, linjer, etiketter, bildtext) och ersätts av bilden.
        y0, y1 = int(2.55 * EMU), int(5.2 * EMU)
        bort = [sh for sh in bild.shapes if y0 <= sh.top < y1]
        if not any(ar_ritning(sh) for sh in bort):
            raise SystemExit(f'{deck} bild {n}: ingen ritning i figurbandet')
        for sh in bort:
            for rid in sh._element.xpath('.//c:chart/@r:id'):
                bild.part.drop_rel(rid)
            sh._element.getparent().remove(sh._element)
        box = (int(0.5 * EMU), int(2.62 * EMU), int(9.5 * EMU), int(5.18 * EMU))
        from PIL import Image
        im = Image.open(UT / f'{namn}.png')
        w, h = box[2] - box[0], box[3] - box[1]
        skala = min(w / im.width, h / im.height)
        pw, ph = int(im.width * skala), int(im.height * skala)
        pic = bild.shapes.add_picture(str(UT / f'{namn}.png'), Emu(box[0] + (w - pw) // 2), Emu(box[1] + (h - ph) // 2), Emu(pw), Emu(ph))
        pic.name = f'schema-{namn}'
        pic._element.nvPicPr.cNvPr.set('descr', f.__doc__)
        prs.save(fil)
        print(f'{deck} bild {n}: {len(bort)} former ersatta av bilden {namn}.png')


if __name__ == '__main__':
    sys.exit(main())
