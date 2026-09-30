#!/usr/bin/env python3
"""Frånskiljning och mätteknik (v38_01): kopplingsscheman och symboler som bilder.

Figurerna på bild 4, 6, 11 och 12 var byggda av lösa linjer, rutor och textrutor. De såg rätt ut i LibreOffice men fel
i PowerPoint. Här ritas samma figurer som SVG, renderas till PNG och ersätter formerna i figurområdet, så att de ser
likadana ut i PowerPoint, Keynote och LibreOffice. Rubrik, sidfot och förklarande text på bilden behålls.

    python3 sjoskolan/verktyg/ac/scheman_v38.py

Omslaget (bild 1) är en ritad effektbrytare i FRÅN-läge med hänglås (verktyg/ac/brytare.mjs).
"""
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

SJO = Path(__file__).resolve().parents[2]
DECK = SJO / 'vecka-38' / 'aktuell' / 'v38_01_Franskiljning_och_matteknik_elev.pptx'
UT = SJO / 'vecka-38' / 'aktuell' / 'figurer'
INK, BLA, GRON = '#163248', '#064f91', '#176844'
B, H = 900, 380  # 100 px per tum
FONT = 'font-family="Arial, Helvetica, sans-serif"'


def t(x, y, s, size=22, farg=INK, anchor='middle', vikt=400):
    return f'<text x="{x}" y="{y}" font-size="{size}" fill="{farg}" text-anchor="{anchor}" font-weight="{vikt}" {FONT}>{s}</text>'


def ln(*p, farg=INK, w=3):
    d = 'M' + ' L'.join(f'{x},{y}' for x, y in p)
    return f'<path d="{d}" fill="none" stroke="{farg}" stroke-width="{w}" stroke-linejoin="round" stroke-linecap="round"/>'


def motstand(cx, cy, farg=INK):  # vågrätt, 64 × 26
    return f'<rect x="{cx - 32}" y="{cy - 13}" width="64" height="26" fill="#fff" stroke="{farg}" stroke-width="3"/>'


def instrument(cx, cy, bokstav, farg, r=30):
    return f'<circle cx="{cx}" cy="{cy}" r="{r}" fill="#fff" stroke="{farg}" stroke-width="3"/>' + t(cx, cy + 8, bokstav, 24, farg, vikt=700)


def prick(x, y):
    return f'<circle cx="{x}" cy="{y}" r="5" fill="{INK}"/>'


def svg(inner):
    return f'<svg xmlns="http://www.w3.org/2000/svg" width="{B}" height="{H}" viewBox="0 0 {B} {H}"><rect width="{B}" height="{H}" fill="#fff"/>{inner}</svg>'


def bild4():
    """Två källor via förreglad omkopplare A / 0 / B till lasten."""
    def ruta(x, y, w, h, rad1, rad2=''):
        s = f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="6" fill="#edf4f9" stroke="{BLA}" stroke-width="3"/>'
        return s + t(x + w / 2, y + (h / 2 - 4 if rad2 else h / 2 + 8), rad1, 22, INK, vikt=700) + (t(x + w / 2, y + h / 2 + 24, rad2, 19, INK) if rad2 else '')
    s = ruta(20, 40, 230, 90, 'Källa A', 'Generator') + ruta(20, 250, 230, 90, 'Källa B', 'Reservmatning')
    s += ruta(360, 140, 240, 100, 'Omkopplare', 'A / 0 / B, förreglad') + ruta(700, 145, 180, 90, 'Last')
    s += ln((250, 85), (305, 85), (305, 170), (360, 170)) + ln((250, 295), (305, 295), (305, 210), (360, 210)) + ln((600, 190), (700, 190))
    return svg(s)


def bild6():
    """Sex symboler i två rader med förklaring under."""
    s = ''
    cols, y1, y2 = (150, 450, 750), 110, 280
    # R: motstånd
    s += ln((cols[0] - 70, y1), (cols[0] - 32, y1)) + motstand(cols[0], y1) + ln((cols[0] + 32, y1), (cols[0] + 70, y1))
    s += t(cols[0], y1 + 60, 'R: motstånd')
    # S: NO-kontakt (slutande), öppen i vila
    x = cols[1]
    s += ln((x - 70, y1), (x - 30, y1)) + prick(x - 30, y1) + ln((x - 30, y1), (x + 26, y1 - 30)) + prick(x + 30, y1) + ln((x + 30, y1), (x + 70, y1))
    s += t(x, y1 + 60, 'S: NO-kontakt (slutande)')
    # V
    s += ln((cols[2] - 70, y1), (cols[2] - 30, y1)) + instrument(cols[2], y1, 'V', INK) + ln((cols[2] + 30, y1), (cols[2] + 70, y1))
    s += t(cols[2], y1 + 60, 'Spänningsmätare')
    # F: säkring
    s += ln((cols[0] - 70, y2), (cols[0] + 70, y2)) + f'<rect x="{cols[0] - 32}" y="{y2 - 13}" width="64" height="26" fill="none" stroke="{INK}" stroke-width="3"/>'
    s += t(cols[0], y2 + 60, 'F: säkring')
    # K: reläspole
    x = cols[1]
    s += ln((x - 70, y2), (x - 30, y2)) + f'<rect x="{x - 30}" y="{y2 - 22}" width="60" height="44" fill="#fff" stroke="{INK}" stroke-width="3"/>' + ln((x + 30, y2), (x + 70, y2))
    s += t(x + 44, y2 - 30, 'K1', 20, BLA, 'start', 700) + t(x, y2 + 60, 'K: reläspole')
    # A
    s += ln((cols[2] - 70, y2), (cols[2] - 30, y2)) + instrument(cols[2], y2, 'A', INK) + ln((cols[2] + 30, y2), (cols[2] + 70, y2))
    s += t(cols[2], y2 + 60, 'Strömmätare')
    return svg(s)


def bild11():
    """12 V DC, amperemeter i serie, R = 1 kΩ, voltmeter parallellt över R."""
    vx, top, bot, hx = 230, 70, 330, 860
    ax, rx = 420, 660
    s = ''
    # Källa: cirkel med + och −
    s += ln((vx, top), (vx, 170)) + f'<circle cx="{vx}" cy="200" r="30" fill="#fff" stroke="{INK}" stroke-width="3"/>' + ln((vx, 230), (vx, bot))
    s += t(vx, 192, '+', 22, INK, vikt=700) + t(vx, 222, '−', 22, INK, vikt=700) + t(vx - 45, 208, '12 V DC', 22, INK, 'end')
    # Övre ledare med A och R
    s += ln((vx, top), (ax - 30, top)) + instrument(ax, top, 'A', BLA) + ln((ax + 30, top), (rx - 90, top))
    s += ln((rx - 90, top), (rx - 32, top)) + motstand(rx, top) + ln((rx + 32, top), (rx + 90, top)) + ln((rx + 90, top), (hx, top), (hx, bot), (vx, bot))
    s += t(rx, top - 24, 'R = 1 kΩ', 22) + t(ax, top + 62, 'A i serie', 22, BLA, vikt=700)
    # Voltmeter parallellt över R
    vy = 180
    s += ln((rx - 90, top), (rx - 90, vy), (rx - 30, vy), farg=GRON) + instrument(rx, vy, 'V', GRON) + ln((rx + 30, vy), (rx + 90, vy), (rx + 90, top), farg=GRON)
    s += prick(rx - 90, top) + prick(rx + 90, top) + t(rx, vy + 62, 'V parallellt över R', 22, GRON, vikt=700)
    return svg(s)


def bild12():
    """Avgränsat motstånd och ohmmeter i en sluten slinga, ingen matning."""
    lx, rx2, top, bot, cx = 260, 520, 115, 300, 390
    s = t(cx, 36, 'Ingen extern matning ansluten', 24, BLA, vikt=700)
    s += ln((lx, top), (cx - 32, top)) + motstand(cx, top) + ln((cx + 32, top), (rx2, top), (rx2, bot), (cx + 36, bot))
    s += instrument(cx, bot, 'Ω', BLA, 36) + ln((cx - 36, bot), (lx, bot), (lx, top))
    s += t(cx, top - 24, 'R = 1 kΩ', 22) + t(rx2 + 30, 205, 'Mätobjektet är avgränsat.', 22, INK, 'start')
    s += t(cx, bot + 62, 'Ohmmätare (Ω)', 20, BLA)
    return svg(s)


FIGURER = {4: bild4, 6: bild6, 11: bild11, 12: bild12}
# Former som behålls: rubriken, sidfotsbilden, principtexten på bild 4, raden ”Principskiss …” och sidnumret.
BEHALL_TEXT = ('Princip: källorna', 'Principskiss för undervisning', 'Kursvecka')
# Figurens ruta (tum). Bild 4 har principtexten under figuren.
RUTA = {4: (0.5, 1.75, 9.0, 3.75), 6: (0.5, 1.8, 9.0, 4.4), 11: (0.5, 1.75, 9.0, 4.5), 12: (0.5, 1.75, 9.0, 4.5)}

RITA = r"""
const { chromium } = require(process.argv[2] + '/playwright');
const fs = require('fs');
(async () => {
  const b = await chromium.launch(); const p = await b.newPage({ deviceScaleFactor: 3 });
  for (const f of process.argv.slice(3)) {
    await p.setContent('<body style="margin:0">' + fs.readFileSync(f, 'utf8') + '</body>');
    await (await p.$('svg')).screenshot({ path: f.replace(/\.svg$/, '.png') });
  }
  await b.close();
})();
"""


def main():
    UT.mkdir(exist_ok=True)
    filer = []
    for n, f in FIGURER.items():
        fil = UT / f'v38_01_bild{n:02d}.svg'
        fil.write_text(f())
        filer.append(str(fil))
    rot = subprocess.run(['npm', 'root', '-g'], capture_output=True, text=True, check=True).stdout.strip()
    skript = UT / '.rita.cjs'
    skript.write_text(RITA)
    try:
        subprocess.run(['node', str(skript), rot, *filer], check=True)
    finally:
        skript.unlink()
    # Omslaget: effektbrytare i FRÅN-läge med hänglås och skylt (verktyg/ac/brytare.mjs) i stället för båtfotot.
    svg = subprocess.run(['node', str(Path(__file__).with_name('brytare.mjs'))], capture_output=True, text=True, check=True).stdout
    (UT / 'v38_01_omslag.svg').write_text(svg)
    skript.write_text(RITA)
    try:
        subprocess.run(['node', str(skript), rot, str(UT / 'v38_01_omslag.svg')], check=True)
    finally:
        skript.unlink()
    prs = Presentation(DECK)
    forsta = prs.slides[0]
    gammal = next(sh for sh in forsta.shapes if sh.name in ('Bildobjekt 1', 'omslag-brytare'))
    ny = forsta.shapes.add_picture(str(UT / 'v38_01_omslag.png'), gammal.left, gammal.top, gammal.width, gammal.height)
    ny.name = 'omslag-brytare'
    ny._element.nvPicPr.cNvPr.set('descr', 'Effektbrytare i FRÅN-läge, låst med rött hänglås och gul skylt: FRÅNSKILD, MANÖVRERA EJ.')
    gammal._element.addprevious(ny._element)  # samma plats i ordningen: bakom rubriken
    gammal._element.getparent().remove(gammal._element)
    for n in FIGURER:
        bild = prs.slides[n - 1]
        titel = bild.shapes.title
        for sh in list(bild.shapes):
            text = sh.text_frame.text if sh.has_text_frame else ''
            behall = (sh.shape_type is not None and 'PICTURE' in str(sh.shape_type) and not sh.name.startswith('schema-')) \
                or sh.name == 'textruta 3' or (titel is not None and sh.shape_id == titel.shape_id) or text.startswith(BEHALL_TEXT)
            if not behall:
                sh._element.getparent().remove(sh._element)
        x, y, w, hmax = RUTA[n]
        h = min(hmax, w * H / B)
        w = h * B / H
        x = (10 - w) / 2
        pic = bild.shapes.add_picture(str(UT / f'v38_01_bild{n:02d}.png'), Emu(int(x * 914400)), Emu(int(y * 914400)), Emu(int(w * 914400)), Emu(int(h * 914400)))
        pic.name = f'schema-bild{n:02d}'
        pic._element.nvPicPr.cNvPr.set('descr', FIGURER[n].__doc__)
    prs.save(DECK)
    print(f'{DECK.name}: {len(FIGURER)} figurer som bilder')


if __name__ == '__main__':
    sys.exit(main())
