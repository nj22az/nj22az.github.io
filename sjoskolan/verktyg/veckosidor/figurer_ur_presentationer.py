#!/usr/bin/env python3
"""Figurerna i vecka 41–45:s presentationer som filer, för delsidorna (innehall/export/delsidor_veckor.py).

Varje bild i en presentation (utom sidfoten) sparas som vecka-NN/aktuell/figurer/<kapitel>-b<bild>-<k>.<ext>.
Delsidan visar figuren vid teoriavsnittet, exemplet eller övningen. Kör efter att presentationerna har ändrats.

    python3 sjoskolan/verktyg/veckosidor/figurer_ur_presentationer.py
"""
import sys
from pathlib import Path

from pptx import Presentation

SJO = Path(__file__).resolve().parents[2]
VECKOR = range(41, 46)
EMU = 914400


def figurer(bild):
    """Bildens figurer i läsordning: alla bilder utom sidfoten (hela bredden längst ned)."""
    ut = []
    for sh in bild.shapes:
        if sh.shape_type is not None and 'PICTURE' in str(sh.shape_type) and not (sh.top > 6.4 * EMU and sh.width > 9 * EMU):
            ut.append(sh)
    return sorted(ut, key=lambda s: (s.top, s.left))


def namn(deck, nr, k, ext):
    return f'{deck[:6]}-b{nr:02d}-{k}.{ext}'


def main():
    n = 0
    for v in VECKOR:
        mapp = SJO / f'vecka-{v}' / 'aktuell'
        ut = mapp / 'figurer'
        ut.mkdir(exist_ok=True)
        for fil in sorted(mapp.glob('v*_elev.pptx')):
            for nr, bild in enumerate(Presentation(fil).slides, 1):
                if nr == 1:  # omslaget visas inte på delsidan
                    continue
                for k, sh in enumerate(figurer(bild), 1):
                    mal = ut / namn(fil.name, nr, k, sh.image.ext)
                    if not mal.exists() or mal.read_bytes() != sh.image.blob:
                        mal.write_bytes(sh.image.blob)
                        n += 1
    print(f'figurer: {n} filer skrivna')


if __name__ == '__main__':
    sys.exit(main())
