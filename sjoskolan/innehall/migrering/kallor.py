"""Läser de gamla källorna vid migreringen. Används bara av migrering/extrahera.py.

Källfilerna läses ur git vid revisionen före migreringen (KALLREV), så att extraktionen går att upprepa.
Boken läses ur den dekrypterade EPUB-filen (BOK_EPUB_DIR), som aldrig checkas in.
"""
import html as _html
import io
import os
import re
import subprocess
from pathlib import Path

REPO = Path(__file__).resolve().parents[3]
KALLREV = '8cabde3'
KAPITEL = {  # kursens kapitel-id -> bokens kapitelnummer
    'v39_01': 4, 'v39_02': 5, 'v40_01': 7, 'v40_02': 8, 'v40_03': 9, 'v41_01': 10, 'v41_02': 11, 'v41_03': 12,
    'v42_01': 13, 'v42_02': 14, 'v42_03': 15, 'v43_01': 16, 'v43_02': 17, 'v43_03': 18,
    'v44_01': 19, 'v44_02': 20, 'v44_03': 21, 'v45_01': 22, 'v45_02': 23, 'v45_03': 24,
}
BOK_KAPITEL_FIL = {k: f'ch{k + 5 if k <= 6 else (k + 6 if k <= 12 else (k + 7 if k <= 18 else k + 8)):03d}.xhtml' for k in range(1, 25)}
BOK_LOSNING_FIL = {k: f'ch{33 + k:03d}.xhtml' for k in range(1, 25)}
DECK = {
    'v39_01': 'vecka-39/aktuell/v39_01_Effekt_och_energi_elev.pptx',
    'v39_02': 'vecka-39/aktuell/v39_02_Kirchhoffs_lagar_med_nodforklaring_v2_elev.pptx',
    'v41_01': 'vecka-41/aktuell/v41_01_Trefassystemets_grunder_elev.pptx',
    'v41_02': 'vecka-41/aktuell/v41_02_Y_och_trefaseffekt_elev.pptx',
    'v41_03': 'vecka-41/aktuell/v41_03_Fysisk_traff_och_matning_elev.pptx',
    'v42_01': 'vecka-42/aktuell/v42_01_Elektriska_risker_elev.pptx',
    'v42_02': 'vecka-42/aktuell/v42_02_Regler_ansvar_och_arbetsmetoder_elev.pptx',
    'v42_03': 'vecka-42/aktuell/v42_03_Riskbedomning_och_skydd_elev.pptx',
    'v43_01': 'vecka-43/aktuell/v43_01_Komponenter_och_skydd_elev.pptx',
    'v43_02': 'vecka-43/aktuell/v43_02_Transformatorer_och_motorer_elev.pptx',
    'v43_03': 'vecka-43/aktuell/v43_03_Elscheman_och_dokumentation_elev.pptx',
    'v44_01': 'vecka-44/aktuell/v44_01_Lagspanningssystem_elev.pptx',
    'v44_02': 'vecka-44/aktuell/v44_02_Hogspanningssystem_elev.pptx',
    'v44_03': 'vecka-44/aktuell/v44_03_Fordjupad_matteknik_elev.pptx',
    'v45_01': 'vecka-45/aktuell/v45_01_Systematisk_felsokning_elev.pptx',
    'v45_02': 'vecka-45/aktuell/v45_02_Analys_av_matresultat_elev.pptx',
    'v45_03': 'vecka-45/aktuell/v45_03_Repetition_infor_tentamen_elev.pptx',
}


def git_bytes(path, rev=KALLREV):
    return subprocess.run(['git', 'show', f'{rev}:sjoskolan/{path}'], cwd=REPO, capture_output=True, check=True).stdout


def git_text(path, rev=KALLREV):
    return git_bytes(path, rev).decode('utf-8')


def inline(fragment):
    """XHTML/HTML-fragment -> innehållstext med _{…}/^{…}."""
    s = re.sub(r'<sub>(.*?)</sub>', lambda m: '_{' + re.sub(r'<[^>]+>', '', m.group(1)) + '}', fragment, flags=re.S)
    s = re.sub(r'<sup>(.*?)</sup>', lambda m: '^{' + re.sub(r'<[^>]+>', '', m.group(1)) + '}', s, flags=re.S)
    s = re.sub(r'<br\s*/?>', ' ', s)
    s = re.sub(r'<[^>]+>', '', s)
    s = _html.unescape(s)
    return re.sub(r'\s+', ' ', s).strip()


# ---------------------------------------------------------------- kursens övningssidor
def kurs_ovningar():
    out = {}
    for vecka in range(39, 46):
        s = git_text(f'vecka-{vecka}/aktuell/Formelstod_och_ovningar.html')
        for aid, body in re.findall(r'<article class="exercise" id="([^"]+)">(.*?)</article>', s, flags=re.S):
            h3 = re.search(r'<h3>(.*?)</h3>', body, flags=re.S).group(1)
            m = re.match(r'Övning (\d+): (.*)', inline(h3))
            fields = {'nr': int(m.group(1)), 'titel': m.group(2)}
            ps = re.findall(r'<p( class="formula")?>(.*?)</p>', re.sub(r'<details.*?</details>', '', body, flags=re.S), flags=re.S)
            fields['fraga'] = inline(ps[0][1])
            for cls, p in ps[1:]:
                t = inline(p)
                if cls:
                    fields['samband'] = t
                elif t.startswith('Symboler och innebörd:'):
                    fields['symboler'] = t.split(':', 1)[1].strip()
                elif t.startswith('Förutsättningar:'):
                    fields['givet'] = t.split(':', 1)[1].strip()
                elif t.startswith('Arbetsgång:'):
                    fields['arbetsgang'] = t.split(':', 1)[1].strip()
            fac = re.search(r'<details class="facit">.*?<p>(.*?)</p>', body, flags=re.S)
            fields['facit'] = inline(fac.group(1)) if fac else None
            out[aid] = fields
    return out


def kurs_kapitelrubrik():
    out = {}
    for vecka in range(39, 46):
        s = git_text(f'vecka-{vecka}/aktuell/Formelstod_och_ovningar.html')
        for kid, body in re.findall(r'<section class="chapter" id="([^"]+)">(.*?)</section>', s, flags=re.S):
            out[kid] = {'titel': inline(re.search(r'<h2>(.*?)</h2>', body).group(1)),
                        'teori': [inline(h) for h in re.findall(r'<h3>(.*?)</h3>', body) if not inline(h).startswith('Övning')]}
    return out


# ---------------------------------------------------------------- presentationerna
def pptx_text(shape):
    from pptx.oxml.ns import qn
    paras = []
    for p in shape.text_frame.paragraphs:
        buf = []
        for r in p.runs:
            rpr = r._r.find(qn('a:rPr'))
            base = int(rpr.get('baseline', '0')) if rpr is not None else 0
            t = r.text
            buf.append(('_{%s}' % t) if base < 0 and t.strip() else ('^{%s}' % t) if base > 0 and t.strip() else t)
        paras.append(''.join(buf))
    return re.sub(r'\s+', ' ', ' / '.join(x for x in paras if x.strip())).strip()


def deck_ovningar():
    from pptx import Presentation
    out = {}
    for kid, path in DECK.items():
        prs = Presentation(io.BytesIO(git_bytes(path)))
        for i, s in enumerate(prs.slides, 1):
            texts = [(sh.name, pptx_text(sh), sh.shape_id) for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            if not texts:
                continue
            t0 = texts[0][1]
            m = re.fullmatch(r'Stöd till övning (\d+)', t0)
            if m:
                d = out.setdefault(f'{kid}-q{m.group(1)}', {})
                d['stod_bild'] = i
                rest = [t for _, t, _ in texts[1:] if not re.fullmatch(r'\d+', t)]
                d['stod_shapes'] = {sid: t for _, t, sid in texts[1:]}
                d['samband'] = rest[0] if rest else None
                for t in rest[1:]:
                    if t.startswith('Förutsättningar:'):
                        d['givet'] = t.split(':', 1)[1].strip()
                    elif t.startswith('Arbetsgång:'):
                        d['arbetsgang'] = t.split(':', 1)[1].strip()
                    else:
                        d['symboler'] = t
                continue
            m = re.fullmatch(r'Övning (\d+)', t0)
            if m:
                d = out.setdefault(f'{kid}-q{m.group(1)}', {})
                d['ovning_bild'] = i
                rest = [t for _, t, _ in texts[1:] if not re.fullmatch(r'\d+', t)]
                d['ovning_shapes'] = {sid: t for _, t, sid in texts[1:]}
                d['titel'] = rest[0] if rest else None
                d['fraga'] = rest[1] if len(rest) > 1 else None
                d['instruktion'] = rest[2] if len(rest) > 2 else None
    return out


# ---------------------------------------------------------------- boken (EPUB)
def bok_dir():
    d = os.environ.get('BOK_EPUB_DIR')
    if not d or not Path(d, 'EPUB/text/ch001.xhtml').exists():
        raise SystemExit('Ange BOK_EPUB_DIR till en uppackad, dekrypterad bok.epub (se innehall/README.md).')
    return Path(d, 'EPUB/text')


def _rad_html(body, cls):
    m = re.search(r'<div class="rad rad-%s">\s*<span class="lbl">[^<]*</span>(.*?)</div>' % cls, body, flags=re.S)
    return m.group(1).strip() if m else None


def _rad(body, cls):
    h = _rad_html(body, cls)
    return inline(h) if h is not None else None


def _rad_rader(body, cls):
    """Raden uppdelad på radbrytningar (<br />), som i bokens Samband."""
    h = _rad_html(body, cls)
    return [inline(x) for x in re.split(r'<br\s*/?>', h) if inline(x)] if h is not None else None


def bok_ovningar():
    d = bok_dir()
    out = {}
    for k in range(1, 25):
        s = (d / BOK_KAPITEL_FIL[k]).read_text(encoding='utf-8')
        sol = (d / BOK_LOSNING_FIL[k]).read_text(encoding='utf-8')
        for n, body in re.findall(r'<section id="ch%d-q(\d+)" class="level3 uppgift"[^>]*>(.*?)</section>' % k, s, flags=re.S):
            n = int(n)
            h3 = re.search(r'<h3>(.*?)</h3>', body, flags=re.S).group(1)
            titel = inline(re.search(r'<span class="titel">(.*?)</span>', h3, flags=re.S).group(1))
            tag = re.search(r'data-tag="([a-z]+)"', s[s.find(f'id="ch{k}-q{n}"') - 5: s.find(f'id="ch{k}-q{n}"') + 200]).group(1)
            ps_html = re.findall(r'<p>(.*?)</p>', re.sub(r'<figure.*?</figure>', '', body, flags=re.S), flags=re.S)
            ps = [inline(p) for p in ps_html]
            figs = [{'src': m.group(2), 'klass': m.group(1), 'alt': _html.unescape(m.group(3)), 'nr': inline(m.group(4)) if m.group(4) else None,
                     'bildtext': inline(m.group(5)) if m.group(5) else None}
                    for m in re.finditer(r'<figure class="([^"]*)"[^>]*>\s*<img src="([^"]+)"[^>]*alt="([^"]*)"[^>]*/>\s*<figcaption[^>]*>(?:<span class="fignr">(.*?)</span>)?(.*?)</figcaption>', body, flags=re.S)]
            smb = re.search(r'<section id="sol%d-s%d" class="level3 losning">(.*?)</section>' % (k, n), sol, flags=re.S)
            steg, svar, kontroll, sfig, sol_html = [], None, None, [], None
            if smb:
                sb = smb.group(1)
                sb2 = re.sub(r'<div class="(?:rad|lank).*?</div>', '', sb, flags=re.S)
                sb2 = re.sub(r'<h3>.*?</h3>', '', sb2, flags=re.S)
                steg = [inline(x) for x in re.findall(r'<(?:p|li)>(.*?)</(?:p|li)>', sb2, flags=re.S)]
                svar = _rad(sb, 'svar')
                kontroll = _rad(sb, 'kontroll')
                sfig = [{'src': m.group(2), 'klass': m.group(1), 'alt': _html.unescape(m.group(3)), 'bildtext': inline(m.group(4)) if m.group(4) else None}
                        for m in re.finditer(r'<figure class="([^"]*)"[^>]*>\s*<img src="([^"]+)"[^>]*alt="([^"]*)"[^>]*/>\s*<figcaption[^>]*>(.*?)</figcaption>', sb, flags=re.S)]
                sol_html = sb
            out[f'{k}.{n}'] = {'kapitel': k, 'nr': n, 'titel': titel, 'niva': tag, 'fraga': ' '.join(ps), 'stycken': ps,
                               'samband': _rad(body, 'samband'), 'samband_rader': _rad_rader(body, 'samband'), 'givet': _rad(body, 'forutsattningar'), 'metod': _rad(body, 'metod'),
                               'figurer': figs, 'steg': steg, 'svar': svar, 'kontroll': kontroll, 'losningsfigurer': sfig,
                               'html': {'uppgift': body, 'losning': sol_html},
                               'id_uppgift': f'ch{k}-q{n}', 'id_losning': f'sol{k}-s{n}'}
    return out


# ---------------------------------------------------------------- alla presentationer: former per bild
ALLA_DECK = {
    'v37_01': 'vecka-37/aktuell/v37_01_Forberedelsefragor_1A_elev.pptx', 'v37_02': 'vecka-37/aktuell/v37_02_Elens_grunder_ombord_elev.pptx',
    'v37_03': 'vecka-37/aktuell/v37_03_Forberedelsefragor_1B_elev.pptx', 'v37_04': 'vecka-37/aktuell/v37_04_Elsakerhet_och_riskbedomning_elev.pptx',
    'v38_01': 'vecka-38/aktuell/v38_01_Franskiljning_och_matteknik_elev.pptx', 'v39_03': 'vecka-39/aktuell/v39_03_Multimeter_och_matfel_elev.pptx',
    'v39_04': 'vecka-39/aktuell/v39_04_Kirchhoff_seminarium_elev.pptx', **DECK,
}


def _platt(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _platt(sh.shapes)
        else:
            yield sh


def deck_former(kid):
    """{bildnummer: [(form-id, [stycken])]} för alla textformer, grupper uppplattade."""
    from pptx import Presentation
    prs = Presentation(io.BytesIO(git_bytes(ALLA_DECK[kid])))
    ut = {}
    for i, s in enumerate(prs.slides, 1):
        rad = []
        for sh in _platt(s.shapes):
            if sh.has_text_frame and sh.text_frame.text.strip():
                paras = []
                for p in sh.text_frame.paragraphs:
                    class _S:  # pptx_text tar en form; återanvänd för ett stycke
                        pass
                    f = _S()
                    f.text_frame = type('T', (), {'paragraphs': [p]})()
                    paras.append(pptx_text(f))
                rad.append((sh.shape_id, paras))
        ut[i] = rad
    return ut
