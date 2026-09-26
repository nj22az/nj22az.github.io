#!/usr/bin/env python3
"""Enhetlig notation: varje storhet skrivs på ett enda sätt i allt elevmaterial.

Den kanoniska formen och de avrådda varianterna står i beteckningar.json (fältet visa och avradda).
Ett index skrivs alltid med markering, X_{L}, och renderas nedsänkt (HTML <sub>, PowerPoint, SVG, canvas).
Låtsasindex (Xᴸ, Uꜰ, Iɴ) och ihopskrivna former (XL, Upp, Urms, U_{fas}) är inte tillåtna.

    python3 sjoskolan/innehall/notation.py kontrollera     lista avvikelser (körs också av innehall.py kontrollera)
    python3 sjoskolan/innehall/notation.py skriv-om        skriv om källor som inte genereras (sidtext, labbarnas och filmernas
                                                           texter, genomgången); poster skrivs om med --poster,
                                                           presentationerna (pptx) med --presentationer

Presentationerna kontrolleras via bildspelens data.json (texten med index), så att CI inte behöver python-pptx.
"""
import html
import json
import re
import sys
from pathlib import Path

ROT = Path(__file__).resolve().parent
SJO = ROT.parent
PSEUDO = {'ᴸ': 'L', 'ᶜ': 'C', 'ᴿ': 'R', 'ꜰ': 'F', 'ɴ': 'N'}
GRANS, SLUT = r'(?<![\wÅÄÖåäö{}_])', r'(?![\wÅÄÖåäö}])'

# Filer med elevtext som inte genereras ur databasen. HTML: bara textnoder utanför genererade regioner.
HTML = ['gemensamt/Formelblad_och_begrepp.html', 'gemensamt/Underlagskort.html', 'gemensamt/Bildgalleri.html', 'tentamen.html', 'index.html',
        'vecka-*/aktuell/*.html', 'vecka-*/aktuell/*/*.html', 'filmer/index.html', 'filmer/spela.html',
        'multimetersimulator/index.html', 'vaxelstromslabbet/index.html', 'trefaslabbet/index.html', 'hallkretslabbet/index.html', 'isolationslabbet/index.html']
# JavaScript med elevtext i strängar: bara strängar skrivs om, och en sträng som bara är ett ord (en nyckel som 'UF') lämnas.
JS = ['trefaslabbet/app.mjs', 'trefaslabbet/lessons.mjs', 'trefaslabbet/model.mjs', 'trefaslabbet/stationB-protokoll.mjs',
      'vaxelstromslabbet/app.mjs', 'vaxelstromslabbet/lessons.mjs', 'vaxelstromslabbet/model.mjs', 'vaxelstromslabbet/stationB-protokoll.mjs',
      'vaxelstromslabbet/guided.mjs', 'vaxelstromslabbet/equipment.mjs', 'vaxelstromslabbet/equipment-state.mjs',
      'hallkretslabbet/app.mjs', 'hallkretslabbet/stationC-protokoll.mjs', 'isolationslabbet/app.mjs', 'isolationslabbet/protokoll.mjs',
      'multimetersimulator/app.mjs', 'multimetersimulator/stationA-protokoll.mjs', 'gemensamt/labbprotokoll.mjs',
      '*/uppgifter.gen.mjs', 'vecka-40/aktuell/kontrollfragor.gen.mjs', 'gemensamt/*.gen.mjs',
      'verktyg/figurer/figs*.py',  # figurernas etiketter: Pythonsträngar läses som JS-strängar
      'filmer/films/*.mjs', 'vecka-40/aktuell/lektioner.mjs', 'vecka-40/aktuell/visuals.mjs', 'vecka-40/aktuell/lararstod.mjs']
BILDSPEL = ['bildspel/*/data.json']
PPTX = ['vecka-*/aktuell/*.pptx']
UNDANTAG_HTML = {'vecka-40/aktuell/Lararstod.html', 'vecka-41/aktuell/Simulerade_stationer_larare.html'}  # krypterade lärarsidor
# Postfält som är kod (nycklar och lägen), inte text.
KODFALT = {'simulator.fraga.storhet', 'simulator.fraga.falt', 'simulator.dolj', 'simulator.initial', 'simulator.validering', 'simulator.alias',
           'labb.steg.kod', 'labb.matningar.kod', 'labb.protokoll.lagen', 'labb.protokoll.rigg', 'labb.protokoll.nyckel', 'labb.protokoll.felsokning',
           'bok.html', 'bok.figurer.fil', 'bok.losningsfigurer.fil', 'referenser', 'parametrar', 'id', 'larare.facit_rader.funktion'}


def regler():
    """[(regex, ersättning i markering, beskrivning)] ur beteckningar.json."""
    d = json.loads((ROT / 'beteckningar.json').read_text(encoding='utf-8'))
    ut = []
    for b in d['beteckningar']:
        kanon = b.get('kanon', b['visa'].split(',')[0].strip())
        for a in b.get('avradda', []):
            fran, _, till = a.partition('→')
            till = till or kanon
            ut.append((re.compile(GRANS + re.escape(fran) + SLUT), till, f'{fran} → {till}'))
    ut.sort(key=lambda r: -len(r[0].pattern))
    return ut


def pseudo_till_markering(s):
    return re.sub('([A-Za-zΔφ])([ᴸᶜᴿꜰɴ]+)', lambda m: m.group(1) + '_{' + ''.join(PSEUDO[c] for c in m.group(2)) + '}', s)


def kanonisera(s, rg):
    """Text med markering -> kanonisk text med markering."""
    s = pseudo_till_markering(s)
    for rx, till, _ in rg:
        s = rx.sub(lambda m: till, s)
    return s


def avvikelser(s, rg):
    s2 = pseudo_till_markering(s)
    ut = [f'låtsasindex {m.group(0)}' for m in re.finditer('[A-Za-zΔφ][ᴸᶜᴿꜰɴ]+', s)]
    for rx, till, besk in rg:
        ut += [besk for _ in rx.finditer(s2)]
    return ut


# ---------------------------------------------------------------- HTML: textnoder
def html_markering(s):
    return re.sub(r'<sub>([^<]*)</sub>', r'_{\1}', s)


def markering_html(s):
    return re.sub(r'_\{([^{}]*)\}', r'<sub>\1</sub>', s)


def html_textdelar(s):
    """Delar sidan i (är_text, bit). Skript, stilar, taggar och genererade regioner är inte text."""
    delar, pos = [], 0
    for m in re.finditer(r'<!-- innehall:start.*?<!-- innehall:slut [^>]*-->|<script.*?</script>|<style.*?</style>|<[^>]+>', s, flags=re.S):
        if m.start() > pos:
            delar.append((True, s[pos:m.start()]))
        delar.append((False, m.group(0)))
        pos = m.end()
    delar.append((True, s[pos:]))
    # <sub>…</sub> ska höra till texten, så att X<sub>L</sub> och XL behandlas lika: slå ihop text + <sub> + text
    ihop, buf = [], None
    i = 0
    while i < len(delar):
        t, b = delar[i]
        if not t and b == '<sub>' and i + 2 < len(delar) and delar[i + 2][1] == '</sub>' and delar[i + 1][0]:
            buf = (buf or '') + '<sub>' + delar[i + 1][1] + '</sub>'
            i += 3
            continue
        if t:
            buf = (buf or '') + b
        else:
            if buf is not None:
                ihop.append((True, buf))
                buf = None
            ihop.append((False, b))
        i += 1
    if buf is not None:
        ihop.append((True, buf))
    return ihop


def skriv_om_html(s, rg):
    ut = []
    for t, b in html_textdelar(s):
        if t:
            n = markering_html(kanonisera(html_markering(html.unescape(b)) if '&' not in b else html_markering(b), rg))
            ut.append(n)
        else:
            ut.append(b)
    return ''.join(ut)


def html_avvikelser(s, rg):
    return [a for t, b in html_textdelar(s) if t for a in avvikelser(html_markering(b), rg)]


# ---------------------------------------------------------------- JavaScript: strängar
STRANG = re.compile(r"'(?:[^'\\\n]|\\.)*'|`(?:[^`\\]|\\.)*`|\"(?:[^\"\\\n]|\\.)*\"")


def ar_nyckel(lit):
    inner = lit[1:-1]
    return re.fullmatch(r'[A-Za-z0-9_]+', inner) is not None


def skriv_om_js(s, rg):
    def ers(m):
        lit = m.group(0)
        if ar_nyckel(lit) and not re.search('[ᴸᶜᴿꜰɴ]', lit):
            return lit
        if lit.startswith('`'):
            # mallsträngar: skriv inte om inuti ${…}
            delar = re.split(r'(\$\{[^}]*\})', lit)
            return ''.join(d if d.startswith('${') else kanonisera(d, rg) for d in delar)
        return kanonisera(lit, rg)
    return STRANG.sub(ers, s)


def js_avvikelser(s, rg):
    ut = []
    for m in STRANG.finditer(s):
        lit = m.group(0)
        if ar_nyckel(lit) and not re.search('[ᴸᶜᴿꜰɴ]', lit):
            continue
        delar = re.split(r'(\$\{[^}]*\})', lit) if lit.startswith('`') else [lit]
        for d in delar:
            if not d.startswith('${'):
                ut += avvikelser(d, rg)
    return ut


# ---------------------------------------------------------------- presentationer
def strangar(o):
    if isinstance(o, str):
        yield o
    elif isinstance(o, dict):
        for v in o.values():
            yield from strangar(v)
    elif isinstance(o, list):
        for v in o:
            yield from strangar(v)


def pptx_stycken(prs):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def former(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from former(sh.shapes)
            elif sh.has_text_frame:
                yield from sh.text_frame.paragraphs
            elif getattr(sh, 'has_table', False) and sh.has_table:
                for rad in sh.table.rows:
                    for c in rad.cells:
                        yield from c.text_frame.paragraphs
    for bild in prs.slides:
        yield from former(bild.shapes)


def skriv_om_pptx(fil, rg):
    """Skriver om presentationens text. Varje körning skrivs om för sig och delas vid index, så att formateringen
    behålls. Går en avvikelse över flera körningar (U + nedsänkt fas) skrivs stycket om med första körningens formatering."""
    import copy
    from pptx import Presentation
    from pptx.oxml.ns import qn
    sys.path[:0] = [str(ROT / 'lib'), str(ROT / 'export')]
    import presentationer as P
    prs = Presentation(str(fil))
    andrat = False
    for p in pptx_stycken(prs):
        fore = P.stycke_text(p)
        if kanonisera(fore, rg) == fore:
            continue
        for r in list(p.runs):
            rpr = r._r.find(qn('a:rPr'))
            if rpr is not None and int(rpr.get('baseline', '0')) != 0:
                continue
            ny = kanonisera(r.text, rg)
            if ny == r.text:
                continue
            bitar = [('sub', d[2:-1]) if d.startswith('_{') else ('text', d) for d in re.split(r'(_\{[^{}]*\})', ny) if d]
            for typ, t in reversed(bitar):  # addnext lägger direkt efter r, därför baklänges
                k = copy.deepcopy(r._r)
                kr = k.find(qn('a:rPr'))
                if kr is None:
                    kr = k.makeelement(qn('a:rPr'), {})
                    k.insert(0, kr)
                if typ == 'sub':
                    kr.set('baseline', '-25000')
                k.find(qn('a:t')).text = t
                r._r.addnext(k)
            r._r.getparent().remove(r._r)
        if kanonisera(P.stycke_text(p), rg) != P.stycke_text(p):
            P.satt_stycke(p, kanonisera(P.stycke_text(p), rg))
        andrat = True
    if andrat:
        prs.save(str(fil))
    return andrat


# ---------------------------------------------------------------- poster
def textfalt(o, vag=''):
    if any(vag == k or vag.startswith(k + '.') for k in KODFALT):
        return
    if isinstance(o, str):
        yield vag, o
    elif isinstance(o, dict):
        for k, v in o.items():
            yield from textfalt(v, f'{vag}.{k}' if vag else k)
    elif isinstance(o, list):
        for v in o:
            yield from textfalt(v, vag)


def skriv_om_post(o, rg, vag=''):
    if any(vag == k or vag.startswith(k + '.') for k in KODFALT):
        return o
    if isinstance(o, str):
        return kanonisera(o, rg)
    if isinstance(o, dict):
        return {k: skriv_om_post(v, rg, f'{vag}.{k}' if vag else k) for k, v in o.items()}
    if isinstance(o, list):
        return [skriv_om_post(v, rg, vag) for v in o]
    return o


def filer(monster):
    ut = []
    for m in monster:
        ut += sorted(p for p in SJO.glob(m) if p.is_file())
    return ut


def ordlistefel():
    """En avrådd form får inte samtidigt vara en godkänd form (former) eller den kanoniska."""
    d = json.loads((ROT / 'beteckningar.json').read_text(encoding='utf-8'))
    godkanda = {f for b in d['beteckningar'] for f in b.get('former', []) + [b.get('kanon', b['visa'].split(',')[0].strip())]}
    return [('beteckningar.json', f'{b["id"]}: {a.partition("→")[0]} är både avrådd och godkänd')
            for b in d['beteckningar'] for a in b.get('avradda', []) if a.partition('→')[0] in godkanda]


def kontrollera(poster=None):
    """Alla avvikelser som (fil eller id, beskrivning). poster: dict id -> post (katalogens)."""
    rg = regler()
    fel = ordlistefel()
    for f in filer(HTML):
        rel = str(f.relative_to(SJO))
        if rel in UNDANTAG_HTML or '/arkiv/' in rel:
            continue
        s = f.read_text(encoding='utf-8')
        if 'lock-data' in s:
            continue  # krypterad sida
        fel += [(rel, a) for a in html_avvikelser(s, rg)]
    for f in filer(JS):
        fel += [(str(f.relative_to(SJO)), a) for a in js_avvikelser(f.read_text(encoding='utf-8'), rg)]
    for f in filer(BILDSPEL):
        d = json.loads(f.read_text(encoding='utf-8'))
        fel += [(str(f.relative_to(SJO)), a) for t in strangar(d) for a in avvikelser(t, rg)]
    for i, p in (poster or {}).items():
        for vag, t in textfalt(p):
            fel += [(f'{i} {vag}', a) for a in avvikelser(t, rg)]
    return fel


def main():
    cmd = sys.argv[1] if len(sys.argv) > 1 else 'kontrollera'
    rg = regler()
    if cmd == 'skriv-om':
        n = 0
        for f in filer(HTML):
            rel = str(f.relative_to(SJO))
            if rel in UNDANTAG_HTML or '/arkiv/' in rel:
                continue
            s = f.read_text(encoding='utf-8')
            if 'lock-data' in s:
                continue
            ny = skriv_om_html(s, rg)
            if ny != s:
                f.write_text(ny, encoding='utf-8')
                n += 1
                print('html', rel)
        for f in filer(JS):
            s = f.read_text(encoding='utf-8')
            ny = skriv_om_js(s, rg)
            if ny != s:
                f.write_text(ny, encoding='utf-8')
                n += 1
                print('js  ', f.relative_to(SJO))
        if '--presentationer' in sys.argv:
            for f in filer(PPTX):
                if skriv_om_pptx(f, rg):
                    n += 1
                    print('pptx', f.relative_to(SJO))
        if '--poster' in sys.argv:
            for f in sorted((ROT / 'ovningar').glob('EL-*.json')):
                p = json.loads(f.read_text(encoding='utf-8'))
                ny = skriv_om_post(p, rg)
                if ny != p:
                    f.write_text(json.dumps(ny, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
                    n += 1
            for f in sorted((ROT / '.skyddat').glob('*.json')):
                d = json.loads(f.read_text(encoding='utf-8'))
                ny = {**d, 'poster': {i: skriv_om_post(p, rg) for i, p in d['poster'].items()}, 'falt': {i: skriv_om_post(p, rg) for i, p in d['falt'].items()}}
                if ny != d:
                    f.write_text(json.dumps(ny, ensure_ascii=False, sort_keys=True, indent=1), encoding='utf-8')
                    print('skyddat', f.name)
        print(f'{n} filer omskrivna')
    else:
        sys.path.insert(0, str(ROT / 'lib'))
        import katalog as K
        kat = K.Katalog()
        fel = kontrollera(kat.poster)
        for plats, a in fel:
            print(f'{plats}: {a}')
        print(f'{len(fel)} avvikelser')
        return 1 if fel else 0


if __name__ == '__main__':
    sys.exit(main() or 0)
