"""Vecka 41–45: en sida per del, som vecka 40. Översikt → Teori → Exempel → Övningar → Labben.

Teori och exempel kommer ur presentationen (samma text och figurer som läraren visar), övningarna ur databasen (ytan
kurs-formelstod). Varje övning har en ledtrådstrappa: ledtråd 1 visar sambandet och vilket teoriavsnitt som förklarar
det med ett uttryckligt teorikort, ledtråd 2 hur man börjar, ledtråd 3 exemplet som räknar samma sorts uppgift med andra tal. Facit sist. Figurerna
ligger i vecka-NN/aktuell/figurer (verktyg/veckosidor/figurer_ur_presentationer.py).
"""
import html
import json
import re

from pptx import Presentation

import rendera as R
from delsidor import CSS as CSS40, hjalp

PUBLIK = 'elev'
V = '20261011a'
VECKOR = range(41, 46)
# Vecka 45:s delar undervisas i vecka 43 och 44 (verktyg/veckosidor/bygg.py).
# Filerna ligger kvar i vecka-45/; sidan visar veckan och delen eleven ser. Övningarna behåller sina nummer.
FLYTTAD = {(45, 1): (43, 4, 4), (45, 2): (44, 4, 5), (45, 3): (44, 5, 5)}
EMU = 914400
HOPPA = ('Dagens två pass', 'Pass 2')
# Lektionsbilderna (verktyg/lektionsbilder, från vecka 42) har egna bilder för lektionens upplägg. De är inte teori.
LEKTIONSBILDER = ('Kursplanen i dag', 'Startfråga', 'Så går lektionen till', 'Övningar tillsammans', 'Paus', 'Sammanfattning',
                  'Det här gör du nu', 'Veckans labb')
# Lektionsbilderna har ingen bild med lösningsmallen; delsidan visar samma mall som de äldre presentationerna.
METOD = [('formel', '1. Givet och sökt'), ('p', 'Skriv av värden, enheter och vad du ska bestämma.'),
         ('formel', '2. Samband och villkor'), ('p', 'Välj formel eller princip. Kontrollera koppling, driftfall och antaganden.'),
         ('formel', '3. Uträkning och kontroll'), ('p', 'Omvandla enheter, sätt in värden och tolka resultatet.'),
         ('p', 'I resonemangsfall: koppla observation till en motiverad slutsats.')]
_FIGURHASH = {}


def figur_med_samma_innehall(figmapp, kap, blob):
    """Figurfilen i vecka-NN/aktuell/figurer med exakt samma innehåll som bilden i presentationen. Används när filnamnet
    (kap-bNN-k) inte finns, till exempel när lektionsbilderna har flyttat figuren till en annan bild."""
    import hashlib
    nyckel = (str(figmapp), kap)
    if nyckel not in _FIGURHASH:
        _FIGURHASH[nyckel] = {hashlib.sha256(f.read_bytes()).hexdigest(): f.name for f in sorted(figmapp.glob(f'{kap}-b*'))}
    return _FIGURHASH[nyckel].get(hashlib.sha256(blob).hexdigest())
CSS = CSS40.replace('/* Vecka 40, en sida per del: Översikt, Teori, Exempel, Övningar, Labben (genereras av innehall/export/delsidor.py) */',
                    '/* En sida per del: Översikt, Teori, Exempel, Övningar, Labben (innehall/export/delsidor_veckor.py; vecka 40 har egen kopia) */') + '''
/* Läsbarhet (Zenn): löptext med radavstånd 1,8. */
.del-sida p,.del-sida li,.del-sida dd{line-height:1.8}
.del-figur{margin:14px 0 18px}.del-figur img{display:block;width:100%;max-width:460px;height:auto;border:1px solid var(--sj-line,#cad8e2);border-radius:8px;background:#fff}
.del-figur figcaption{font-size:14px;color:var(--sj-muted,#4d6579);margin-top:4px}
.del-formel{font-size:21px;font-weight:700;color:var(--sj-accent,#064f91);margin:14px 0}
.del-teori{margin:18px 0 30px;padding-bottom:10px;border-bottom:1px solid var(--sj-line,#cad8e2)}
.del-exempel{margin:18px 0 30px;padding:10px 18px;border-left:8px solid var(--sj-accent,#064f91);background:var(--sj-soft,#edf4f9);border-radius:0 12px 12px 0}
.del-start{padding:10px 16px;border:1px dashed var(--sj-accent,#064f91);border-radius:10px}
.del-ovning .formula{white-space:normal;overflow-wrap:anywhere}
/* Kort för teori och exempel: rubriken och formeln syns, texten öppnas ett avsnitt i taget */
.del-verktyg{margin:6px 0 14px}.del-alla{min-height:44px;padding:6px 14px;border:1px solid var(--sj-line,#cad8e2);border-radius:8px;background:#fff;font:inherit;font-weight:700;color:var(--sj-accent,#064f91);cursor:pointer}
details.del-kort{margin:10px 0;padding:0;border:1px solid var(--sj-line,#cad8e2);border-radius:12px;background:#fff}
details.del-kort.del-exempel{border-left:8px solid var(--sj-accent,#064f91)}
.course-reading details.del-kort>summary{display:flex;align-items:flex-start;gap:12px;padding:12px 16px;list-style:none;min-height:44px;font-weight:400}
details.del-kort>summary::-webkit-details-marker{display:none}
details.del-kort>summary::after{content:"Öppna";margin-left:auto;align-self:center;padding-left:8px;font-size:14px;font-weight:700;color:var(--sj-accent,#064f91);white-space:nowrap}
details.del-kort[open]>summary::after{content:"Stäng"}
details.del-kort[open]>summary{border-bottom:1px solid var(--sj-line,#cad8e2)}
.kort-nr{flex:none;display:inline-grid;place-items:center;min-width:32px;height:32px;padding:0 6px;border-radius:16px;background:var(--sj-soft,#edf4f9);color:var(--sj-accent,#064f91);font-weight:700}
.kort-text{display:flex;flex-direction:column;gap:4px;min-width:0}.kort-titel{font-weight:700;font-size:19px;line-height:1.3}
.kort-formel{font-weight:700;color:var(--sj-accent,#064f91);font-size:16px;overflow-wrap:anywhere}
.kort-kropp{padding:4px 18px 12px}
.del-ova{margin:16px 0 4px;padding:10px 14px;border-radius:10px;background:var(--sj-soft,#edf4f9)}
@media print{details.del-kort>summary::after{content:none}.steg-knappar{display:none}}
.del-kontroll{margin:28px 0;padding:14px 18px;border:2px solid var(--sj-accent,#064f91);border-radius:12px;background:var(--sj-soft,#edf4f9)}.del-kontroll h3{margin-top:0}
.kf{margin:14px 0;padding:12px 14px;background:#fff;border:1px solid var(--sj-line,#cad8e2);border-radius:10px}.kf fieldset{border:0;margin:0;padding:0}.kf legend{padding:0;margin-bottom:6px}
.kf label{display:flex;gap:10px;align-items:flex-start;min-height:40px;padding:6px 4px;cursor:pointer}.kf input{margin-top:5px}.kf .kf-led{font-size:14px;color:var(--sj-muted,#4d6579)}
.kf-svara{min-height:44px;padding:6px 16px;border:1px solid var(--sj-accent,#064f91);border-radius:8px;background:var(--sj-accent,#064f91);color:#fff;font:inherit;font-weight:700;cursor:pointer}
.kf.ratt{border-color:var(--sj-ok,#176844)}.kf.fel{border-color:#9a330f}.kf.ratt .kf-utfall{color:var(--sj-ok,#176844);font-weight:700}.kf.fel .kf-utfall{color:#9a330f;font-weight:700}
.steg-knappar{margin:10px 0 4px;display:flex;flex-wrap:wrap;gap:8px}
'''


# Interaktiva sidor som hör till ett teoriavsnitt: (presentation, avsnittets rubrik, länk, text).
INTERAKTIV = [
    ('v41_01', 'Tre sinusspänningar', '../../generatorn/', 'Vrid rotorn själv: Generatorn i 3D'),
]


def e(s):
    return html.escape(str(s), quote=False)


def run_text(r):
    t = r.text
    b = r.font._rPr.get('baseline') if r.font._rPr is not None else None
    if b and int(b) < 0 and t.strip():
        return '_{' + t + '}'
    return t


def las_bild(bild, deck, nr, figmapp):
    """Bildens rubrik och innehåll i läsordning: ('p', text) | ('formel', text) | ('fig', fil, alt)."""
    former = sorted(bild.shapes, key=lambda s: (s.top, s.left))
    titel, block, k, forra = '', [], 0, None
    for sh in former:
        if sh.shape_type is not None and 'PICTURE' in str(sh.shape_type):
            if sh.top > 6.4 * EMU and sh.width > 9 * EMU:
                continue
            k += 1
            fil = f'{deck[:6]}-b{nr:02d}-{k}.{sh.image.ext}'
            if not (figmapp / fil).exists():
                fil = figur_med_samma_innehall(figmapp, deck[:6], sh.image.blob) or fil
            if (figmapp / fil).exists():
                alt = sh._element.nvPicPr.cNvPr.get('descr') or re.sub(r'^Figur:\s*', '', sh.name).rstrip(';').strip() or 'Figur'
                block.append(('fig', fil, alt))
            continue
        # Dagens namn på titelbilden (lektionsbilder ”Dag: Måndag 12 oktober”) är ingen rubrik.
        if not sh.has_text_frame or sh.top > 6.9 * EMU or sh.name.startswith('Dag: '):
            continue
        for para in sh.text_frame.paragraphs:
            t = ''.join(run_text(r) for r in para.runs).replace('\xa0', ' ').strip()
            if not t or re.fullmatch(r'\d+', t):
                continue
            if not titel:
                titel = t
                continue
            rs = [r for r in para.runs if r.text.strip()]
            formel = rs and all(r.font.bold for r in rs) and all(r.font.color and r.font.color.type and str(r.font.color.rgb) == '064F91' for r in rs)
            typ = 'formel' if formel else 'p'
            if typ == 'p' and block and block[-1][0] == 'p' and forra is sh and not re.search(r'[.!?:]$', block[-1][1]):
                block[-1] = ('p', block[-1][1] + ' ' + t)
            else:
                block.append((typ, t))
            forra = sh
    return titel, block


def block_html(block, nummerlista=False):
    ut, lista = [], []
    for b in block:
        if nummerlista and b[0] == 'p' and re.match(r'^\d+\.\s', b[1]):
            punkt = re.sub(r'^\d+\.\s*', '', b[1])
            lista.append(f'<li>{R.h(punkt)}</li>')
            continue
        if lista:
            ut.append(f'<ol>{"".join(lista)}</ol>')
            lista = []
        if b[0] == 'fig':
            ut.append(f'<figure class="del-figur"><img src="figurer/{b[1]}" alt="{R.attr(b[2])}" loading="lazy"></figure>')
        elif b[0] == 'formel':
            ut.append(f'<p class="del-formel">{R.h(b[1])}</p>')
        else:
            ut.append(f'<p>{R.h(b[1])}</p>')
    if lista:
        ut.append(f'<ol>{"".join(lista)}</ol>')
    return ''.join(ut)


def teori_for(pl, teori):
    """Slå upp placeringens exakta teorirubrik inom delen. Gissa aldrig vid saknad/tvetydig referens.

    Rubriken är nyckeln, inte bild- eller kortnumret: flyttade bilder får nya ankarlänkar automatiskt.
    En ändrad rubrik kräver ett uttryckligt beslut i placeringen och stoppar annars bygg/kontrollera.
    """
    var = f'{pl["plats"]} / {pl["del"]} / {pl["ovning_id"]}'
    titel = pl.get('teorikort')
    if not isinstance(titel, str) or not titel.strip():
        raise ValueError(f'{var}: saknar teorikort i placeringar/kurs-formelstod.json')
    match = [(i, t) for i, (_, t, _) in enumerate(teori, 1) if t == titel]
    if len(match) != 1:
        raise ValueError(f'{var}: teorikort {titel!r} måste matcha exakt ett teoriavsnitt i delen (fann {len(match)})')
    return match[0]


def veckans_delar(vv):
    """Delarna eleven ser i vecka vv, i ordning: [(mappvecka, filens delnummer, delnummer i veckan)].
    Veckans egna presentationer först, sedan delarna som flyttats dit från en annan veckas mapp (FLYTTAD)."""
    egna = [(vv, d, d) for d in sorted(int(x.name[5]) for x in (R.SJO / f'vecka-{vv}' / 'aktuell').glob('v*_elev.pptx'))
            if (vv, d) not in FLYTTAD]
    flyttade = sorted(((fv, d, t[1]) for (fv, d), t in FLYTTAD.items() if t[0] == vv), key=lambda x: x[2])
    return egna + flyttade


def deltitel_for(fv, d):
    """Rubriken på första bilden i delens presentation (samma som delsidans rubrik)."""
    mapp = R.SJO / f'vecka-{fv}' / 'aktuell'
    deck = next(mapp.glob(f'v{fv}_0{d}_*_elev.pptx'))
    return las_bild(Presentation(deck).slides[0], deck.name, 1, mapp / 'figurer')[0]


def sida(a, v, deck, alla_decks, pl_ovn, pres_bild, veckotitel, kontroll=()):
    mapp = R.SJO / f'vecka-{v}' / 'aktuell'
    figmapp = mapp / 'figurer'
    kap = deck[:6]
    dnr = int(kap[-1])
    prs = Presentation(mapp / deck)
    bilder = [(nr, *las_bild(b, deck, nr, figmapp)) for nr, b in enumerate(prs.slides, 1)]
    deltitel = bilder[0][1]
    mal = next((b for b in bilder if b[1].startswith('Mål')), None)
    metod = next((b for b in bilder if b[1].startswith('En lösning som går att följa')), None)
    avslut = next((b for b in bilder if b[1].startswith('Avslut')), None)
    kallor = next((b for b in bilder if b[1].startswith('Källor')), None)
    exempel = [b for b in bilder if b[1].startswith('Exempel')]
    lektionsbilder = any(b[1] == 'Kursplanen i dag' for b in bilder)
    if lektionsbilder and not metod:
        metod = (0, 'En lösning som går att följa', METOD)
    teori = [b for b in bilder[1:] if b not in (mal, metod, avslut, kallor) and b not in exempel and not b[1].startswith(HOPPA + LEKTIONSBILDER)
             and not re.match(r'^(Stöd till övning|Övning) \d', b[1])]
    ovn = sorted(pl_ovn.get(kap, []), key=lambda x: x[0]['ordning'])
    poster = [(pl, p) for pl, p in ovn]
    kort = hjalp([' '.join([p['uppgift']['fraga'], *p['uppgift'].get('samband', []), p['uppgift'].get('givet', '')]) for _, p in poster])

    teorireferenser = {pl['ankare']: teori_for(pl, teori) for pl, _ in poster}

    titel = f'Del {dnr}: {deltitel}'
    flytt = FLYTTAD.get((int(v), dnr))
    vv, idx, inl = (flytt[0], f'../../vecka-{flytt[0]}/aktuell/index.html', f'../../vecka-{flytt[0]}/aktuell/Inlamning.html') if flytt else (v, 'index.html', 'Inlamning.html')
    if flytt:
        titel = deltitel
    # Delarna eleven ser i veckan, också de som ligger i en annan veckas mapp (vecka 43: del 1–4, vecka 44: del 1–5).
    delar = veckans_delar(int(vv))
    antal = len(delar)
    plats = next(i for i, (fv, d, _) in enumerate(delar) if (fv, d) == (int(v), dnr))
    kicker = f'Vecka {vv} · Del {delar[plats][2]} av {antal}'
    flyttnotis = (f'<p class="sj-muted">Övningarna i den här delen är numrerade {dnr}.1–{dnr}.10.</p>' if flytt else '')
    h = [f'''<!doctype html><html lang="sv"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{e(titel)} · Vecka {vv} · Sjöskolan</title><meta name="description" content="{R.attr(veckotitel)}: {R.attr(deltitel)}. Teori, exempel och övningar med ledtrådar och facit."><link rel="canonical" href="https://nj22az.github.io/sjoskolan/vecka-{v}/aktuell/Del_{dnr}.html"><link rel="icon" href="/assets/images/apple-touch-icon.png"><link rel="stylesheet" href="/sjoskolan/gemensamt/sjoskolan.css?v=20260926"><link rel="stylesheet" href="/sjoskolan/course.css?v=20260928"><link rel="stylesheet" href="../../gemensamt/delsida.css?v={V}"><link rel="stylesheet" href="../../gemensamt/beteckningar.css?v=20260928"><link rel="stylesheet" href="../../gemensamt/raknehjalp.css?v=20260929"><script src="/sjoskolan/gemensamt/oversattning.js?v=20260927" defer></script></head>
<body class="course"><nav class="school-nav" aria-label="Sjöskolan"><a href="/sjoskolan/"><strong>SJÖSKOLAN</strong></a><a href="{idx}">Vecka {vv}</a><a href="{inl}">Inlämning</a></nav>
<main id="main-content" class="course-main del-main"><div class="course-breadcrumb"><a href="{idx}">← Vecka {vv}</a></div>
<article class="course-reading del-sida">
<p class="course-kicker">{kicker}</p><h1>{e(titel)}</h1>{flyttnotis}
<nav class="del-steg" aria-label="Delens ordning"><a href="#oversikt"><b>1</b>Översikt</a><a href="#teori"><b>2</b>Teori</a><a href="#exempel"><b>3</b>Exempel</a><a href="#ovningar"><b>4</b>Övningar</a><a href="#labb"><b>5</b>Labben</a></nav>''']
    # Vecka som visas som dagar (verktyg/veckosidor/dagsvy.py): rutan Din del med dagens färg, presentationen, läxorna
    # efter lektionen och det som gäller hela veckan, samma data som veckosidan.
    if not flytt:
        h.append(dagsvy().delpanel_html(int(v), dnr, sjalvrattning_modul()))

    # 1 Översikt
    malrader = [b for b in (mal[2] if mal else []) if b[0] != 'fig']
    start = next((i for i, b in enumerate(malrader) if b[1].lower().startswith('startfråga')), None)
    kunna = malrader[:start] if start is not None else malrader
    fraga = malrader[start + 1:] if start is not None else []
    if lektionsbilder and not mal:
        # ”Efter lektionen kan du” på bilden Kursplanen i dag och frågan på bilden Startfråga.
        kp = next(b[2] for b in bilder if b[1] == 'Kursplanen i dag')
        efter = next((i for i, b in enumerate(kp) if b[1] == 'Efter lektionen kan du'), None)
        kunna = [('p', re.sub(r'^\d+\.\s*', '', b[1])) for b in kp[efter + 1:] if b[0] != 'fig'] if efter is not None else []
        sf = next((b[2] for b in bilder if b[1] == 'Startfråga'), [])
        fraga = [b for b in sf if b[0] != 'fig'][:1]
    h.append(f'''<section id="oversikt" class="del-block"><h2><span>1</span> Översikt</h2>
<div class="sj-panel soft"><h3>Det här ska du kunna</h3><ul>{''.join(f'<li>{R.h(b[1])}</li>' for b in kunna)}</ul></div>
{f'<p class="del-start"><strong>Fundera först:</strong> {" ".join(R.h(b[1]) for b in fraga)} Svaret finns i teorin nedan.</p>' if fraga else ''}
<h3>Så gör du</h3><ol class="del-gor">
<li><strong>Läs teorin ett avsnitt i taget</strong> ({len(teori)} korta avsnitt). Sist i varje avsnitt står vilka övningar du kan göra direkt.</li>
<li><strong>Följ exemplen</strong> med papper och räknare, steg för steg. Gissa nästa steg innan du läser det.</li>
<li><strong>Gör övningarna</strong> {dnr}.1–{dnr}.{len(poster)}. Fastnar du: <em>Ledtråd 1</em> visar vilket samband du behöver och var i teorin det förklaras. <em>Ledtråd 2</em> visar hur du börjar och <em>Ledtråd 3</em> exemplet som räknar samma sorts uppgift med andra tal. Öppna facit när du har ett eget svar.</li>
<li><strong>Lämna in</strong> veckans uppgifter: <a href="Inlamning.html">Veckans inlämning</a>.</li></ol>''')
    if metod:
        h.append(f'<details class="sj-panel"><summary><strong>Så skriver du en lösning som går att följa</strong></summary>{block_html(metod[2])}</details>')
    h.append(f'<p class="del-mer"><a href="../../bildspel/?d={e(deck[:-5].removesuffix("_elev"))}">Lärarens presentation</a> · <a href="{e(deck[:-5])}.pdf">PDF</a> · <a href="../../gemensamt/Raknehjalp.html">Räknarhjälp</a> · <a href="../../gemensamt/Formelblad_och_begrepp.html">Formelblad</a></p>')
    texter = [t for _, t, bl in bilder for t in [t, *[x[1] for x in bl if x[0] == 'p']]] + [p['uppgift']['fraga'] for _, p in poster]
    formler = [x[1] for _, _, bl in bilder for x in bl if x[0] == 'formel'] + [x for _, p in poster for x in p['uppgift'].get('samband', [])]
    h.append(R.beteckningar_html(R.beteckningar_i(texter, a.beteckningar(), formler), 'Förkortningar och beteckningar i den här delen', trefas=True))
    h.append('</section>')

    # 2 Teori: ett kort per avsnitt. Rubriken och avsnittets formel syns direkt, texten öppnas ett avsnitt i taget,
    # och varje avsnitt slutar med de övningar som bygger på det (läs lite, öva direkt).
    ova = {}
    for pl, p in poster:
        ref = teorireferenser[pl['ankare']]
        if ref:
            ova.setdefault(ref[0], []).append(pl)
    h.append('<section id="teori" class="del-block"><h2><span>2</span> Teori</h2><p>Ett avsnitt i taget: öppna det, läs, och gör sedan övningarna som står sist i avsnittet. '
             f'Formeln bredvid rubriken är det viktigaste i avsnittet.</p><p class="del-verktyg"><button type="button" class="del-alla" data-mal="teori" aria-pressed="false">Öppna alla {len(teori)} avsnitt</button></p>')
    for i, (nr, t, bl) in enumerate(teori, 1):
        formel = next((x[1] for x in bl if x[0] == 'formel'), '')
        forhand = f'<span class="kort-formel">{R.h(formel)}</span>' if formel else ''
        lank = ', '.join(f'<a href="#{x["ankare"]}">{e(x["nummer"])}</a>' for x in ova.get(i, []))
        ovning = f'<p class="del-ova"><strong>Öva nu:</strong> {lank}</p>' if lank else ''
        interaktiv = ''.join(f'<p class="del-mer"><a href="{e(href)}">{e(txt)}</a></p>' for d_, t_, href, txt in INTERAKTIV if deck.startswith(d_) and t == t_)
        h.append(f'<details class="del-kort del-teori" id="teori-{i}"{" open" if i == 1 else ""}><summary><span class="kort-nr">{i}</span><span class="kort-text"><span class="kort-titel">{R.h(t)}</span>{forhand}</span></summary>'
                 f'<div class="kort-kropp">{block_html(bl)}{interaktiv}{ovning}</div></details>')
    h.append('</section>')

    # 3 Exempel: samma kort. Det första är öppet.
    h.append('<section id="exempel" class="del-block"><h2><span>3</span> Exempel</h2><p>Följ exemplet steg för steg med papper och räknare. Försök själv med nästa steg innan du läser det. Övningarna använder samma metod med andra tal.</p>')
    for i, (nr, t, bl) in enumerate(exempel, 1):
        h.append(f'<details class="del-kort del-exempel" id="exempel-{i}"{" open" if i == 1 else ""}><summary><span class="kort-nr">E{i}</span><span class="kort-text"><span class="kort-titel">{R.h((lambda x: x[:1].upper() + x[1:])(t.removeprefix("Exempel:").strip()))}</span></span></summary>'
                 f'<div class="kort-kropp">{block_html(bl, nummerlista=True)}</div></details>')
    h.append('</section>')

    # 4 Övningar
    nyckel = f'sj-ovningar:/sjoskolan/vecka-{v}/aktuell/Formelstod_och_ovningar.html'
    h.append('<section id="ovningar" class="del-block"><h2><span>4</span> Övningar</h2><p>Räkna på papper. Bocka av en övning när du har kontrollerat ditt svar mot facit. Avbockningen sparas i den här webbläsaren.</p><p id="del-status" class="del-status" aria-live="polite"></p>')
    for k, ((pl, p), rh) in enumerate(zip(poster, kort), 1):
        u, los = p['uppgift'], p.get('losning') or {}
        pass_ = pl['ordning'] > 5
        ref = teorireferenser[pl['ankare']]
        formel = ' och '.join(f'<span class="formula">{R.h(x)}</span>' for x in u.get('samband', []))
        resonemang = p.get('typ') == 'resonemang'
        led1 = ((f'Utgå från {formel}. ' if resonemang else f'Använd {formel}. ') if formel else '') + (f'Läs <a href="#teori-{ref[0]}">Teori {ref[0]}: {R.h(ref[1])}</a>.' if ref else '')
        led2 = ' '.join(R.h(x['text']) for x in p.get('ledtradar', []) if x.get('niva') in ('begrepp', 'metod'))
        ex = exempel[1 if pass_ and len(exempel) > 1 else 0] if exempel else None
        led3 = ((f'Samma sätt att resonera: <a href="#exempel-{exempel.index(ex) + 1}">{R.h(ex[1])}</a>. Gör likadant med fallet i den här uppgiften och motivera varje steg.' if resonemang
                 else f'Samma metod med andra tal: <a href="#exempel-{exempel.index(ex) + 1}">{R.h(ex[1])}</a>. Följ stegen där och gör sedan likadant med talen i den här uppgiften.') if ex else '')
        steg = ' '.join(R.h(x['text']) for x in p.get('ledtradar', []) if x.get('niva') == 'nasta-steg')
        if steg:  # ledtråd 3 ur databasen: metoden steg för steg med andra tal
            led3 = steg
        figs = [x for x in (pres_bild.get(p['id']) or []) if x[0] == 'fig']
        facit = ''.join(f'<p>{R.h(x)}</p>' for x in str(los.get('text', 'Facit saknas.')).split('\n\n') if x.strip())
        svar = ' · '.join(f'{R.h(sv["storhet"])} ≈ {e(R.svarsvarde(sv))} {e(sv.get("enhet", ""))}' for sv in los.get('svar', []) if isinstance(sv.get('varde'), (int, float)))
        h.append(f'''<article class="del-ovning" id="{pl["ankare"]}"><h3>{e(pl["nummer"])}: {R.h(p["titel"])}</h3>
{f'<p>{R.h(u["scenario"])}</p>' if u.get("scenario") else ''}<p>{R.h(u["fraga"])}</p>{f'<p class="del-givet"><strong>Givet:</strong> {R.h(u["givet"])}</p>' if u.get("givet") else ''}
{block_html(figs)}
{f'<details class="ledtrad"><summary>Ledtråd 1: {"vilken princip?" if resonemang else "vilket samband?"}</summary><p>{led1}</p></details>' if led1 else ''}
{f'<details class="ledtrad"><summary>Ledtråd 2: hur börjar jag?</summary><p>{led2}</p></details>' if led2 else ''}
{f'<details class="ledtrad"><summary>Ledtråd 3: {("steg för steg" if resonemang else "steg för steg med andra tal") if steg else "samma metod i ett exempel"}</summary><p>{led3}</p></details>' if led3 else ''}
{rh}
<details class="facit"><summary>Facit</summary>{facit}{f'<p><strong>Svar:</strong> {svar}</p>' if svar else ''}</details>
<button type="button" class="del-klar" data-id="{pl["ankare"]}">Markera som klar</button></article>''')
    if kontroll:  # kontrollfrågor (placeringar/kontrollfragor.json): rätt alternativ står i losning.kommentar
        fr = []
        for pl, p in kontroll:
            alt, ratt = p['uppgift']['alternativ'], p['losning']['kommentar']
            val = ''.join(f'<label><input type="radio" name="{pl["ankare"]}" value="{i}"> {R.h(x)}</label>' for i, x in enumerate(alt))
            led = ' '.join(R.h(x['text']) for x in p.get('ledtradar', []))
            fr.append(f'<form class="kf" id="{pl["ankare"]}" data-ratt="{alt.index(ratt)}"><fieldset><legend><b>{e(pl["nummer"])}.</b> {R.h(p["uppgift"]["fraga"])}</legend>{val}</fieldset>'
                      f'<p class="kf-knapp"><button type="submit" class="kf-svara">Svara</button> <span class="kf-led">Ledtråd: {led}</span></p>'
                      f'<div class="kf-svar" hidden><p class="kf-utfall"></p><p>{R.h(p["losning"]["text"])}</p></div></form>')
        h.append(f'<section class="del-kontroll" id="kontroll"><h3>Kontrollfrågor: har du förstått delen?</h3><p>Svara på de tre frågorna innan du går vidare. Ditt första svar sparas och följer med i QR-koden på sidan <a href="Resultat.html">Skicka resultat</a>. Det är ingen poäng, det visar vad som behöver repeteras.</p>{"".join(fr)}</section>')
    if avslut:
        h.append(f'<section class="del-tur" id="klart"><h3>Klart med delen?</h3>{block_html(avslut[2])}</section>')
    h.append('</section>')

    # 5 Labben
    h.append(f'<section id="labb" class="del-block"><h2><span>5</span> Labben</h2><p>Labben kommer sist i veckan, när du har gått igenom del 1–{antal}. Du använder samma metod som i övningarna: <strong>räkna först, läs av, jämför och förklara</strong>.</p>'
             f'<p><a class="sj-btn" href="{idx}#labb">Veckans labbar</a> <a class="sj-btn" href="{inl}">Veckans inlämning</a></p>')
    if kallor:
        h.append(f'<details class="del-mer"><summary>Källor och fortsatt läsning</summary>{block_html(kallor[2])}</details>')
    h.append('</section>')

    kf_nyckel = f'sj-kontroll:vecka-{v}'
    def lank(i, pil):
        # Länk till del nummer i (0-baserat) i veckans ordning. En del i en annan mapp får också sin rubrik.
        fv, d, n = delar[i]
        if fv == int(v):
            return f'<a href="Del_{d}.html">{"← " if pil == "<" else ""}Del {n}{" →" if pil == ">" else ""}</a>'
        text = f'Del {n}: {e(deltitel_for(fv, d))}'
        return f'<a href="../../vecka-{fv}/aktuell/Del_{d}.html">{"← " + text if pil == "<" else text + " →"}</a>'
    fore = lank(plats - 1, '<') if plats > 0 else f'<a href="{idx}">← Vecka {vv}</a>'
    efter = lank(plats + 1, '>') if plats + 1 < antal else f'<a href="{idx}#labb">Labben →</a>'
    # Vecka med självrättande inlämning: sist på sidan bekräftar eleven att hen har läst delen (sj-last:vecka-NN).
    if not flytt:
        h.append(sjalvrattning().lasbekraftelse_html(int(v), dnr))
    h.append(f'''<nav class="art-nav" aria-label="Delar"><span>{fore}</span><span><a href="{idx}">Veckans översikt</a></span><span>{efter}</span></nav>
</article></main><footer class="school-nav">Sjöskolan · Elteknik och ellära · Nils Johansson</footer>
<script>
// Avbockade övningar sparas per vecka i den här webbläsaren.
(()=>{{const K='{nyckel}';let klara={{}};try{{klara=JSON.parse(localStorage.getItem(K)||'{{}}')||{{}};}}catch{{}}
const knappar=[...document.querySelectorAll('.del-klar')];
const rita=()=>{{let n=0;for(const b of knappar){{const k=!!klara[b.dataset.id];n+=k;b.textContent=k?'✓ Klar':'Markera som klar';b.setAttribute('aria-pressed',String(k));b.closest('.del-ovning').classList.toggle('klar',k);}}document.getElementById('del-status').textContent=`${{n}} av ${{knappar.length}} övningar klara.`;}};
for(const b of knappar)b.addEventListener('click',()=>{{klara[b.dataset.id]=!klara[b.dataset.id];try{{localStorage.setItem(K,JSON.stringify(klara));}}catch{{}}rita();}});
rita();}})();
// Kontrollfrågor: första svaret sparas per fråga ({{ankare: valt alternativ}}) och följer med i resultatkoden.
(()=>{{const K='{kf_nyckel}';let sv={{}};try{{sv=JSON.parse(localStorage.getItem(K)||'{{}}')||{{}};}}catch{{}}
const visa=(f,val)=>{{const ok=val===Number(f.dataset.ratt);const u=f.querySelector('.kf-svar');u.hidden=false;u.querySelector('.kf-utfall').textContent=ok?'Rätt.':'Inte rätt. Rätt svar: '+f.querySelectorAll('label')[f.dataset.ratt].textContent.trim()+'.';f.classList.toggle('ratt',ok);f.classList.toggle('fel',!ok);}};
for(const f of document.querySelectorAll('form.kf')){{if(Number.isInteger(sv[f.id])){{const r=f.querySelector(`input[value="${{sv[f.id]}}"]`);if(r)r.checked=true;visa(f,sv[f.id]);}}
f.addEventListener('submit',ev=>{{ev.preventDefault();const r=f.querySelector('input:checked');if(!r)return;const val=Number(r.value);if(!Number.isInteger(sv[f.id])){{sv[f.id]=val;try{{localStorage.setItem(K,JSON.stringify(sv));}}catch{{}}}}visa(f,val);}});}}}})();
// Teori och exempel är kort som öppnas ett i taget. En länk dit (ledtråd, Öva nu, stegraden) öppnar kortet.
(()=>{{const oppna=h=>{{if(!h||h.length<2)return;const m=document.getElementById(decodeURIComponent(h.slice(1)));if(!m)return;for(let d=m;d;d=d.parentElement?.closest('details'))if(d.tagName==='DETAILS')d.open=true;}};
document.addEventListener('click',ev=>{{const a=ev.target.closest('a[href^="#"]');if(a)oppna(a.getAttribute('href'));}});
addEventListener('hashchange',()=>oppna(location.hash));oppna(location.hash);
// Exemplen visas ett steg i taget: eleven försöker själv innan nästa steg visas. Utan JavaScript syns alla steg.
for(const ol of document.querySelectorAll('.del-exempel ol')){{const steg=[...ol.children];if(steg.length<2)continue;const efter=[...ol.parentElement.children].slice([...ol.parentElement.children].indexOf(ol)+1);
let n=1;const p=document.createElement('p');p.className='steg-knappar';const nasta=document.createElement('button'),alla=document.createElement('button');nasta.type=alla.type='button';nasta.className='sj-btn primary';alla.className='sj-btn';alla.textContent='Visa alla steg';
const visa=()=>{{steg.forEach((li,i)=>li.hidden=i>=n);efter.forEach(x=>x.hidden=n<steg.length);nasta.hidden=alla.hidden=n>=steg.length;nasta.textContent=`Visa steg ${{n+1}} av ${{steg.length}}`;}};
nasta.onclick=()=>{{n++;visa();}};alla.onclick=()=>{{n=steg.length;visa();}};p.append(nasta,' ',alla);ol.after(p);visa();}}
for(const b of document.querySelectorAll('.del-alla'))b.addEventListener('click',()=>{{const kort=[...document.querySelectorAll('#'+b.dataset.mal+' .del-kort')];const alla=!kort.every(d=>d.open);kort.forEach(d=>d.open=alla);b.setAttribute('aria-pressed',String(alla));b.textContent=alla?'Stäng alla avsnitt':'Öppna alla '+kort.length+' avsnitt';}});}})();
</script></body></html>
''')
    return ''.join(h)


_SJALV = []


def sjalvrattning():
    """verktyg/inlamning/sjalvrattning.py (veckorna med självrättande inlämning och rutan Har du läst hela del N?)."""
    if not _SJALV:
        import importlib.util
        spec = importlib.util.spec_from_file_location('sjalvrattning_delsidor', R.SJO / 'verktyg' / 'inlamning' / 'sjalvrattning.py')
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        _SJALV.append(mod)
    return _SJALV[0]


_DAGSVY = []


def dagsvy():
    """verktyg/veckosidor/dagsvy.py: rutan Din del på delsidorna för veckorna i DAGSVY."""
    if not _DAGSVY:
        import importlib.util
        spec = importlib.util.spec_from_file_location('dagsvy_delsidor', R.SJO / 'verktyg' / 'veckosidor' / 'dagsvy.py')
        mod = importlib.util.module_from_spec(spec)
        import sys
        sokvag = list(sys.path)
        sys.path.insert(0, str(R.SJO / 'verktyg' / 'veckosidor'))  # dagsvy.py importerar kursplan.py bredvid sig
        try:
            spec.loader.exec_module(mod)
        finally:
            sys.path[:] = sokvag
        _DAGSVY.append(mod)
    return _DAGSVY[0]


def sjalvrattning_modul():
    """Adressen till gemensamt/sjalvrattning.mjs relativt dagsvy.mjs, med innehållets hash (samma som veckosidan)."""
    import hashlib
    modul = R.SJO / 'gemensamt' / 'sjalvrattning.mjs'
    return f'./sjalvrattning.mjs?v={hashlib.sha1(modul.read_bytes()).hexdigest()[:10]}' if modul.exists() else None


def filer(a):
    ut = {'gemensamt/delsida.css': CSS}
    placeringar = json.loads((R.SJO / 'innehall' / 'placeringar' / 'presentation.json').read_text())['placeringar']
    for v in VECKOR:
        mapp = R.SJO / f'vecka-{v}' / 'aktuell'
        decks = sorted(x.name for x in mapp.glob('v*_elev.pptx'))
        pl_ovn = {}
        for pl, p in a.placeringar('kurs-formelstod', f'vecka-{v}/aktuell/Formelstod_och_ovningar.html'):
            pl_ovn.setdefault(pl['del'], []).append((pl, p))
        # Övningens figur: bilden i presentationen där övningen står.
        pres_bild = {}
        for pl in placeringar:
            if pl['plats'].startswith(f'vecka-{v}/') and pl.get('bild') and not pl.get('kontroll'):
                deck = pl['plats'].split('/')[-1]
                bild = Presentation(mapp / deck).slides[pl['bild'] - 1]
                pres_bild[pl['ovning']] = las_bild(bild, deck, pl['bild'], mapp / 'figurer')[1]
        veckotitel = f'Vecka {v}'
        kf = {}
        for pl, p in a.placeringar('kontrollfragor'):
            if pl['plats'].startswith(f'vecka-{v}/'):
                kf.setdefault(pl['del'], []).append((pl, p))
        for deck in decks:
            ut[f'vecka-{v}/aktuell/Del_{int(deck[5])}.html'] = sida(a, v, deck, decks, pl_ovn, pres_bild, veckotitel, kf.get(deck[:6], ()))
        if kf:  # kontrollfrågornas rätta svar för lärarsidan med resultatkoder (larare/resultat-vNN.html)
            rader = [{'id': pl['ankare'], 'nummer': pl['nummer'], 'fraga': p['uppgift']['fraga'], 'alternativ': p['uppgift']['alternativ'],
                      'ratt': p['uppgift']['alternativ'].index(p['losning']['kommentar'])} for d in sorted(kf) for pl, p in kf[d]]
            ut[f'vecka-{v}/aktuell/kontrollfragor.gen.mjs'] = ('// GENERERAD FIL · ur sjoskolan/innehall (placeringar/kontrollfragor.json, innehall.py bygg delsidor_veckor).\n'
                                                              '// Redigera posterna i innehall/ovningar/, inte här. Samma frågor och rätta svar som på delsidorna.\n'
                                                              'export const KONTROLLFRAGOR = ' + json.dumps(rader, ensure_ascii=False, indent=1) + ';\n')
    return ut
