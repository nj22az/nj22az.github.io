"""Presentationerna: skriver övningstexterna på bilderna ur databasen (ytorna presentation och presentation-fragor),
gör om PDF:erna med LibreOffice och bygger om webbildspelen.

Tung exportör: körs med `innehall.py bygg presentationer`. Bara presentationer vars text skiljer sig från databasen skrivs om.
Bildernas formatering (typsnitt, storlek, färg) behålls: första körningen i varje textform används som mall,
index X_{L} blir nedsänkta körningar. Former märkta kontroll=true skrivs inte om (text i figurer); de jämförs bara.
"""
import copy
import re
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'migrering'))
import rendera as R
import text as T

PUBLIK = 'elev'
YTOR = ('presentation', 'presentation-fragor')


def falt_text(pl, p, namn):
    """Texten som ska stå i formen `namn` (en lista med stycken)."""
    u, roll = p['uppgift'], pl.get('roll')
    led = {l['niva']: l['text'] for l in p.get('ledtradar', [])}
    if namn == 'titel':
        if pl['del'] == 'v39_04':
            return [f'{pl["nummer"]}. {p["titel"]}']
        if pl['del'] == 'v37_04':
            return [f'{pl["nummer"]}: {p["titel"]}']
        return [p['titel']]
    if namn == 'fraga':
        return [u['fraga']]
    if namn == 'kortfraga':
        k = u.get('kortfraga') or u['fraga']
        return [f'{pl["nummer"]}: {k}'] if pl['del'] == 'v37_02' else [k]
    if namn == 'scenario':
        return [u.get('scenario', '')]
    if namn == 'instruktion':
        return [u.get('instruktion', '')]
    if namn == 'stod/samband':
        return list(u.get('samband', []))
    if namn == 'stod/begrepp':
        return [led.get('begrepp', '')]
    if namn == 'stod/givet':
        return [f'Förutsättningar: {u.get("givet", "")}']
    if namn == 'stod/metod':
        return [f'Arbetsgång: {led.get("metod", "")}']
    if namn == 'stod/losning':
        return [(p.get('losning') or {}).get('text', '')]
    raise ValueError(namn)


def _former(slide):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    ut = {}

    def g(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                g(sh.shapes)
            elif sh.has_text_frame:
                ut[sh.shape_id] = sh
    g(slide.shapes)
    return ut


def stycke_text(p):
    """Ett styckes text med markering, som kallor.pptx_text läser det."""
    from pptx.oxml.ns import qn
    buf = []
    for r in p.runs:
        rpr = r._r.find(qn('a:rPr'))
        base = int(rpr.get('baseline', '0')) if rpr is not None else 0
        t = r.text
        buf.append(('_{%s}' % t) if base < 0 and t.strip() else ('^{%s}' % t) if base > 0 and t.strip() else t)
    return re.sub(r'\s+', ' ', ''.join(buf)).strip()


def satt_stycke(p, text):
    """Skriver text med markering till ett stycke och behåller första körningens formatering."""
    from pptx.oxml.ns import qn
    runs = p.runs
    if not runs:
        p.text = T.ren_text(text)
        return
    mall = copy.deepcopy(runs[0]._r)
    for r in runs:
        r._r.getparent().remove(r._r)
    # a:endParaRPr ska ligga sist; körningar läggs före den
    end = p._p.find(qn('a:endParaRPr'))
    for typ, t in T.delar(text):
        r = copy.deepcopy(mall)
        rpr = r.find(qn('a:rPr'))
        if rpr is None:
            rpr = r.makeelement(qn('a:rPr'), {})
            r.insert(0, rpr)
        if typ == 'sub':
            rpr.set('baseline', '-25000')
        elif typ == 'sup':
            rpr.set('baseline', '30000')
        elif 'baseline' in rpr.attrib:
            del rpr.attrib['baseline']
        tt = r.find(qn('a:t'))
        tt.text = t
        if end is not None:
            end.addprevious(r)
        else:
            p._p.append(r)


def satt_form(sh, stycken, bara_stycke=None):
    """Skriver stycken till en form. bara_stycke: skriv bara stycke nr n (1-baserat)."""
    tf = sh.text_frame
    paras = list(tf.paragraphs)
    if bara_stycke:
        satt_stycke(paras[bara_stycke - 1], stycken[0])
        return
    mall = copy.deepcopy(paras[0]._p)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    satt_stycke(paras[0], stycken[0])
    for s in stycken[1:]:
        ny = copy.deepcopy(mall)
        tf._txBody.append(ny)
        from pptx.text.text import _Paragraph
        satt_stycke(_Paragraph(ny, tf), s)


def form_text(sh, bara_stycke=None):
    paras = list(sh.text_frame.paragraphs)
    if bara_stycke:
        return [stycke_text(paras[bara_stycke - 1])]
    return [stycke_text(p) for p in paras if stycke_text(p)]


def skriv_deck(a, plats, placeringar):
    from pptx import Presentation
    fil = R.SJO / plats
    prs = Presentation(fil)
    slides = list(prs.slides)
    andrat, avvik = 0, []
    for pl, p in placeringar:
        former = pl.get('former') or {}
        # Kontrollformer: en text kan vara delad på flera former (fraga/1, fraga/2 …); delarna jämförs ihopslagna.
        if pl.get('kontroll'):
            grupper = {}
            for namn, ref in former.items():
                bas = re.sub(r'/\d+$', '', namn)
                grupper.setdefault(bas, []).append(ref)
            for bas, refs in grupper.items():
                bild = pl['stodbild'] if bas.startswith('stod/') else pl['bild']
                delar = []
                for ref in refs:
                    sh = _former(slides[bild - 1]).get(int(ref))
                    delar += form_text(sh) if sh is not None else ['(saknas)']
                nu, mal = T.normalisera(' '.join(delar)), T.normalisera(' '.join(x for x in falt_text(pl, p, bas) if x))
                if nu != mal:
                    avvik.append(f'{plats} bild {bild} ({p["id"]} {bas}): bilden har {nu!r}, databasen {mal!r}')
            continue
        for namn, ref in former.items():
            bild = pl['stodbild'] if namn.startswith('stod/') else pl['bild']
            sid, _, nr = ref.partition('/')
            sh = _former(slides[bild - 1]).get(int(sid))
            if sh is None:
                avvik.append(f'{plats} bild {bild}: form {sid} saknas')
                continue
            mal = [T.normalisera(y) for x in falt_text(pl, p, namn) for y in x.split(' / ') if y]
            nu = [T.normalisera(x) for x in form_text(sh, int(nr) if nr else None)]
            if nu == mal:
                continue
            satt_form(sh, mal, int(nr) if nr else None)
            andrat += 1
    if andrat:
        prs.save(fil)
    return andrat, avvik


def pdf_och_bildspel(filer):
    if not filer:
        return
    for f in filer:
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', str(f.parent), str(f)], check=True, capture_output=True)
    subprocess.run([sys.executable, str(R.SJO / 'verktyg' / 'bildspel' / 'bygg.py')], check=True, capture_output=True)


def bygg(a, kat):
    decks = {}
    for yta in YTOR:
        for plats in a.platser(yta):
            decks.setdefault(plats, []).extend(a.placeringar(yta, plats))
    omskrivna, avvik, rev = [], [], {}
    for plats, pls in sorted(decks.items()):
        n, av = skriv_deck(a, plats, pls)
        avvik += av
        if n:
            omskrivna.append(R.SJO / plats)
        for pl, p in pls:
            rev[p['id']] = p['revision']
    pdf_och_bildspel(omskrivna)
    for x in avvik:
        print('  obs', x)
    return {'sammanfattning': f'{len(decks)} presentationer, {len(omskrivna)} omskrivna (PDF och bildspel byggda om), {len(avvik)} avvikelser i kontrollformer',
            'omskrivna': [str(f.relative_to(R.SJO)) for f in omskrivna], 'avvikelser': avvik, 'revisioner': rev}
